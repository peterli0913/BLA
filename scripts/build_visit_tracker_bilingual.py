#!/usr/bin/env python3
"""Build bilingual and English-only copies of the Lilly visit action tracker."""

from __future__ import annotations

import copy
import math
import re
import shutil
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font

SRC = Path("/workspace/Asymchem_Visit_Action_Tracker_中英对照-分工更新-09-25-2026.xlsx")
OUT_BI = Path("/workspace/deliverables/Asymchem_Visit_Action_Tracker_中英对照_09-25-2026.xlsx")
OUT_EN = Path("/workspace/deliverables/Asymchem_Visit_Action_Tracker_EN_09-25-2026.xlsx")

CJK = re.compile(r"[\u3000-\u303f\u3400-\u9fff\uff00-\uffef]")

NAMES = {
    "李延龙": "Yanlong Li", "刘维江": "Weijiang Liu", "刘林冲": "Linchong Liu",
    "郭宏杰": "Hongjie Guo", "张俊峰": "Junfeng Zhang", "刘美佳": "Meijia Liu",
    "郑明智": "Mingzhi Zheng", "范兴涛": "Xingtao Fan", "王琪": "Qi Wang",
    "马辉": "Hui Ma", "胡春鹏": "Chunpeng Hu", "熊熊": "Xiong Xiong",
    "刘凯": "Kai Liu", "刘恺": "Kai Liu", "孙志强": "Zhiqiang Sun",
    "朱新义": "Xinyi Zhu", "张忠尧": "Zhongyao Zhang", "王帆": "Fan Wang",
    "杨明旸": "Mingyang Yang", "张文艺": "Wenyi Zhang", "于生辉": "Shenghui Yu",
    "张国睿": "Guorui Zhang", "陈明志": "Mingzhi Chen", "唐安垚": "Anyao Tang",
    "史永钢": "Yonggang Shi", "王志恒": "Zhiheng Wang",
}
DEPTS = {
    "生产": "Production", "工程": "Engineering", "QA": "QA", "IEPE": "IEPE",
    "项目组": "Project Team", "罐区": "Tank Farm", "生产设备部": "Production Equipment Dept.",
    "研发": "R&D",
}
CATEGORIES = {
    "清洁相关": "Cleaning", "其他": "Other", "设备改造": "Equipment modification",
    "物流通道": "Material logistics corridor", "纯化水相关": "Purified water (PW)",
    "DIPEA相关": "DIPEA", "钝化": "Passivation", "CCS": "CCS",
    "过滤策略相关": "Filtration strategy", "CNC湿度": "CNC humidity",
}

# (English, Chinese). Asymchem-authored English is lightly corrected for spelling and grammar;
# Lilly's own wording in columns C and L is left as written.
T: dict[str, tuple[str, str | None]] = {
    # headers
    "D1": ("Implementation Timing", "实施时机"),
    "F1": ("Proposed target completion date", "建议目标完成日期"),
    "G1": ("Existing approach", "现行做法"),
    "H1": ("Action Plan", "行动计划"),
    "I1": ("Questions", "问题"),
    "M1": ("Lead Owner", "牵头负责人"),
    "N1": ("QA Contact", "QA对接人"),
    "O1": ("Completion Date", "完成时间"),
    "P1": ("Asymchem Internal Category", "Asymchem内部分类"),
    # item 2
    "F3": ("Has been completed.", "已完成。"),
    # item 3
    "H4": ("Theoretical calculation by IEPE; on-site testing.", "IEPE理论计算；现场测试"),
    # item 4
    "G5": ("Modification and installation of the C3 and C4 lines are complete, and slope verification is in progress. "
           "The C1 and C2 lines were previously assessed as having no slope requirement, and only drain points were provided; "
           "the slope requirement needs to be re-discussed and confirmed.",
           "C3、C4管路进行改造安装完毕，正在进行坡度的确认中；C1、 C2管路之前评估没有坡度的要求，仅设置了排净点，坡度需要重新沟通确认。"),
    "H5": ("Assess and inspect each process step and each line individually.", "逐个步骤及逐条管路进行评估并进行排查；"),
    "I5": ("The requirement for the C1 and C2 lines is drainability, with no specific slope requirement. How should the slope be confirmed?",
           "C1、C2管路可排尽的要求，没有坡度的具体要求，这个坡度如何进行确认？"),
    # item 5
    "C6": ("Missing sight-glass lights", "缺少视灯"),
    "G6": ("The entire equipment train is being reviewed; any missing sight-glass lights will be procured.",
           "整体设备链进行确认，缺少的进行采购"),
    "H6": ("An inspection checklist is being established.", "正在建立排查清单"),
    # item 6
    "H7": ("Provide the material-of-construction certificate and adjust the current dip-tube length.",
           "提供材质证明，现有长度进行调整。"),
    # item 7
    "G8": ("No detailed requirement in the MOR or SOP.", "MOR 或 SOP 中无详细要求。"),
    "H8": ("Update the MOR to include the detailed requirement.", "更新 MOR，纳入详细要求。"),
    # item 8
    "H9": ("To be carried out together with the overall piping construction.", "同增体管路施工同步进行"),
    # item 9
    "H10": ("Submit the change document to carry out renovation on the factory logistics passage. "
            "An airlock will be added before raw materials are transferred into the CNC area.",
            "提交变更文件，对厂区物流通道进行改造；原料转入 CNC 区域前增设气闸。"),
    "I10": ("Prior to Pre-PPQ for which specific step? Step 1 starts on 10/05/2026, so this timing cannot be met.",
            "具体步骤的pre-PPQ前？Step1于10/05/2026启动，时间上无法满足"),
    # item 10
    "G11": ("One-month service life for dedicated PW hoses.", "专用 PW 软管使用寿命为一个月。"),
    "H11": ("Perform a use/disinfection-period study on dedicated PW hoses and redefine their lifecycle.",
            "开展专用 PW 软管使用/消毒周期研究，并重新定义其生命周期。"),
    # item 11
    "H12": ("Revise the SOP to define the corresponding requirements.", "更改SOP，明确对应要求"),
    # item 12
    "F13": ("To be confirmed by Linchong Liu", "刘林冲确认"),
    "H13": ("Before Step 3 Demo production, perform one rinse with process solvent and document it in the corresponding record.",
            "计划在step3 Demo生产前，使用工艺溶剂做一次工艺溶剂淋洗，并有对应记录。"),
    # item 13
    "F14": ("Domestic: 3 weeks\nImported: 2 months", "国产：3周\n进口：2月"),
    "G14": ("Platinum-cured silicone hoses are used.", "使用铂金硫化硅胶软管"),
    "H14": ("Replacement is planned.", "计划替换"),
    "I14": ("Confirm whether domestic hoses meet the requirements.", "确认国产是否满足要求"),
    # item 14
    "F15": ("09/30/2026 (after confirmation of analytical methods)", "09/30/2026（分析方法确认后）"),
    "G15": ("Supplier samples received. Preliminary testing passed (< 50 ppm). Waiting for analytical method alignment for final release.",
            "已收到供应商样品，初步检测合格（< 50 ppm）。待分析方法对齐后最终放行。"),
    "H15": ("Purchase (completed), then analyze.", "采购（已完成），随后检测。"),
    "I15": ("Analytical method TBD.", "分析方法待定。"),
    # item 15
    "G16": ("Lab-scale trials showed some effectiveness in reducing acetaldehyde.",
            "实验室规模试验显示对降低乙醛有一定效果。"),
    "H16": ("We will directly purchase the low-acetaldehyde DIPEA from the supplier; no further purification will be performed at Asymchem.",
            "将直接从供应商采购低乙醛 DIPEA，Asymchem 不再进行额外纯化。"),
    # item 16
    "G17": ("i) Experimental results show that a long washing time does not cause DIPEA decomposition. "
            "ii) Compared with the RS batch, the main difference in exposure time came from the smaller DIPEA packages used in the RS batch; "
            "frequent container changes extended the operation time. Larger DIPEA containers (200 L) will be used to reduce DIPEA exposure time.",
            "i) 实验结果表明，较长的洗涤时间不会导致 DIPEA 分解。ii) 与 RS 批相比，暴露时间的主要差异来自 RS 批使用的 DIPEA 包装较小，"
            "频繁更换容器延长了操作时间。将改用更大的 DIPEA 容器（200 L）以缩短 DIPEA 暴露时间。"),
    "H17": ("The DIPEA package will be changed to larger drums.", "DIPEA 包装将改为更大规格的桶。"),
    # item 17
    "G18": ("17 m³/h has been proven effective for HCN removal at the 22 kg batch scale, and the flow rate can be adjusted up to 30 m³/h.",
            "已证明 17 m³/h 可有效去除 22 kg 规模批次的 HCN，流量最高可调至 30 m³/h。"),
    "H18": ("Relevant documents will be updated accordingly.", "相关文件将相应更新。"),
    # item 18
    "G19": ("Agreement has been reached that sweeping/vacuum is not recommended.", "双方已达成一致：不建议采用吹扫/抽真空。"),
    # item 19
    "F20": ("For details, refer to the Notes for each project.", "详见各项目备注。"),
    "G20": ("No passivation on the inner surface of process lines.", "工艺管路内表面未进行钝化。"),
    "H20": ("Perform passivation of GG917 and LAARA process lines.", "对 GG917 与 LAARA 工艺管路进行钝化。"),
    "L20": ("L1 Step 1 Completed; Step 2: 09/29/2026\nL2 Step 1 Completed; Step 2: 09/29/2026\n"
            "L3 Step 1 Completed; Step 2: 09/29/2026; Step 3: 10/07/2026\nGG917 Step 1 Completed; Step 2: 10/02/2026; Step 3: Completed",
            "L1：Step 1 已完成；Step 2：09/29/2026\nL2：Step 1 已完成；Step 2：09/29/2026\n"
            "L3：Step 1 已完成；Step 2：09/29/2026；Step 3：10/07/2026\nGG917：Step 1 已完成；Step 2：10/02/2026；Step 3 已完成"),
    # item 20
    "G21": ("The Phase I tank farm will not be passivated for now; metal-ion testing is performed instead.",
            "一期罐区暂时不钝化,执行金属离子检测；"),
    "H21": ("After the Phase II tank farm is handed over, its lines and tanks will be passivated, and all subsequent projects will use solvents from the Phase II tank farm.",
            "待二期罐区交付后，进行管路和储罐钝化，后续项目全部使用二期罐区溶剂；"),
    "I21": ("None", "无"),
    # item 21
    "G22": ("There is an administrative requirement that, once material arrives at the production workshop, a dedicated operator must transfer it to the CNC area.",
            "现有管理要求：物料到达生产车间后，须由专人转运至 CNC 区域。"),
    "H22": ("Update the SOP to include detailed material lifecycle management to avoid gaps during material transfer.",
            "更新 SOP，纳入详细的物料生命周期管理，避免物料转移环节出现缺口。"),
    # item 22
    "H23": ("Submit a change control to add a buffer room.", "提交变更，增加缓冲间"),
    # item 23
    "G24": ("The CCS document lacks a detailed description.", "CCS 文件缺少详细描述。"),
    "H24": ("Add a detailed description to the CCS document.", "在 CCS 文件中补充详细描述。"),
    # item 24
    "G25": ("Silicone water hoses are currently used.", "目前使用硅胶水管。"),
    "H25": ("Replace silicone hoses with PTFE hoses.", "将硅胶软管更换为 PTFE 软管。"),
    # item 25
    "G26": ("No clear time requirement in the SOP for disconnecting water hoses after use.", "SOP 对用水软管用后断开无明确时间要求。"),
    "H26": ("Update the SOP to require that water hoses be disconnected immediately once water charging is finished.",
            "更新 SOP，要求加水完成后立即断开用水软管。"),
    # item 26
    "G27": ("The detailed filtration principle is not included in the SOP or CCS document.", "SOP 与 CCS 文件中未包含详细的过滤原则。"),
    "H27": ("Update the SOP and CCS document to include the detailed filtration principle.", "更新 SOP 与 CCS 文件，纳入详细的过滤原则。"),
    # item 27
    "G28": ("No detailed requirement or assessment for liquid material filtration.", "对液体物料过滤无详细要求和评估。"),
    "H28": ("Add a detailed description of liquid material filtration to the SOP and CCS.", "在 SOP 与 CCS 中补充液体物料过滤的详细描述。"),
    # item 28
    "G29": ("Water is used for glovebox cleaning.", "目前用水清洁手套箱。"),
    "H29": ("After water cleaning, an alcohol wipe shall be performed to avoid any water residue.", "水清洁后须用酒精擦拭，避免水残留。"),
    # item 29
    "H30": ("Modify the air handling units by adding cooling coils, and complete qualification during a production gap.",
            "对空调机组改造，增加表冷器，在生产间歇完成验证；"),
    "I30": ("The affected rooms must stop production during the air handling unit modification.", "空调机组改造期间，对应房间需要停产"),
    # item 30
    "F31": ("To be uploaded before the end of this week", "本周周末前上传"),
    # item 31
    "G32": ("No detailed assessment in the CCS document.", "CCS 文件中无详细评估。"),
    "H32": ("Add a detailed assessment to the CCS document.", "在 CCS 文件中补充详细评估。"),
    # item 32
    "G33": ("No permeable film is used to cover the ends of process and water hoses.", "目前未使用透气膜覆盖工艺及用水软管端部。"),
    "H33": ("Use permeable film to cover the ends of process and water hoses.", "使用透气膜覆盖工艺及用水软管端部。"),
    "N33": ("N/A", None),
    # item 33
    "F34": ("Completed", "已完成"),
    "G34": ("This capability is already in place.", "具备该条件"),
    # item 34
    "G35": ("The PD pump is rinse-sampled for residue testing.", "目前对 PD 泵采用淋洗取样进行残留检测。"),
    "H35": ("Update the SOP to include swab sampling for residue testing.", "更新 SOP，纳入残留检测的擦拭取样。"),
    # item 35
    "G36": ("After use, both ends of the pump are drained and blown dry; a detailed operating procedure will be provided.",
            "使用结束后进行泵两端进行排空吹干。并提供具体操作流程。"),
    "H36": ("The corresponding hygienic centrifugal pumps are being procured.", "对应卫生级离心泵采购中"),
    # item 36
    "H37": ("Redesign is in progress. Once the design is confirmed, the key items will be discussed with the Lilly team before execution.",
            "重新设计中，待方案确认后重点与礼来团队讨论执行。"),
    "I37": ("Does the drainability design of the material lines need to meet 2D or 3D? For some tanks, the modification plan has been confirmed; "
            "we would like the Lilly team to help confirm these plans.",
            "物料管路排尽的设计是否满足要求2Dor3D；部分储罐已经确认改造方案的需要礼来团队帮忙确认；"),
    # item 37
    "F38": ("Same as Item 1.", "同第一条。"),
    # item 38
    "F39": ("Lead time: 6 weeks\nModification: 3 weeks\nPassivation: 1 week\nValidation: 3 weeks",
            "货期6周\n改造3周\n钝化1周\n验证3周"),
    "G39": ("Temporary hoses are used to connect water outlets to the points of use.", "使用临时水管从水点连接到使用点"),
    "I39": ("Completing this before Pre-PPQ is tight. Could it be completed after PPQ, with water used per the current practice in the meantime?",
            "pre-PPQ前完成较为紧张，是否在PPQ后进行，按照现有的执行方式进行用水。"),
    # item 39
    "G40": ("No requirement for filtration of purified water supplied by hose.", "对经软管供应的纯化水无过滤要求。"),
    "H40": ("Add filtration and filter replacement-interval requirements to the MOR.", "在 MOR 中增加过滤要求及更换周期要求。"),
    # item 40
    "F41": ("Same as row 39 (Item 38).", "同39行"),
    # item 41
    "F42": ("Zhiqiang Sun to prepare a PPT and assess whether swabbing is required.", "志强写个PPT，评估是否需要擦拭"),
    "G42": ("Equipment is disassembled for cleaning, and ARL performs Level 2 testing of the reflux solvent. "
            "The line from the activation tank to the synthesizer is reused after solvent flushing.",
            "设备进行拆卸清洗，ARL进行回流溶剂的level2检测；激活罐到合成仪的管路存在溶剂淌洗的方式进行复用。"),
    # item 42
    "H43": ("The SOP will be revised.", "将修订 SOP。"),
    # item 43
    "G44": ("Spray coverage testing is performed on the equipment.", "设备进行喷淋覆盖率的验证。"),
    # item 44
    "G45": ("Asymchem has an internal gowning procedure for classified areas. The gowning procedure needs to be revisited.",
            "Asymchem 已有洁净级别区域内部更衣程序，需重新审视该程序。"),
    "H45": ("Update the gowning procedure for classified areas.", "更新洁净级别区域更衣程序。"),
    # item 45
    "G46": ("Asymchem has an internal procedure for access control.", "Asymchem 已有门禁控制内部程序。"),
    "H46": ("Present the Asymchem access-control procedure to Lilly.", "向 Lilly 介绍 Asymchem 门禁控制程序。"),
    # item 47
    "G48": ("Conductivity is monitored downstream of the solution-preparation mixer: acceptable mobile phase is routed to the mobile-phase tank, "
            "and out-of-range mobile phase is diverted to waste. Solution preparation is metered by mass flow meters, so temperature changes have "
            "little effect on the metered composition before mixing. pH probes are installed at the mixer and on the receiving mobile-phase tank for monitoring.",
            "配液系统混合器后有电导监测响应，合格流动相会进流动相储罐，不在范围内会响应切换排废；配液采用的质量流量计计量，温度变化对其混合前组成的计量影响较小。"
            "混合器和接收后的流动相储罐安装pH，有监测功能。"),
    # item 48
    "G49": ("Asymchem has an internal procedure for filter integrity testing.", "Asymchem 已有过滤器完整性测试内部程序。"),
    "H49": ("Present the Asymchem filter integrity testing strategy to Lilly.", "向 Lilly 介绍 Asymchem 过滤器完整性测试策略。"),
    # item 49
    "F50": ("Timing to be provided by Zepeng Guo", "郭泽鹏反馈时间"),
    # item 51
    "G52": ("No requirement for the height of purified water hose ends above the floor.", "对纯化水软管端部离地高度无要求。"),
    "H52": ("Update the SOP to add a minimum height requirement of 6 inches.", "更新 SOP，增加离地至少 6 inches 的高度要求。"),
    # item 52
    "F53": ("To be provided by Linchong Liu", "刘林冲提供"),
    # item 54
    "F55": ("Timing to be provided by Zepeng Guo", "郭泽鹏反馈时间"),
    "H55": ("Install a sampling valve on the 100 L preparation vessel.", "100L制备桶配置取样阀"),
    # item 55
    "F56": ("Procurement lead time: 6 weeks", "采购周期6周，"),
}

SUMMARY_CN = {
    "A1": "BLA 来访行动项跟踪汇总", "A3": "指标", "B3": "数量", "D3": "优先级", "E3": "数量",
    "A4": "行动项总数", "A5": "高优先级", "A6": "未关闭", "A7": "进行中", "A8": "阻塞",
    "A9": "已完成", "D4": "高", "D5": "中", "D6": "低",
}
INSTRUCTIONS_CN = {
    "A1": "本跟踪表使用说明",
    "A3": "• 在建议负责人一栏指定具体责任人；如适用，替换建议的职能。",
    "A4": "• 为每个高优先级事项设定截止日期。",
    "A5": "• 使用“状态”下拉列表更新进展。",
    "A6": "• 使用表格筛选，按日期、区域、优先级、责任人或状态查看行动项。",
    "A7": "• 优先级与责任人为根据来访记录建议的组织字段，需由项目团队确认。",
}

# Original Summary used structured references that were already #REF! in the uploaded file.
SUMMARY_FORMULAS = {
    "B4": "=COUNTA('Action Tracker'!C2:C56)",
    "B5": "=COUNTIF('Action Tracker'!E2:E56,\"High*\")",
    "B6": "=COUNTIF('Action Tracker'!K2:K56,\"Open\")",
    "B7": "=COUNTIF('Action Tracker'!K2:K56,\"In Progress\")",
    "B8": "=COUNTIF('Action Tracker'!K2:K56,\"Blocked\")",
    "B9": "=COUNTIF('Action Tracker'!K2:K56,\"Completed\")",
    "E4": "=COUNTIF('Action Tracker'!E2:E56,\"High*\")",
    "E5": "=COUNTIF('Action Tracker'!E2:E56,\"Med*\")",
    "E6": "=COUNTIF('Action Tracker'!E2:E56,\"Low*\")",
}

NEUTRAL = re.compile(r"^(N/A|[\d/]+|(\s*(2D|3D|step\d[AB]?):?\s*[\d/]*\s*)+)$", re.I)


def has_cjk(s: str) -> bool:
    return bool(CJK.search(s))


def split_bilingual(text: str) -> tuple[str, str | None]:
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if has_cjk(line):
            if i == 0:
                raise ValueError(f"Chinese-only text needs a translation: {text!r}")
            return "\n".join(lines[:i]).rstrip(), "\n".join(lines[i:])
    return text, None


def owner_en(text: str) -> str:
    out = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if "/" in line:
            dept, people = line.split("/", 1)
            dept_en = DEPTS[dept.strip()]
        else:
            dept_en, people = None, line
        names = [n for n in re.split(r"[、，,]", people) if n.strip()]
        names_en = ", ".join(NAMES[n.strip()] for n in names) or "TBD"
        out.append(f"{dept_en} / {names_en}" if dept_en else names_en)
    return "\n".join(out)


def build_pairs(ws) -> dict[str, tuple[str, str | None]]:
    pairs: dict[str, tuple[str, str | None]] = {}
    missing = []
    for row in ws.iter_rows(min_row=1, max_row=56, max_col=16):
        for cell in row:
            v = cell.value
            if not isinstance(v, str):
                continue
            coord, col = cell.coordinate, cell.column_letter
            if coord in T:
                pairs[coord] = T[coord]
            elif col in ("M", "N") and cell.row > 1:
                if v.strip() == "N/A":
                    pairs[coord] = ("N/A", None)
                else:
                    pairs[coord] = (owner_en(v), v)
            elif col == "P" and cell.row > 1:
                en = CATEGORIES[v.strip()]
                pairs[coord] = (en, None if en == v.strip() else v)
            elif has_cjk(v):
                pairs[coord] = split_bilingual(v)
            else:
                if col in "FGHIL" and cell.row > 1 and not NEUTRAL.match(v.strip()):
                    missing.append((coord, v))
                pairs[coord] = (v, None)
    if missing:
        for m in missing:
            print("NEEDS CHINESE:", m, file=sys.stderr)
        raise SystemExit("untranslated English-only cells")
    return pairs


def text_lines(text: str, width: float) -> int:
    n = 0
    for para in str(text).split("\n"):
        units = sum(2.0 if has_cjk(ch) else 1.1 for ch in para)
        n += max(1, math.ceil(units / max(width - 1.5, 4)))
    return n


def fit_rows(ws, keep_taller: bool) -> None:
    widths = {k: (d.width or 9) for k, d in ws.column_dimensions.items()}
    hidden = {k for k, d in ws.column_dimensions.items() if d.hidden}
    for r in range(1, 57):
        need = 1
        for cell in ws[r][:16]:
            if cell.value is None or cell.column_letter in hidden:
                continue
            need = max(need, text_lines(cell.value, widths.get(cell.column_letter, 9)))
        height = need * 16.5 + 10
        cur = ws.row_dimensions[r].height or 15
        ws.row_dimensions[r].height = round(max(height, cur) if keep_taller else max(height, 30), 1)


def fix_status_cf(ws) -> None:
    fills = {}
    for rng, rules in list(ws.conditional_formatting._cf_rules.items()):
        if str(rng.sqref) == "K2:K56":
            for rule in rules:
                key = "Completed" if "Completed" in rule.formula[0] else "Blocked"
                fills[key] = rule.dxf
            del ws.conditional_formatting._cf_rules[rng]
    for key, dxf in fills.items():
        rule = FormulaRule(formula=[f'$K2="{key}"'])
        rule.dxf = dxf
        ws.conditional_formatting.add("K2:K56", rule)


def summary_values(ws) -> dict[str, int]:
    effort = [str(ws[f"E{r}"].value or "") for r in range(2, 57)]
    status = [str(ws[f"K{r}"].value or "") for r in range(2, 57)]
    high = sum(e.startswith("High") for e in effort)
    return {
        "B4": sum(1 for r in range(2, 57) if ws[f"C{r}"].value not in (None, "")),
        "B5": high,
        "B6": status.count("Open"),
        "B7": status.count("In Progress"),
        "B8": status.count("Blocked"),
        "B9": status.count("Completed"),
        "E4": high,
        "E5": sum(e.startswith("Med") for e in effort),
        "E6": sum(e.startswith("Low") for e in effort),
    }


def inject_cached_values(path: Path, sheet_xml: str, values: dict[str, int]) -> None:
    # openpyxl writes formulas without results; add them so previews that do not recalculate still show counts.
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == sheet_xml:
                xml = data.decode("utf-8")
                for coord, val in values.items():
                    xml, n = re.subn(
                        rf'(<c r="{coord}"[^>]*>)(<f>.*?</f>)(<v>.*?</v>)?(</c>)',
                        rf"\g<1>\g<2><v>{val}</v>\g<4>",
                        xml,
                        count=1,
                        flags=re.S,
                    )
                    if n != 1:
                        raise SystemExit(f"could not cache {coord}")
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def build(lang: str, out: Path) -> None:
    wb = load_workbook(SRC)
    ws = wb["Action Tracker"]
    pairs = build_pairs(ws)
    body_font = copy.copy(ws["G5"].font)

    for coord, (en, cn) in pairs.items():
        cell = ws[coord]
        cell.value = en if lang == "en" or not cn else f"{en}\n{cn}"
        if cell.column_letter in "MNOP" and cell.row > 1:
            cell.font = copy.copy(body_font)
        if not cell.alignment.wrap_text:
            al = cell.alignment
            cell.alignment = Alignment(horizontal=al.horizontal, vertical=al.vertical, wrap_text=True)

    table = ws.tables["BLAActionTracker"]
    for col in table.tableColumns:
        idx = table.tableColumns.index(col) + 2
        col.name = ws.cell(1, idx).value

    fix_status_cf(ws)
    fit_rows(ws, keep_taller=(lang == "bi"))

    summary = wb["Summary"]
    for coord, formula in SUMMARY_FORMULAS.items():
        summary[coord] = formula
    instructions = wb["Instructions"]
    if lang == "bi":
        for sheet, labels in ((summary, SUMMARY_CN), (instructions, INSTRUCTIONS_CN)):
            for coord, cn in labels.items():
                c = sheet[coord]
                c.value = f"{c.value}\n{cn}"
                al = copy.copy(c.alignment)
                c.alignment = Alignment(horizontal=al.horizontal, vertical=al.vertical or "center", wrap_text=True)
        for r in range(1, 10):
            summary.row_dimensions[r].height = 36 if summary[f"A{r}"].value or summary[f"D{r}"].value else None
        for r in (1, 3, 4, 5, 6, 7):
            instructions.row_dimensions[r].height = 36

    wb.calculation.fullCalcOnLoad = True
    values = summary_values(ws)
    summary_index = wb.sheetnames.index("Summary") + 1
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    inject_cached_values(out, f"xl/worksheets/sheet{summary_index}.xml", values)
    print("saved", out, values)


if __name__ == "__main__":
    build("bi", OUT_BI)
    build("en", OUT_EN)
