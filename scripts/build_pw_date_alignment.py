#!/usr/bin/env python3
"""Tracker vs. purified-water PPT date conflicts, for internal alignment before sending to Lilly.

Usage: python3 scripts/build_pw_date_alignment.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "纯水日期对齐_Tracker与PPT对照_09-25-2026.xlsx"

HEAD = ["Tracker 行", "行动项", "负责人", "Tracker 里", "纯水 PPT 里", "统一口径（待填）"]
ROWS = [
    (11, "论证专用 PW 软管一个月使用寿命的依据", "胡春鹏", "11/30，做周期研究", "10/20，每周清洁、每半年更换"),
    (25, "硅胶水管更换为 PTFE 软管或提供技术论证", "胡春鹏", "10/20", "国产 10/30 / 进口 11/30"),
    (26, "用水软管用后立即断开并受控存放", "胡春鹏", "10/04", "10/20"),
    (33, "软管干燥存放期间用透气膜覆盖端部", "胡春鹏", "10/04", "10/30"),
    (39, "纯化水总管改造（使用点排废冲洗、排尽干燥）", "胡春鹏、郭宏杰", "6+3+1+3 周",
     "6+3+2+3+2 周，01/30/2027"),
    (40, "经软管供水的使用点过滤器（NMT 24 h 更换）", "胡春鹏", "10/20", "09/30"),
    (41, "新建使用点排尽口代表性取样评估", "胡春鹏", "同 39 行", "10/30"),
    (52, "软管端部离地至少 6 inches 并设目视标识", "胡春鹏", "10/04", "09/30 / 11/30"),
]
WIDTHS = [11, 40, 16, 20, 28, 30]
NOTE = ("说明：39 行 Tracker 为 货期6周+改造3周+钝化1周+验证3周；PPT 为 采购6周+施工3周+酸洗钝化2周+"
        "连续取样验证并培养3周+管路干燥验证2周。")


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "日期对齐"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(wrap_text=True, vertical="center", horizontal="center")

    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor="2F5597")
        c.alignment = center
        c.border = border
    for row in ROWS:
        ws.append(list(row) + [""])
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in r:
            c.font = Font(name="微软雅黑", size=10.5)
            c.alignment = center if c.column in (1, 3) else wrap
            c.border = border
        r[5].fill = PatternFill("solid", fgColor="FFF2CC")
        ws.row_dimensions[r[0].row].height = 34
    ws.row_dimensions[1].height = 24

    note_row = ws.max_row + 2
    ws.cell(note_row, 1, NOTE).font = Font(name="微软雅黑", size=9.5, color="595959")
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(HEAD))
    ws.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[note_row].height = 32

    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(OUT)
    print(OUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
