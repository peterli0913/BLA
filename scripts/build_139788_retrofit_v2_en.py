#!/usr/bin/env python3
"""English version of 139788项目改造工作说明V2.pptx (retrofit options, equipment plan, piping schematics).

Logos and the "Confidential | 保密文件" / website footers stay as in the source; all content text becomes
English in Times New Roman. Chinese annotations inside the slide-4 CAD pictures are covered with editable
white-filled English text boxes.

Usage: python3 scripts/build_139788_retrofit_v2_en.py
"""
import copy
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "139788项目改造工作说明V2.pptx"
OUT = ROOT / "deliverables" / "139788项目改造工作说明V2_EN.pptx"

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
FONT = "Times New Roman"
RED = "FF0000"
# a:rPr children that must come after latin/ea/cs
RPR_AFTER_FONTS = {A + t for t in ("sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst")}

TITLES = {
    1: "139788 BLA Retrofit Plan",
    2: "Three Retrofit Options: Completion Dates, Pros and Cons",
    3: "Equipment and Tank Retrofit / Replacement Plan",
    4: "Piping Retrofit Options (Schematic)",
    5: "THANK YOU!",
}
OPTION_LABELS = {"方案一": "Option 1", "方案二": "Option 2", "方案三": "Option 3"}

# ---------------------------------------------------------------- slide 2: options table
# One list of paragraphs per cell; [[...]] marks text that is red in the source.
UPGRADE = ("All associated piping, instruments, valves, metering pumps, diaphragm pumps, etc. to be changed to "
           "centrifugal pumps, quaternary diaphragm pumps, instruments, valves, pH meters and flowmeters")
PRECIP = ("Precipitation area: replace the 80 L and 600 L agitated tanks with BPE-compliant 100 L and 600 L "
          "agitated tanks; retrofit the nozzles of the DN1200 AFD")
PROS_SCOPE = ("preparative solution-prep, preparative fraction and precipitation-area equipment replaced "
              "(old for new) and some tank nozzles retrofitted to BPE; piping components changed to centrifugal "
              "pumps, quaternary diaphragm pumps, instruments, valves, pH meters and flowmeters that meet BPE "
              "with 2D/3D drainability")
CONS_MATERIAL = "materials of some AFD, reactor and tank equipment are not replaced to meet BPE"
CONS_CIP = "most connecting piping is not retrofitted for CIP; the resulting audit risk needs to be assessed"

OPTIONS = [
    [["Option", "Scope of Retrofit", "Completion Date", "Pros and Cons after Retrofit"]],
    [
        ["Option 1", "(Piping retrofit with 2D drainability; shortest option)"],
        ["Hastelloy reactor R05: replace with a new BPE-standard unit. Solution-preparation tanks: replace 7 tanks "
         "with 316L tanks. Preparative operation: replace 8 receiving drums with 8 small side-sampling receivers; "
         "retrofit the nozzles of 5 tanks in this area to BPE. UF area: replace the 1000 L feed tank with a "
         "BPE-compliant 1500 L agitated tank; [[piping not designed to establish a CIP system]]. "
         + PRECIP.replace("DN1200 AFD", "DN1200 AFD (agitated filter dryer)")
         + "; other main equipment unchanged. " + UPGRADE + " that meet BPE; "
         "[[tank connections and the water system do not meet 2D/3D drainability]]; for mobile-phase lines, "
         "2D is not mandatory and 3D is considered; 316L material considered."],
        ["12/17/2026"],
        ["Pros: only " + PROS_SCOPE + "; short retrofit time and fewer equipment validation items.",
         "Cons: " + CONS_MATERIAL + "; " + CONS_CIP + "."],
    ],
    [
        ["Option 2", "(Piping retrofit with CIP cleaning valves reserved)"],
        ["Hastelloy reactor R05: replace with a new BPE-standard unit. Solution-preparation tanks: replace 7 tanks "
         "with 316L tanks. Preparative operation: replace 8 receiving drums with 8 small side-sampling receivers; "
         "retrofit the nozzles of 5 tanks in this area to BPE. UF area: replace the 1000 L feed tank with a "
         "BPE-compliant 1500 L agitated tank; [[piping arranged to establish a CIP system]]. "
         + PRECIP + "; other main equipment unchanged. " + UPGRADE.replace(" and flowmeters", ",")
         + " [[flowmeters that meet BPE with 2D/3D drainability; designed for CIP, with valve connections "
         "reserved for CIP cleaning lines]]; 316L material considered."],
        ["01/20/2027"],
        ["Pros: " + PROS_SCOPE + "; valve connections reserved for CIP cleaning lines make it easy to add a CIP "
         "system later; short retrofit time and fewer equipment validation items.",
         "Cons: " + CONS_MATERIAL + "; " + CONS_CIP + ". In addition, the reserved CIP valves have a long "
         "procurement lead time."],
    ],
    [
        ["Option 3", "(Piping retrofit with a new CIP system)"],
        ["Hastelloy reactor R05: replace with a new BPE-standard unit. Solution-preparation tanks: replace 7 tanks "
         "with 316L tanks; retrofit the piping to establish a CIP system. Preparative operation: replace 8 "
         "receiving drums with 8 small side-sampling receivers; retrofit the nozzles of 5 tanks in this area to "
         "BPE; [[piping arranged to establish a CIP system.]] UF area: replace the 1000 L feed tank with a "
         "BPE-compliant 1500 L agitated tank; retrofit the piping to establish a CIP system. "
         + PRECIP + "; other main equipment unchanged; retrofit the piping to establish a CIP system. "
         + UPGRADE + " [[that meet BPE with 2D/3D drainability; CIP piping and valve cleaning system "
         "established;]] 316L material considered."],
        ["04/25/2027"],
        ["Pros: " + PROS_SCOPE + "; CIP piping and valve cleaning system established, ensuring compliance to the "
         "greatest extent.",
         "Cons: " + CONS_MATERIAL + "; long supply lead time for the CIP piping retrofit, many construction items "
         "and more subsequent validation, i.e. a long procurement, construction and validation cycle."],
    ],
]
OPTIONS_COL_W = [1900000, 9800000, 1550000, 7677292]
OPTIONS_HEADER_SZ, OPTIONS_LABEL_SZ, OPTIONS_BODY_SZ = 20, 18, 17
OPTIONS_CELL_MAR = 54864  # 0.06 in; the source's 0.5 pt lets English text touch the cell borders

# ---------------------------------------------------------------- slide 3: equipment table
EQUIP_HEADER = ["Process", "No.", "Equipment Tag", "Size", "Material", "Status", "Vendor", "Equipment Use",
                "Arrival Date", "Drawings Sent to CIP Designer for Clarification", "Remarks"]
SUPPLIER = " – supplier info required"
EQUIP_TERMS = {
    "制备系统储罐汉邦": ["Prep system tanks", "Hanbon"],
    "D级洁净区汉邦": ["Grade D area", "Hanbon"],
    "超滤系统-东富龙": ["UF system", "Tofflon"],
    "转盐、析晶、结晶釜、三合一亚光": ["Salt conversion & crystallization reactors, AFD", "Yaguang"],
    "哈氏合金": ["Hastelloy"],
    "需更换": ["Replace"],
    "改造罐口": ["Retrofit nozzles"],
    "无需改造": ["No retrofit"],
    "诚信": ["Chengxin"],
    "东富龙": ["Tofflon"],
    "英德": ["Yingde"],
    "汉邦": ["Hanbon"],
    "亚光": ["Yaguang"],
    "是": ["Yes"],
    "溶料配液": ["Dissolution / solution prep"],
    "浓盐罐": ["Concentrated salt tank"],
    "配盐罐": ["Salt solution prep tank"],
    "纯水储罐": ["PW storage tank"],
    "乙腈储罐": ["Acetonitrile tank"],
    "流动相A储罐": ["Mobile phase A tank"],
    "流动相B储罐": ["Mobile phase B tank"],
    "废液储罐": ["Waste liquid tank"],
    "异丙醇储罐": ["IPA tank"],
    "组分罐": ["Fraction tank"],
    "上样罐": ["Load tank"],
    "组分罐(缓冲)": ["Fraction tank (buffer)"],
    "组分接收罐(合格)": ["Fraction receiver (in-spec)"],
    "顶样罐": ["Chase tank"],
    "匀浆罐": ["Slurry tank"],
    "超滤接收罐": ["UF receiving tank"],
    "NaOH+MeOH储存溶液": ["NaOH+MeOH storage solution"],
    "转盐釜": ["Salt conversion reactor"],
    "析晶釜": ["Crystallization reactor"],
    "三合一烘料": ["AFD drying"],
    "更换储罐-需要提供供应商": ["Replace tank" + SUPPLIER],
    "改造罐口-需要提供供应商": ["Retrofit nozzles" + SUPPLIER],
    "罐口无改造-需要提供供应商": ["No nozzle retrofit" + SUPPLIER],
    "11/07/2026到货": ["[[11/07/2026 arrival]]"],
    "两个方案：方案1：现场改造-时间来不及；暂定按照方案2：按照中式区-更换诚信新釜-供应商新图纸（AT04）": [
        "[[Two options:]]",
        "[[Option 1: on-site retrofit – not enough time;]]",
        "[[Tentatively Option 2: follow the pilot-plant area – replace with a new Chengxin reactor – "
        "supplier's new drawing (AT04)]]"],
    "确定异丙醇能否使用T06储罐（方案取消）-刘林冲-已经确定将储罐更换为英德1500L 备用储罐-2026.09.27": [
        "Confirm whether IPA can use tank T06 (option cancelled) – Liu Linchong – confirmed: tank to be replaced "
        "with the Yingde 1500 L spare tank – 2026.09.27"],
}
EQUIP_DATE = re.compile(r"^(\d\d/\d\d/\d{4})(到货|改完)$")
EQUIP_DATE_WORD = {"到货": "arrival", "改完": "retrofitted"}
EQUIP_COL_W = [1300000, 560000, 2244436, 900000, 1150000, 1350000, 1000000, 2400000, 1950000, 1700000]
EQUIP_HEADER_SZ, EQUIP_PROCESS_SZ = 14, 13

# ---------------------------------------------------------------- slide 4: CAD picture annotations
# picture name -> [(Chinese ink box in source-image pixels, English box in source-image pixels, lines, pt, anchor)]
# English boxes avoid the leader lines and symbols; each is white-filled and drawn over the union of both boxes.
ANNOT_PT = 10
NOTE_LINES = ["Main valve horizontal, self-draining;", "branch valve mounted vertically upward",
              "on the right; main/branch valve diaphragm", "orientation similar to HVL02"]
BLOCK_LINES = ["Block valve / GMP-type valve, L ≤ 2D"]
HIGH_LINES = ["Piping high point:", "new cleaning valve set added"]
ANNOTATIONS = {
    "图片 3": [
        ((321, 298, 494, 334), (200, 296, 650, 334), BLOCK_LINES, ANNOT_PT, "b"),
        ((279, 360, 682, 436), (200, 358, 682, 470), NOTE_LINES, ANNOT_PT, "t"),
        ((780, 360, 872, 390), (745, 345, 960, 392), ["Highest point:", "connect cleaning water"], 9, "b"),
        ((991, 401, 1234, 434), (991, 378, 1290, 436), HIGH_LINES, ANNOT_PT, "b"),
    ],
    "图片 8": [
        ((314, 314, 488, 346), (220, 310, 660, 347), BLOCK_LINES, ANNOT_PT, "b"),
        ((273, 376, 676, 452), (220, 374, 690, 486), NOTE_LINES, ANNOT_PT, "t"),
        ((985, 417, 1227, 450), (985, 394, 1285, 452), HIGH_LINES, ANNOT_PT, "b"),
    ],
}
# Source label box "纯化水/CCA/溶剂 | 来自清洗站" (arrow tip on the green line) -> redrawn as a pentagon arrow.
SOURCE_LABEL = ("图片 8", (215, 208, 395, 258), ["PW / CCA / solvent", "from wash station"], 8)


def set_fonts(rpr):
    for tag in ("latin", "ea", "cs"):
        old = rpr.find(A + tag)
        if old is not None:
            rpr.remove(old)
    anchor = next((c for c in rpr if c.tag in RPR_AFTER_FONTS), None)
    for tag in ("latin", "ea", "cs"):
        el = rpr.makeelement(A + tag, {"typeface": FONT})
        if anchor is None:
            rpr.append(el)
        else:
            anchor.addprevious(el)


def clean_rpr(rpr, size=None, color=None):
    rpr.set("lang", "en-US")
    for k in ("altLang", "dirty", "err", "smtClean"):
        rpr.attrib.pop(k, None)
    if size is not None:
        rpr.set("sz", str(int(size * 100)))
    if color is not None:
        fill = rpr.find(A + "solidFill")
        if fill is None:
            fill = rpr.makeelement(A + "solidFill", {})
            ln = rpr.find(A + "ln")
            if ln is None:
                rpr.insert(0, fill)
            else:
                ln.addnext(fill)
        for c in list(fill):
            fill.remove(c)
        fill.append(fill.makeelement(A + "srgbClr", {"val": color}))
    set_fonts(rpr)
    return rpr


def segments(text):
    """'a [[b]] c' -> [('a ', False), ('b', True), (' c', False)]"""
    out = []
    for i, part in enumerate(re.split(r"\[\[|\]\]", text)):
        if part:
            out.append((part, i % 2 == 1))
    return out


def fill_paragraph(p, text, base, size=None, black=None):
    for child in list(p):
        if child.tag in (A + "r", A + "br", A + "fld"):
            p.remove(child)
    end = p.find(A + "endParaRPr")
    for seg, red in segments(text):
        r = p.makeelement(A + "r", {})
        color = RED if red else black
        r.append(clean_rpr(copy.deepcopy(base), size, color))
        t = r.makeelement(A + "t", {})
        t.text = seg
        r.append(t)
        if end is None:
            p.append(r)
        else:
            end.addprevious(r)
    if end is not None:
        clean_rpr(end, size)


def set_text(txbody, paras, size=None, black=None):
    """Replace all paragraphs of a txBody with `paras`, reusing the first paragraph's pPr and run format."""
    ps = txbody.findall(A + "p")
    first = ps[0]
    base = copy.deepcopy(next(first.iter(A + "rPr")))
    for p in ps[1:]:
        txbody.remove(p)
    tmpl = copy.deepcopy(first)
    for i, text in enumerate(paras):
        p = first if i == 0 else copy.deepcopy(tmpl)
        if i:
            txbody.append(p)
        fill_paragraph(p, text, base, size, black)


def cell_key(cell):
    return "".join(t.text or "" for t in cell._tc.iter(A + "t")).strip()


def text_width_pt(text, size, bold=True):
    """Width of `text` in Times New Roman (metric-compatible Liberation Serif), or a rough estimate."""
    try:
        from PIL import ImageFont

        style = "Bold" if bold else "Regular"
        font = ImageFont.truetype(f"/usr/share/fonts/truetype/liberation/LiberationSerif-{style}.ttf", 1000)
        return font.getlength(text) * size / 1000
    except OSError:
        return len(text) * size * 0.55


def translate_title(shape, text, keep_right=True):
    tb = shape.text_frame._txBody
    set_text(tb, [text])
    if keep_right:
        right = shape.left + shape.width
        size = int(next(tb.iter(A + "rPr")).get("sz", "4800")) / 100
        shape.width = Emu(int(Pt(text_width_pt(text, size) * 1.08 + 12)))
        shape.left = Emu(right - shape.width)


def translate_options(table):
    for i, w in enumerate(OPTIONS_COL_W):
        table.columns[i].width = Emu(w)
    for r, row in enumerate(OPTIONS):
        for c in range(4):
            cell = table.cell(r, c)
            paras = row[0][c:c + 1] if r == 0 else row[c]
            size = OPTIONS_HEADER_SZ if r == 0 else OPTIONS_LABEL_SZ if c in (0, 2) else OPTIONS_BODY_SZ
            set_text(cell._tc.txBody, paras, size, black="000000")
            cell.margin_left = cell.margin_right = Emu(OPTIONS_CELL_MAR)
            if r and c == 1:
                for p in cell._tc.txBody.findall(A + "p"):
                    ppr = p.find(A + "pPr")
                    if ppr is not None and ppr.get("algn") == "just":
                        ppr.set("algn", "l")


def translate_equipment(table):
    widths = list(EQUIP_COL_W)
    widths.append(sum(c.width for c in table.columns) - sum(widths))
    for i, w in enumerate(widths):
        table.columns[i].width = Emu(w)
    missing = []
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            if cell.is_spanned:
                continue
            key = cell_key(cell)
            if r == 0:
                set_text(cell._tc.txBody, [EQUIP_HEADER[c]], EQUIP_HEADER_SZ)
                continue
            if not re.search(r"[\u3400-\u9fff（）：；]", key):
                set_text(cell._tc.txBody, [key])
                continue
            m = EQUIP_DATE.match(key)
            if m:
                paras = [f"{m.group(1)} {EQUIP_DATE_WORD[m.group(2)]}"]
            elif key in EQUIP_TERMS:
                paras = EQUIP_TERMS[key]
            else:
                missing.append((r, c, key))
                continue
            set_text(cell._tc.txBody, paras, EQUIP_PROCESS_SZ if c == 0 else None)
    assert not missing, missing


def px_to_emu(pic):
    """Map source-image pixels to slide EMU for a cropped, scaled picture."""
    from PIL import Image
    import io

    im = Image.open(io.BytesIO(pic.image.blob))
    w, h = im.size
    l, t, r, b = pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom
    sx = pic.width / ((1 - l - r) * w)
    sy = pic.height / ((1 - t - b) * h)
    return lambda x, y: (int(pic.left + (x - l * w) * sx), int(pic.top + (y - t * h) * sy))


def add_label(slide, x0, y0, x1, y1, lines, size, anchor, shape_type=MSO_SHAPE.RECTANGLE, outline=False):
    shp = slide.shapes.add_shape(shape_type, Emu(x0), Emu(y0), Emu(x1 - x0), Emu(y1 - y0))
    shp.shadow.inherit = False
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if outline:
        shp.line.color.rgb = RGBColor(0, 0, 0)
        shp.line.width = Pt(0.5)
    else:
        shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Emu(int(Pt(2)))
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "b": MSO_ANCHOR.BOTTOM, "m": MSO_ANCHOR.MIDDLE}[anchor]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = 1.0
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        clean_rpr(run._r.get_or_add_rPr())
    return shp


def annotate_pictures(slide):
    pics = {s.name: s for s in slide.shapes if s.shape_type == 13}
    for name, items in ANNOTATIONS.items():
        to_emu = px_to_emu(pics[name])
        for ink, box, lines, size, anchor in items:
            x0, y0 = to_emu(min(ink[0], box[0]), min(ink[1], box[1]))
            x1, y1 = to_emu(max(ink[2], box[2]), max(ink[3], box[3]))
            add_label(slide, x0, y0, x1, y1, lines, size, anchor)
    name, (bx0, by0, bx1, by1), lines, size = SOURCE_LABEL
    to_emu = px_to_emu(pics[name])
    x0, y0 = to_emu(bx0, by0)
    x1, y1 = to_emu(bx1, by1)
    add_label(slide, x0, y0, x1, y1, lines, size, "m", MSO_SHAPE.PENTAGON, outline=True)


def main():
    prs = Presentation(SRC)
    for n, slide in enumerate(prs.slides, 1):
        for shape in list(slide.shapes):
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if n in (1, 5) and shape.name == "Text Box 3" and re.search(r"[\u3400-\u9fff！]", text):
                    translate_title(shape, TITLES[n], keep_right=False)
                elif shape.name == "TextBox 5" and text in OPTION_LABELS:
                    set_text(shape.text_frame._txBody, [OPTION_LABELS[text]])
                elif shape.name == "TextBox 5":
                    translate_title(shape, TITLES[n])
            elif shape.has_table:
                if n == 2:
                    translate_options(shape.table)
                elif n == 3:
                    translate_equipment(shape.table)
        if n == 4:
            annotate_pictures(slide)
    prs.save(OUT)
    print(OUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
