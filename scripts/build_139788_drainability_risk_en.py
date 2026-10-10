#!/usr/bin/env python3
"""English version of 139788项目step3管路管路无法排空风险分析(1).pptx (one slide, two tables).

All review comments are removed. Text is Times New Roman; the option table is re-laid out (column widths,
left-aligned bodies, narrower numbering indent) so that the English fits at presentation size and the
timeline table no longer overlaps it. The third option row is labelled 方案2 in the source but is Option 3
(see the timeline table); the English uses Option 3.

Usage: python3 scripts/build_139788_drainability_risk_en.py
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_139788_retrofit_v2_en import A, set_fonts, text_width_pt  # noqa: E402
from build_uf_retrofit_en import fill_cell, times_new_roman  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "139788项目step3管路管路无法排空风险分析(1).pptx"
OUT = ROOT / "deliverables" / "139788项目step3管路无法排空风险分析_EN.pptx"

TITLE = "139788 Step 3 Piping Dead Legs and Residue: Retrofit Timeline"
TITLE_SZ = 44

OPTIONS_HEADER = ["Option", "Description", "Potential Risk", "Mitigation Plan", "Residual Risk", "Due Date"]
OPTIONS = [
    [
        ["Option 1"],
        ["20 locations do not meet the 2D/3D dead-leg requirement, and 11 diaphragm valves are not "
         "self-draining.",
         "Some lines are not BPE tubing."],
        ["Residue in the lines reduces product yield; residual material may degrade or support microbial "
         "growth.",
         "No defined requirements for rinsing, draining, purging or disassembly after each use of the lines."],
        ["After each material transfer, add a process-solvent rinse so that the residue is a "
         "non-growth-promoting medium.",
         "Update the operating and cleaning procedures to define rinsing, draining, purging and disassembly "
         "requirements."],
        ["Low", "(with enhanced procedural controls)"],
        ["10/15/2026"],
    ],
    [
        ["Option 2"],
        ["Retrofit lines and tanks with zero-dead-leg tee valves and short-outlet (3D) tees to control dead "
         "legs, slope and drainage."],
        ["New change control; impact on product quality must be assessed.",
         "Pre-use and post-use operations not sufficiently defined."],
        ["Qualify the retrofitted lines.",
         "Define the operating and cleaning procedures and update the related forms."],
        ["Low", "(with retrofitted hardware and defined procedures)"],
        ["12/27/2026"],
    ],
    [
        ["Option 3"],
        ["Retrofit lines and tanks with zero-dead-leg tee valves and GMP block valves to control dead legs, "
         "and install a CIP system for cleaning in place."],
        ["Inadequate design.",
         "Inadequate control of the CIP program."],
        ["Qualify the equipment as a whole.",
         "Achieve proper slope and full drainage and eliminate dead legs.",
         "Provide pre-use flushing, cleaning, sanitization and purging, with full CIP capability.",
         "Update SOPs to incorporate these requirements into routine operation."],
        ["Low", "(with improved piping hardware and CIP program)"],
        ["04/25/2027"],
    ],
]
# x positions in pt on the 1920 x 1080 pt slide
OPTIONS_TOP = 150
OPTIONS_COL_W = [150, 410, 420, 500, 230, 150]
OPTIONS_HEADER_SZ, OPTIONS_BODY_SZ, OPTIONS_NOTE_SZ = 26, 24, 20
BULLET_MAR = 30  # pt, hanging indent for numbered items

TIMELINE = [
    ["Process Step", "Service", "Demo", "Pre-PPQ", "PPQ", "Commercial"],
    ["Preparative Solution Prep", "Process materials", "Option 1", "Option 2", "Option 2", "Option 3"],
    ["Preparative Purification", "Process materials", "Option 1", "Option 2", "Option 2", "Option 3"],
    ["TFF", "Process materials", "Option 1", "Option 2", "Option 2", "Option 3"],
    ["Precipitation", "Process materials", "Option 1", "Option 2", "Option 2", "Option 3"],
]
TIMELINE_COL_W = [420, 300, 190, 190, 190, 220]
TIMELINE_SZ = 24
TIMELINE_ROW_H = 42
TIMELINE_BOTTOM = 1012  # pt, just above the footer band

COMMENTS_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
AUTHORS_RT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors"


def pt(v):
    return Emu(int(Pt(v)))


def drop_comments(prs):
    for slide in prs.slides:
        for rid, rel in list(slide.part.rels.items()):
            if rel.reltype == COMMENTS_RT:
                slide.part.drop_rel(rid)
    for rid, rel in list(prs.part.rels.items()):
        if rel.reltype == AUTHORS_RT:
            prs.part.drop_rel(rid)


def set_size(tc_or_txbody, size):
    for rpr in tc_or_txbody.iter(A + "rPr", A + "endParaRPr"):
        rpr.set("sz", str(int(size * 100)))
        rpr.attrib.pop("altLang", None)
        rpr.set("lang", "en-US")


def style_paragraphs(txbody, align):
    for p in txbody.findall(A + "p"):
        ppr = p.find(A + "pPr")
        if ppr is None:
            ppr = p.makeelement(A + "pPr", {})
            p.insert(0, ppr)
        ppr.set("algn", align)
        if ppr.find(A + "buAutoNum") is not None:
            ppr.set("marL", str(int(Pt(BULLET_MAR))))
            ppr.set("indent", str(-int(Pt(BULLET_MAR))))


def translate_options(shape):
    shape.left, shape.top = pt((1920 - sum(OPTIONS_COL_W)) / 2), pt(OPTIONS_TOP)
    table = shape.table
    for i, w in enumerate(OPTIONS_COL_W):
        table.columns[i].width = pt(w)
    rows = [[[h] for h in OPTIONS_HEADER]] + OPTIONS
    for r, row in enumerate(rows):
        table.rows[r].height = pt(40)
        for c, paras in enumerate(row):
            cell = table.cell(r, c)
            body = cell._tc.txBody
            fill_cell(body, paras)
            set_size(body, OPTIONS_HEADER_SZ if r == 0 else OPTIONS_BODY_SZ)
            if r and c == 4:
                for rpr in body.findall(f"{A}p")[1].iter(A + "rPr", A + "endParaRPr"):
                    rpr.set("sz", str(OPTIONS_NOTE_SZ * 100))
            style_paragraphs(body, "l" if r and c in (1, 2, 3) else "ctr")
            cell.margin_left = cell.margin_right = pt(8)
    times_new_roman(table._tbl)


def translate_timeline(shape):
    table = shape.table
    for i, w in enumerate(TIMELINE_COL_W):
        table.columns[i].width = pt(w)
    width = sum(TIMELINE_COL_W)
    shape.left = pt((1920 - width) / 2)
    shape.top = pt(TIMELINE_BOTTOM - TIMELINE_ROW_H * len(TIMELINE))
    for r, row in enumerate(TIMELINE):
        table.rows[r].height = pt(TIMELINE_ROW_H)
        for c, text in enumerate(row):
            body = table.cell(r, c)._tc.txBody
            fill_cell(body, [text])
            set_size(body, TIMELINE_SZ)
            style_paragraphs(body, "ctr")
    times_new_roman(table._tbl)


def translate_title(shape):
    body = shape.text_frame._txBody
    fill_cell(body, [TITLE])
    set_size(body, TITLE_SZ)
    for rpr in body.iter(A + "rPr", A + "endParaRPr"):
        set_fonts(rpr)
    right = shape.left + shape.width
    shape.width = pt(text_width_pt(TITLE, TITLE_SZ) + 30)
    shape.left = Emu(right - shape.width)


def main():
    prs = Presentation(SRC)
    drop_comments(prs)
    slide = prs.slides[0]
    tables = [s for s in slide.shapes if s.has_table]
    for shape in slide.shapes:
        if shape.name == "TextBox 5":
            translate_title(shape)
    translate_options(tables[0])
    translate_timeline(tables[1])
    prs.save(OUT)
    print(OUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
