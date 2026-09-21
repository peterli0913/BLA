#!/usr/bin/env python3
"""Revised BLA retrofit meeting minutes as Word — follow original outline, correct owners."""

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from docx.shared import Cm, Pt, RGBColor

OUT = "/workspace/meetings/2026/2026-09-21-BLA改造沟通会会议纪要-修订.docx"

NAVY = RGBColor(0x1B, 0x4F, 0x72)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
INK = RGBColor(0x1C, 0x28, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5D, 0x6D, 0x7E)
ROW_ALT = "F4F8FB"
HEAD_BG = "1B4F72"
WARN_BG = "FDEBD0"


def set_run_font(run, name="Arial", size=10.5, bold=False, color=INK):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), "微软雅黑")


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color="D5D8DC"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def cell_text(cell, text, *, bold=False, size=9.5, color=INK, center=False, fill=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        shade(cell, fill)
    set_cell_border(cell)
    # clear default para
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_heading_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "1B4F72")
    pBdr.append(bottom)
    pPr.append(pBdr)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=True, color=NAVY)


def add_body(doc, text, *, bold=False, size=10.5, space_after=6, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep_with_next
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=INK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.75)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)


def set_col_widths(table, widths_cm):
    table.autofit = False
    table.allow_autofit = False
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)
    sec.top_margin = Cm(1.3)
    sec.bottom_margin = Cm(1.3)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("内部资料 · POA / 项目 139788 · 2026-09-21")
    set_run_font(hr, size=8, color=MUTED)

    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("修订依据：总指挥张俊峰会前分工。人员以该消息为准，不以转写口语或初稿职称为准。")
    set_run_font(fr, size=8, color=MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    tr = title.add_run("BLA改造沟通会会议纪要")
    set_run_font(tr, size=18, bold=True, color=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(8)
    sr = sub.add_run("修订稿  ·  校正人员、压缩表述")
    set_run_font(sr, size=11, bold=True, color=GOLD)

    # meta
    add_heading_line(doc, "一、会议时间")
    add_body(doc, "2026年9月21日（周一）09:03–09:29")

    add_heading_line(doc, "二、参会人员")
    add_body(
        doc,
        "张俊峰、范兴涛、刘美佳、皮轶杰、郭泽鹏、张冬、韩瑞衡、王志恒、张国睿、"
        "田子才、冯毅、宋金柱、刘维江；记录：李涛、李丰江。",
        space_after=4,
    )
    add_body(doc, "刘维江会中明确后天到岗。初稿另列人员本次未写入职责。", size=9.5, space_after=2)

    add_heading_line(doc, "三、记录人员")
    add_body(doc, "李涛、李丰江")

    add_heading_line(doc, "四、职责（按总指挥会前分工）")
    add_body(doc, "今日现场：工序负责人跟王志恒和厂家技术走车间，既是执行，也是学习。QA 等 DH 人员到场后再定主责。", size=10)

    roles = [
        ("角色", "人员", "职责"),
        ("总指挥", "张俊峰", "总体安排；配液利旧/推倒两套方案列出后报领导"),
        ("IEPE", "王志恒、张国睿", "设计整体负责"),
        ("改造清单 / 厂家对表", "王志恒", "打印清单。上午东富龙：制备配液+超滤，逐管路设备查漏补缺；13:00 亚光：沉淀，更新同一清单"),
        ("客户问题举一反三", "范兴涛", "客户来访全部问题，核三个片段是否都有整改要求"),
        ("制备配液 / 制备工序", "皮轶杰", "本片段事项；今日跟现场勘察"),
        ("超滤工序", "张冬", "本片段事项；今日跟现场勘察"),
        ("沉淀工序", "韩瑞衡", "本片段事项；13:00 跟亚光对表"),
        ("现场配合", "郭泽鹏", "50% 精力在整改；今日跟王志恒和厂家走现场，不统筹三个片段"),
        ("工程", "宋金柱", "工程整体负责；今日交出 139788 采购总台账三类"),
        ("采购日报", "冯毅、田子才", "即日起群里每日上报进展"),
        ("QA", "待 DH 支持", "今日不指定 QA 主责"),
    ]
    t = doc.add_table(rows=len(roles), cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [4.2, 4.0, 9.0]
    set_col_widths(t, widths)
    for i, row in enumerate(roles):
        for j, val in enumerate(row):
            if i == 0:
                cell_text(t.cell(i, j), val, bold=True, size=9, color=WHITE, center=True, fill=HEAD_BG)
            else:
                fill = ROW_ALT if i % 2 == 0 else "FFFFFF"
                cell_text(
                    t.cell(i, j),
                    val,
                    bold=(j < 2),
                    size=9,
                    center=(j < 2),
                    fill=fill,
                )

    add_heading_line(doc, "五、会议纪要")
    add_body(doc, "按方案一推进，保证 10 月 15 日开建、11 月 18 日交付。今日把厂家勘察补进同一份改造清单，并出 139788 采购总台账。", space_after=6)

    add_body(doc, "1. 现场勘察与清单", bold=True, space_after=2, keep_with_next=True)
    add_bullet(doc, "昨日东富龙、亚光已做初勘。今日不另起表，只在王志恒打印的改造事项清单上查漏补缺。")
    add_bullet(doc, "上午：王志恒带东富龙技术员，制备配液 + 超滤，逐管路、逐设备过。皮轶杰跟制备，张冬跟超滤，郭泽鹏随队。")
    add_bullet(doc, "13:00：王志恒带亚光技术员过沉淀。韩瑞衡跟本片段。")
    add_bullet(doc, "东富龙须出利旧、推倒重建两套方案做时间和事项对比；配液最终走哪套，列出后由张俊峰报领导，会上未定。")
    add_bullet(doc, "范兴涛：客户来访问题在三个片段是否都举一反三、都有整改要求，结果写回清单。")

    add_body(doc, "2. 工程与采购", bold=True, space_after=2, keep_with_next=True)
    add_bullet(doc, "工程总责：宋金柱。今日必须交出 139788 项目采购总台账，三类写清：已下单采买、询价中、尚未采购。")
    add_bullet(doc, "冯毅、田子才：即日起每天在群里报进展。客户新加项和厂家新识别项，清单出来后再往里补。")

    add_body(doc, "3. 时间节点（方案一）", bold=True, space_after=2, keep_with_next=True)
    add_bullet(doc, "Demo 期 10 月 5 日；开建 10 月 15 日（demo 提前则可能提前 1–2 天拆设备，推迟则开工顺延）。")
    add_bullet(doc, "建设约 11 月 3 日，验证约 11 月 16 日，清洗 2 天，交付 11 月 18 日。一律不得晚于 11 月 18 日。")
    add_bullet(doc, "制备配液最紧。超滤约 11 月 5 日基本完成，不是关键路径。")
    add_bullet(doc, "设计须在 10 月中下旬前完成并报客户。边改边设计来不及。IEPE：王志恒、张国睿。")

    add_body(doc, "4. 风险与沟通", bold=True, space_after=2, keep_with_next=True)
    add_bullet(doc, "刘美佳：腾信 D 区现有搅拌罐改口约 20 天，与 11/18 冲突。继续压工期，资源不够找郭宏杰，完不成报到中午会。")
    add_bullet(doc, "每晚 19:00 改造沟通会，可线上。周一早晨约 30 分钟同步。")
    add_bullet(doc, "湿氮气中午前擦净、消毒、转移至 AFD（三合一）后侧，避开 13:00 客户进车间。郭泽鹏运到门口，宋金柱安排拆卸转移。")

    add_heading_line(doc, "六、待办事项")

    todos = [
        ("事项", "责任人", "时限"),
        ("打印改造事项清单；与东富龙逐管路、逐设备过制备配液和超滤，缺项补进清单", "王志恒（皮轶杰、张冬、郭泽鹏现场跟）", "9/21 上午"),
        ("与亚光过沉淀区域整改项，更新同一清单", "王志恒（韩瑞衡现场跟）", "9/21 13:00"),
        ("客户来访问题在三个片段是否全部举一反三、均有整改要求", "范兴涛", "9/21"),
        ("东富龙出利旧 / 推倒两套方案对比", "王志恒对接厂家", "9/21"),
        ("交出 139788 采购总台账：已下单 / 询价中 / 未采购", "宋金柱", "9/21"),
        ("群里上报采购与到货进展", "冯毅、田子才", "每日"),
        ("湿氮气中午前转移到 AFD（三合一）后侧", "郭泽鹏、宋金柱", "9/21 12:00"),
        ("腾信搅拌罐改口：压工期或正式报出冲突节点", "刘美佳（升级：郭宏杰）", "9/21"),
        ("配液利旧还是推倒，方案列出后报领导", "张俊峰", "方案列出后"),
        ("19:00 改造沟通会", "全员（可线上）", "每日"),
    ]
    t2 = doc.add_table(rows=len(todos), cols=3)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False
    set_col_widths(t2, [9.0, 5.2, 3.0])
    for i, row in enumerate(todos):
        for j, val in enumerate(row):
            if i == 0:
                cell_text(t2.cell(i, j), val, bold=True, size=9, color=WHITE, center=True, fill=HEAD_BG)
            else:
                fill = ROW_ALT if i % 2 == 0 else "FFFFFF"
                if "腾信" in row[0]:
                    fill = WARN_BG
                cell_text(t2.cell(i, j), val, bold=(j > 0), size=9, center=(j > 0), fill=fill)


    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
