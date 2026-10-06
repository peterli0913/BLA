#!/usr/bin/env python3
"""Append Step 3 of 10.6第一页统计(1).xlsx, in English, to the styled 10-06_EN.xlsx.

The styled workbook (repo root, saved from WPS with the user's font and layout) is the baseline:
Times New Roman 24 pt, the existing column widths, row height, fills and the photo sheet are kept
byte for byte. Only the Overall Plan sheet grows, and the column D header changes from "Line" to
"Component / Fitting" because the new source renamed 管路 to 部件/管件.

Usage: python3 scripts/append_1006_step3_en.py
"""
import re
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from PIL import ImageFont

ROOT = Path(__file__).resolve().parents[1]
SRC_EN = ROOT / "10-06_EN.xlsx"
SRC_NEW = ROOT / "10.6第一页统计(1).xlsx"
OUT = ROOT / "deliverables" / "10-06_EN.xlsx"

# cell styles already in the styled workbook (xl/styles.xml)
STYLE_BODY = 5       # TNR 24, centered, full grid border
STYLE_STEP = 9       # TNR 24 bold, blue fill, full grid border (same as Step 1)
STYLE_DATE = 10      # TNR 24, number format "mmm d"
STYLE_MERGE_TOP = 11
STYLE_MERGE_MID = 12
STYLE_MERGE_BOT = 14
STYLE_DONE = 13      # TNR 24, green fill
ROW_H = 37
HEADER_ROW = 1

TERMS = {
    "step3": "Step 3",
    "Head": "Head",
    "进行中": "In progress",
    "已完成": "Completed",
    "N/A": "N/A",
    "哈氏合金搅拌罐，300L": "Hastelloy agitated tank, 300 L",
    "哈氏合金搅拌罐，500L": "Hastelloy agitated tank, 500 L",
    "316L，1000L": "316L, 1000 L",
    "316L，8000L": "316L, 8000 L",
    "316L，3000L": "316L, 3000 L",
    "316L，80L": "316L, 80 L",
    "316L，600L": "316L, 600 L",
    "316L": "316L",
    "哈氏合金反应釜，3000L": "Hastelloy reactor, 3000 L",
    "哈氏合金三合一，DN1200": "Hastelloy agitated filter dryer (AFD), DN1200",
    "DAC1000/316L": "DAC1000 / 316L",
    "加纯化水管路": "Purified water addition line",
    "加纯化水管线": "Purified water addition line",
    "配制流动相A": "Mobile phase A preparation",
    "配制流动相B": "Mobile phase B preparation",
    "自循环": "Self-recirculation",
    "连接到制备系统": "Connection to the preparative system",
    "溶剂管路": "Solvent line",
    "上样流动相A": "Mobile phase A loading",
    "配制顶样液": "Chase solution preparation",
    "上流动相B": "Mobile phase B loading",
    "加甲醇": "Methanol addition",
    "转移甲醇": "Methanol transfer",
    "自循环清洗": "Self-recirculation cleaning",
    "转移顶样液": "Chase solution transfer",
    "转移顶样液（制备系统)": "Chase solution transfer (preparative system)",
    "转移稀释后粗肽": "Diluted crude peptide transfer",
    "上样": "Sample loading",
    "转移粗肽": "Crude peptide transfer",
    "转移合格组分": "In-spec fraction transfer",
    "转移组分": "Fraction transfer",
    "接收组分": "Fraction receiving",
    "接收合格组分": "In-spec fraction receiving",
    "转移不合格组分": "Out-of-spec fraction transfer",
    "转移制备废液": "Preparative waste transfer",
    "转移盐溶液": "Salt solution transfer",
    "待超滤体系转移": "Pre-UF solution transfer",
    "转移物料": "Material transfer",
    "平衡膜包": "Membrane cassette equilibration",
    "转移溶剂": "Solvent transfer",
    "超滤体系转移": "UF solution transfer",
    "转移废液": "Waste transfer",
    "析晶管路": "Crystallization line",
    "转移制备组分": "Preparative fraction transfer",
    "预留": "Reserved",
    "接收溶剂": "Solvent receiving",
}
HEADER = {"部件/管件": "Component / Fitting"}


def tr(value):
    if not isinstance(value, str):
        return value
    key = value.strip()
    if key in TERMS:
        return TERMS[key]
    raise KeyError(key)


def read_step3():
    ws = load_workbook(SRC_NEW)["Sheet1"]
    rows = []
    for r in range(2, ws.max_row + 1):
        rows.append([ws.cell(r, c).value for c in range(1, 7)])
    merges = []
    for m in ws.merged_cells.ranges:
        merges.append((m.min_row, m.min_col, m.max_row, m.max_col))
    return rows, merges


def col_px(width):
    return int(((256 * width + int(128 / 7)) / 256) * 7)


def wrap_lines(text, size, avail_px):
    font = ImageFont.truetype("/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf", size)
    n = 1
    line = ""
    for word in text.split(" "):
        trial = (line + " " + word).strip()
        if line and font.getlength(trial) > avail_px:
            n += 1
            line = word
        else:
            line = trial
    return n


def main():
    rows, merges = read_step3()
    assert rows[0][0] == "step3" and len(rows) == 91, len(rows)
    delta = 29  # source row 2 -> destination row 31

    with zipfile.ZipFile(SRC_EN) as z:
        parts = {n: z.read(n) for n in z.namelist()}

    sst = parts["xl/sharedStrings.xml"].decode("utf-8")
    existing = re.findall(r"<si>(.*?)</si>", sst, re.S)
    strings = [re.sub(r"&amp;", "&", "".join(re.findall(r"<t[^>]*>(.*?)</t>", it, re.S))) for it in existing]
    index = {s: i for i, s in enumerate(strings)}

    def sid(text):
        if text not in index:
            index[text] = len(strings)
            strings.append(text)
        return index[text]

    # the styled header says "Line"; the new source widened the column to 部件/管件
    assert strings[3] == "Line"
    strings[3] = HEADER["部件/管件"]
    index[HEADER["部件/管件"]] = 3

    groups = {}  # (col) -> list of (start, end) in destination rows, 1-based
    for r1, c1, r2, c2 in merges:
        if r1 == 1:
            continue
        groups.setdefault(c1, []).append((r1 + delta, r2 + delta))
    step_span = next(span for span in groups[1])

    avail_d = col_px(52.9166666666667) - 14
    row_xml = []
    for i, (step, equip, material, part, status, when) in enumerate(rows):
        dest = i + 31
        part_en = tr(part)
        status_en = tr(status)
        lines = wrap_lines(part_en, 24, avail_d)
        height = ROW_H if lines == 1 else 8 + lines * 30

        def cell(col, style, value=None, numeric=False):
            ref = f"{col}{dest}"
            if numeric:
                return f'<c r="{ref}" s="{style}"><v>{value}</v></c>'
            if value is None:
                return f'<c r="{ref}" s="{style}"/>'
            return f'<c r="{ref}" s="{style}" t="s"><v>{sid(value)}</v></c>'

        in_b = next((span for span in groups.get(2, []) if span[0] <= dest <= span[1]), None)
        in_c = next((span for span in groups.get(3, []) if span[0] <= dest <= span[1]), None)

        def merge_style(span, value):
            if span is None:
                return STYLE_BODY, value
            start, end = span
            pos = 0 if dest == start else (2 if dest == end else 1)
            style = (STYLE_MERGE_TOP, STYLE_MERGE_MID, STYLE_MERGE_BOT)[pos]
            return style, value if dest == start else None

        b_style, b_val = merge_style(in_b, equip)
        c_style, c_val = merge_style(in_c, None if material is None else tr(material))
        e_style = STYLE_DONE if status_en == "Completed" else STYLE_BODY
        if isinstance(when, (int, float)):
            f = cell("F", STYLE_DATE, int(when), numeric=True)
        else:
            f = cell("F", STYLE_BODY, tr(when))
        cols = [
            cell("A", STYLE_STEP, "Step 3" if dest == step_span[0] else None),
            cell("B", b_style, b_val),
            cell("C", c_style, c_val),
            cell("D", STYLE_BODY, part_en),
            cell("E", e_style, status_en),
            f,
        ]
        row_xml.append(f'<row r="{dest}" ht="{height}" customHeight="1" spans="1:6">{"".join(cols)}</row>')

    sheet = parts["xl/worksheets/sheet1.xml"].decode("utf-8")
    anchor = re.search(r'<row r="30" .*?</row>', sheet, re.S)
    assert anchor, "row 30 not found"
    sheet = sheet[:anchor.end()] + "".join(row_xml) + sheet[anchor.end():]
    last = 30 + len(rows)
    sheet = sheet.replace('ref="A1:F30"', f'ref="A1:F{last}"', 1)
    extra = [f'<mergeCell ref="A{step_span[0]}:A{step_span[1]}"/>']
    for col in (2, 3):
        letter = "BC"[col - 2]
        for start, end in groups[col]:
            extra.append(f'<mergeCell ref="{letter}{start}:{letter}{end}"/>')
    sheet = sheet.replace('<mergeCells count="14">', f'<mergeCells count="{14 + len(extra)}">', 1)
    sheet = sheet.replace("</mergeCells>", "".join(extra) + "</mergeCells>", 1)
    sheet = sheet.replace(
        '<rowBreaks count="1" manualBreakCount="1"><brk id="18" max="16383" man="1"/></rowBreaks>',
        '<rowBreaks count="2" manualBreakCount="2"><brk id="18" max="16383" man="1"/>'
        f'<brk id="30" max="16383" man="1"/></rowBreaks>', 1)
    parts["xl/worksheets/sheet1.xml"] = sheet.encode("utf-8")

    rebuilt = []
    for i, s in enumerate(strings):
        if i == 3 or i >= len(existing):
            rebuilt.append(f"<si><t>{escape(s)}</t></si>")
        else:
            rebuilt.append(f"<si>{existing[i]}</si>")
    sst = re.sub(r'count="\d+" uniqueCount="\d+"', f'count="{len(strings)}" uniqueCount="{len(strings)}"', sst, count=1)
    sst = re.sub(r"<si>.*</si>", "".join(rebuilt), sst, count=1, flags=re.S)
    parts["xl/sharedStrings.xml"] = sst.encode("utf-8")

    with zipfile.ZipFile(OUT, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in parts.items():
            out.writestr(name, data)
    print(OUT.relative_to(ROOT), "rows", last, "strings", len(strings))


if __name__ == "__main__":
    main()
