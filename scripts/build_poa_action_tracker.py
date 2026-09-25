#!/usr/bin/env python3
"""Build the POA retrofit action tracker in the single-sheet register format."""

from datetime import date
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

OUT = "/workspace/meetings/BLA改造行动项跟踪.xlsx"

NAVY = "1B4F72"
WHITE = "FFFFFF"
INK = "1C2833"
LINE = "D5D8DC"
ALT = "F4F8FB"
OPEN_BG = "FBF8F1"
PROG_BG = "D6EAF8"
BLOCK_BG = "FADBD8"
DONE_BG = "D5F5E3"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
thick_bottom = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="medium", color=NAVY),
)

font = Font(name="Arial", size=10, color=INK)
font_b = Font(name="Arial", size=10, bold=True, color=INK)
font_title = Font(name="Arial", size=18, bold=True, color=WHITE)
font_hdr = Font(name="Arial", size=9, bold=True, color=WHITE)
wrap = Alignment(wrap_text=True, vertical="center", horizontal="left")
wrap_c = Alignment(wrap_text=True, vertical="center", horizontal="center")

fill_navy = PatternFill("solid", fgColor=NAVY)
fill_white = PatternFill("solid", fgColor=WHITE)
fill_head = PatternFill("solid", fgColor=NAVY)

STATUSES = ["未开始", "进行中", "阻塞", "已完成", "取消"]
PRIORS = ["高", "中", "低"]
STREAMS = [
    "制备配液",
    "超滤",
    "沉淀",
    "采购到货",
    "设计方案",
    "QA质量",
    "客户沟通",
    "综合协调",
]

# Curated from the uploaded register. Do not attach Excel comments/notes.
ACTIONS = [
    dict(id="A-20260921-01", raised=date(2026, 9, 21), stream="设计方案", origin="张俊峰",
         action="与东富龙技术员逐管路、逐设备过制备配液和超滤，缺项补进同一清单。",
         done=None, owner="王志恒", due=date(2026, 9, 21), status="未开始", pri="高",
         progress="2026-09-21：皮轶杰跟制备，张冬跟超滤，郭泽鹏随队。",
         src="2026-09-21 部署会"),
    dict(id="A-20260921-05", raised=date(2026, 9, 21), stream="客户沟通", origin="张俊峰",
         action="梳理客户来访全部问题，核三个片段涉及整改的内容是否都举一反三、都有整改要求。",
         done=None, owner="范兴涛", due=date(2026, 9, 21), status="未开始", pri="高",
         progress="2026-09-21：",
         src="2026-09-21 部署会"),
    dict(id="A-20260921-06", raised=date(2026, 9, 21), stream="制备配液", origin="张俊峰",
         action="对接东富龙，下午交出利旧与推倒重建两套方案的时间/事项对比。",
         done=None, owner="王志恒", due=date(2026, 9, 21), status="未开始", pri="高",
         progress="2026-09-21：会上要求厂家出两套方案。",
         src="2026-09-21 部署会"),
    dict(id="A-20260921-07", raised=date(2026, 9, 21), stream="沉淀", origin="张俊峰",
         action="13:00 与亚光技术员梳理沉淀区域整改项，更新改造清单。",
         done=None, owner="王志恒", due=date(2026, 9, 21), status="未开始", pri="高",
         progress="2026-09-21：韩瑞衡现场跟。",
         src="2026-09-21 部署会"),
    dict(id="A-20260921-11", raised=date(2026, 9, 21), stream="采购到货", origin="张俊峰",
         action="项目采购总台账，三类写清：已下单采买、询价中、尚未采购。",
         done=None, owner="宋金柱", due=date(2026, 9, 21), status="未开始", pri="高",
         progress="2026-09-21：宋金柱跟采购同事今天出总账。",
         src="2026-09-21 部署会"),
    dict(id="A-20260921-12", raised=date(2026, 9, 21), stream="采购到货", origin="张俊峰",
         action="冯毅、田子才每日在群里上报采购与到货进展。",
         done=None, owner="冯毅", due=date(2026, 9, 21), status="进行中", pri="中",
         progress="2026-09-21：冯毅、田子才后续每日上报",
         src="2026-09-21 部署会"),
    dict(id="A-20260921-13", raised=date(2026, 9, 21), stream="综合协调", origin="张俊峰",
         action="湿氮气擦净、消毒、中午前转移至 AFD（三合一）后侧。",
         done=None, owner="郭泽鹏", due=date(2026, 9, 21), status="未开始", pri="高",
         progress="2026-09-21：郭泽鹏运到门口；宋金柱安排拆卸转移。",
         src="2026-09-21 部署会"),
    dict(id="A-20260921-15", raised=date(2026, 9, 21), stream="制备配液", origin="宋金柱",
         action="向腾信压缩 D 区现有搅拌罐改口工期；达不成则把完不成节点报到中午沟通会。",
         done=None, owner="宋金柱", due=date(2026, 9, 21), status="未开始", pri="高",
         progress="2026-09-21：约 5 台、工期 20 天，与 11/18 冲突。资源不够找郭宏杰。",
         src="2026-09-21 部署会"),
]


def style_range(cell, fill=None, font_=None, align=None, border=None, num=None):
    if fill:
        cell.fill = fill
    cell.font = font_ or font
    cell.alignment = align or wrap
    cell.border = border if border is not None else thin
    if num:
        cell.number_format = num


def apply_cf_status(ws, col, first, last):
    rng = f"{col}{first}:{col}{last}"
    mapping = {
        "未开始": "D5D8DC",
        "进行中": "AED6F1",
        "阻塞": "F5B7B1",
        "已完成": "ABEBC6",
        "取消": "E5E8E8",
    }
    for val, color in mapping.items():
        ws.conditional_formatting.add(
            rng,
            CellIsRule(
                operator="equal",
                formula=[f'"{val}"'],
                fill=PatternFill("solid", fgColor=color),
                font=Font(name="Arial", size=10, bold=True, color=INK),
            ),
        )


def apply_cf_pri(ws, col, first, last):
    rng = f"{col}{first}:{col}{last}"
    for val, color in (("高", "F5B041"), ("中", "F9E79F"), ("低", "D5DBDB")):
        ws.conditional_formatting.add(
            rng,
            CellIsRule(
                operator="equal",
                formula=[f'"{val}"'],
                fill=PatternFill("solid", fgColor=color),
                font=Font(name="Arial", size=10, bold=True, color=INK),
            ),
        )


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "行动项总账"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(0.4, 0.4, 0.5, 0.5)
    ws.print_title_rows = "1:3"
    ws.sheet_properties.tabColor = NAVY

    ws.merge_cells("A1:M2")
    ws["A1"] = "POA 生物制品项目改造 · 行动项总账（持续）"
    style_range(ws["A1"], fill=fill_navy, font_=font_title,
                align=Alignment(vertical="center", horizontal="left"), border=Border())
    for r in (1, 2):
        for c in range(1, 14):
            ws.cell(r, c).fill = fill_navy
            ws.cell(r, c).border = Border()

    headers = [
        "ID",
        "提出日\nRaised",
        "提出方\nOriginator",
        "片断 / 专业\nWorkstream",
        "行动项\nAction",
        "验收标准\nDone means",
        "负责人\nOwner",
        "截止日期\nDue",
        "状态\nStatus",
        "优先级\nPriority",
        "进展日志（只追加）\nProgress log",
        "更新日\nUpdated",
        "来源会议\nSource",
    ]
    for i, h in enumerate(headers, 1):
        style_range(ws.cell(3, i, h), fill=fill_head, font_=font_hdr, align=wrap_c, border=thick_bottom)

    first_data = 4
    fills_status = {
        "未开始": OPEN_BG,
        "进行中": PROG_BG,
        "阻塞": BLOCK_BG,
        "已完成": DONE_BG,
    }
    for i, a in enumerate(ACTIONS):
        r = first_data + i
        values = [
            a["id"], a["raised"], a["origin"], a["stream"], a["action"], a["done"],
            a["owner"], a["due"], a["status"], a["pri"], a["progress"],
            date(2026, 9, 21), a["src"],
        ]
        row_fill = PatternFill("solid", fgColor=fills_status.get(a["status"], ALT if i % 2 else WHITE))
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            num = "YYYY-MM-DD" if c in (2, 8, 12) else None
            al = wrap_c if c in (1, 2, 4, 7, 8, 9, 10, 12) else wrap
            fnt = font_b if c in (1, 7, 9) else font
            style_range(cell, fill=row_fill, font_=fnt, align=al, num=num)
            # never attach cell comments / notes
            cell.comment = None
        ws.row_dimensions[r].height = 56

    last = first_data + len(ACTIONS) - 1
    blank_last = last + 15
    for r in range(last + 1, blank_last + 1):
        for c in range(1, 14):
            cell = ws.cell(r, c, None)
            style_range(cell, fill=fill_white, font_=font, align=wrap)
            if c in (2, 8, 12):
                cell.number_format = "YYYY-MM-DD"
            cell.comment = None
        ws.row_dimensions[r].height = 22

    dv_status = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv_status.error = "请选：未开始 / 进行中 / 阻塞 / 已完成 / 取消"
    dv_status.errorTitle = "状态"
    dv_status.showInputMessage = False
    dv_pri = DataValidation(type="list", formula1='"' + ",".join(PRIORS) + '"', allow_blank=True)
    dv_pri.showInputMessage = False
    dv_stream = DataValidation(type="list", formula1='"' + ",".join(STREAMS) + '"', allow_blank=True)
    dv_stream.showInputMessage = False
    ws.add_data_validation(dv_status)
    ws.add_data_validation(dv_pri)
    ws.add_data_validation(dv_stream)
    dv_status.add(f"I{first_data}:I{blank_last}")
    dv_pri.add(f"J{first_data}:J{blank_last}")
    dv_stream.add(f"D{first_data}:D{blank_last}")

    apply_cf_status(ws, "I", first_data, blank_last)
    apply_cf_pri(ws, "J", first_data, blank_last)

    ws.auto_filter.ref = f"A3:M{blank_last}"
    ws.freeze_panes = "A4"
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 32
    widths = {
        "A": 16, "B": 12, "C": 16, "D": 14, "E": 42, "F": 28, "G": 20,
        "H": 13, "I": 12, "J": 10, "K": 46, "L": 12, "M": 18,
    }
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

    wb.save(OUT)
    print("saved", OUT, "rows", len(ACTIONS))


if __name__ == "__main__":
    build()
