#!/usr/bin/env python3
"""Client English deck: 139788 Step 3 dead-leg and hold-up retrofit schedule."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTS = [
    ROOT / "deliverables" / "pptx" / "2026-10-27-client-139788-step3-deadleg-holdup.pptx",
    ROOT / "139788 Step 3 Dead-leg and Hold-up Retrofit_EN.pptx",
]

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x1A, 0x23, 0x32)
TEAL = RGBColor(0x2D, 0x8B, 0x8B)
SEA = RGBColor(0xA8, 0xDA, 0xDC)
CREAM = RGBColor(0xF1, 0xFA, 0xEE)
INK = RGBColor(0x1C, 0x28, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5D, 0x6D, 0x7E)
CARD = RGBColor(0xF7, 0xFB, 0xFB)
LINE = RGBColor(0xD5, 0xE3, 0xE3)
LOW = RGBColor(0x1F, 0x6B, 0x4A)
FONT = "DejaVu Sans"


def set_run(run, text, size, bold=False, color=INK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = rpr.makeelement(qn(f"a:{tag}"), {})
            rpr.append(el)
        el.set("typeface", font)


def add_text(slide, text, x, y, w, h, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    body_anchor = {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor]
    tf._txBody.bodyPr.set("anchor", body_anchor)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    set_run(p.add_run(), text, size, bold, color)
    return box


def add_lines(slide, lines, x, y, w, h, size=15, color=INK, spacing=8, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(0 if i == 0 else spacing)
        p.space_after = Pt(0)
        set_run(p.add_run(), line, size, bold, color)
    return box


def rect(slide, x, y, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def footer(slide, page, total):
    add_text(slide, "139788 Step 3  ·  Dead-leg and hold-up retrofit", 0.55, 7.12, 8.5, 0.24, 11, color=MUTED)
    add_text(slide, f"{page}  /  {total}", 11.2, 7.12, 1.55, 0.24, 11, color=MUTED, align=PP_ALIGN.RIGHT)


def header_bar(slide):
    rect(slide, 0, 0, 13.333, 0.08, TEAL)


def title(slide, text):
    add_text(slide, text, 0.55, 0.28, 12.2, 0.95, 26, True, NAVY)


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, WHITE)
    header_bar(s)
    return s


def style_cell(cell, text, size, bold, color, fill, align=PP_ALIGN.LEFT):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.margin_left = Inches(0.1)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.06)
    cell.margin_bottom = Inches(0.06)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, bold, color)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "139788 Step 3 Dead-leg and Hold-up Retrofit"
    prs.core_properties.author = "Asymchem"
    prs.core_properties.subject = "Client briefing"
    total = 5

    # Cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, NAVY)
    rect(s, 0, 0, 0.18, 7.5, TEAL)
    add_text(s, "139788  ·  STEP 3", 0.7, 1.55, 10, 0.35, 16, True, SEA)
    add_text(s, "Dead-leg and hold-up retrofit", 0.7, 2.05, 11.5, 1.1, 40, True, WHITE)
    add_text(s, "Schedule for the three options, and where each one applies.", 0.7, 3.3, 10, 0.5, 18, False, CREAM)
    add_text(s, "Asymchem   ·   Client review   ·   October 2026", 0.7, 6.35, 10, 0.3, 14, False, SEA)

    # Matrix
    s = new_slide(prs)
    title(s, "Demo uses Option 1. Pre-PPQ and PPQ use Option 2.")
    add_text(s, "Commercial manufacture uses Option 3. All four steps handle process material.", 0.55, 1.25, 12, 0.35, 15, False, MUTED)
    rows = [
        ["Process step", "Demo", "Pre-PPQ", "PPQ", "Commercial"],
        ["Preparative buffer preparation", "Option 1", "Option 2", "Option 2", "Option 3"],
        ["Preparative chromatography", "Option 1", "Option 2", "Option 2", "Option 3"],
        ["TFF", "Option 1", "Option 2", "Option 2", "Option 3"],
        ["Precipitation", "Option 1", "Option 2", "Option 2", "Option 3"],
    ]
    table = s.shapes.add_table(5, 5, Inches(0.55), Inches(1.85), Inches(12.2), Inches(4.55)).table
    widths = [3.7, 2.1, 2.15, 2.05, 2.2]
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            if r == 0:
                fill, color, bold = NAVY, WHITE, True
            elif c == 0:
                fill, color, bold = CARD, NAVY, True
            elif val == "Option 1":
                fill, color, bold = RGBColor(0xE7, 0xF3, 0xF3), TEAL, True
            elif val == "Option 3":
                fill, color, bold = RGBColor(0xE8, 0xF2, 0xEA), LOW, True
            else:
                fill, color, bold = WHITE, INK, True
            style_cell(table.cell(r, c), val, 16 if r else 14, bold, color, fill, PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    footer(s, 2, total)

    # Option 1
    s = new_slide(prs)
    title(s, "Option 1: procedural control by November 5, 2026.")
    add_text(s, "The piping stays as it is. The control is how the line is used and cleaned.", 0.55, 1.18, 12, 0.28, 15, False, MUTED)
    rect(s, 0.55, 1.55, 12.25, 1.22, RGBColor(0xF8, 0xEE, 0xEE), RGBColor(0xE4, 0xC8, 0xC8))
    add_text(s, "POTENTIAL RISK", 0.75, 1.64, 4, 0.22, 12, True, RGBColor(0x8C, 0x3A, 0x3A))
    add_text(
        s,
        "Hold-up in the piping can degrade the material, support microbial growth, or cause cross-contamination. "
        "The flush, drain, purge, and disassembly required after each use are not defined.",
        0.75, 1.9, 11.85, 0.72, 15, False, INK,
    )
    rect(s, 0.55, 2.95, 6.0, 2.2, CARD, LINE)
    rect(s, 6.75, 2.95, 6.05, 2.2, CARD, LINE)
    add_text(s, "WHAT REMAINS", 0.78, 3.08, 5.5, 0.24, 12, True, TEAL)
    add_lines(s, [
        "Some piping does not meet the 2D/3D dead-leg criterion.",
        "Some diaphragm valves do not meet the self-draining requirement.",
        "Some piping is not ASME BPE piping.",
    ], 0.78, 3.38, 5.5, 1.6, 14, INK, 6)
    add_text(s, "WHAT WE WILL ADD", 6.98, 3.08, 5.5, 0.24, 12, True, TEAL)
    add_lines(s, [
        "After each system transfer, flush or clean with process solvent. Bioburden data support this step.",
        "Write the flush, drain, purge, and disassembly requirements into the procedures.",
    ], 6.98, 3.38, 5.55, 1.6, 14, INK, 6)
    rect(s, 0.55, 5.35, 8.35, 1.45, WHITE, LINE)
    rect(s, 9.05, 5.35, 3.75, 1.45, NAVY)
    add_text(s, "RESIDUAL RISK", 0.78, 5.5, 3, 0.25, 12, True, TEAL)
    add_text(s, "Low, once the tighter procedural controls are in place.", 0.78, 5.85, 7.8, 0.7, 16, False, INK)
    add_text(s, "DUE", 9.28, 5.5, 3.3, 0.25, 12, True, SEA)
    add_text(s, "November 5, 2026", 9.28, 5.9, 3.3, 0.55, 22, True, WHITE)
    footer(s, 3, total)

    # Option 2
    s = new_slide(prs)
    title(s, "Option 2: piping retrofit complete by February 6, 2027.")
    add_text(s, "This is the configuration for Pre-PPQ and PPQ.", 0.55, 1.28, 12, 0.32, 15, False, MUTED)
    rect(s, 0.55, 1.8, 6.0, 3.35, CARD, LINE)
    rect(s, 6.75, 1.8, 6.05, 3.35, CARD, LINE)
    add_text(s, "THE RETROFIT", 0.78, 1.98, 5.5, 0.28, 12, True, TEAL)
    add_lines(s, [
        "Rebuild the piping and vessels with hygienic T-valves and short-outlet tees meeting the 3D criterion.",
        "Control dead legs, slope, and drainability.",
        "Leave valves for a later CIP connection.",
    ], 0.78, 2.4, 5.5, 2.45, 15, INK, 8)
    add_text(s, "RISK AND THE RESPONSE", 6.98, 1.98, 5.5, 0.28, 12, True, TEAL)
    add_lines(s, [
        "This is a new change control. The concern is an effect on product quality, and pre-use and post-use steps that are not yet sufficient.",
        "Qualify the modified piping. Define the operating and cleaning sequence, and update the records.",
    ], 6.98, 2.4, 5.55, 2.45, 15, INK, 10)
    rect(s, 0.55, 5.35, 8.35, 1.45, WHITE, LINE)
    rect(s, 9.05, 5.35, 3.75, 1.45, NAVY)
    add_text(s, "RESIDUAL RISK", 0.78, 5.5, 3, 0.25, 12, True, TEAL)
    add_text(s, "Low, with the modified hardware and the revised sequence.", 0.78, 5.85, 7.8, 0.7, 16, False, INK)
    add_text(s, "DUE", 9.28, 5.5, 3.3, 0.25, 12, True, SEA)
    add_text(s, "February 6, 2027", 9.28, 5.9, 3.3, 0.55, 22, True, WHITE)
    footer(s, 4, total)

    # Option 3
    s = new_slide(prs)
    title(s, "Option 3 adds CIP, five months after Option 2.")
    add_text(s, "This is the configuration for commercial manufacture. It is built on the Option 2 piping.", 0.55, 1.28, 12, 0.32, 15, False, MUTED)
    rect(s, 0.55, 1.8, 6.0, 3.35, CARD, LINE)
    rect(s, 6.75, 1.8, 6.05, 3.35, CARD, LINE)
    add_text(s, "WHAT CIP ADDS", 0.78, 1.98, 5.5, 0.28, 12, True, TEAL)
    add_lines(s, [
        "Clean-in-place on the retrofitted piping.",
        "Slope, drainage, and removal of dead legs.",
        "Pre-use flush, cleaning, disinfection, and purge, all done in place.",
    ], 0.78, 2.4, 5.5, 2.45, 15, INK, 10)
    add_text(s, "RISK AND THE RESPONSE", 6.98, 1.98, 5.5, 0.28, 12, True, TEAL)
    add_lines(s, [
        "The design may be incomplete, and the CIP sequence and controls may be incomplete.",
        "Qualify the system as a whole. Update the SOP and bring the requirements into routine use.",
    ], 6.98, 2.4, 5.55, 2.45, 15, INK, 10)
    rect(s, 0.55, 5.35, 8.35, 1.45, WHITE, LINE)
    rect(s, 9.05, 5.35, 3.75, 1.45, NAVY)
    add_text(s, "RESIDUAL RISK", 0.78, 5.5, 3, 0.25, 12, True, TEAL)
    add_text(s, "Low, with the revised piping and the CIP cycle.", 0.78, 5.85, 7.8, 0.7, 16, False, INK)
    add_text(s, "DUE", 9.28, 5.48, 3.3, 0.22, 12, True, SEA)
    add_lines(s, ["5 months after", "Option 2"], 9.28, 5.78, 3.3, 0.85, 18, WHITE, 0, True)
    footer(s, 5, total)

    one = one_page()
    outs = OUTS + [
        ROOT / "deliverables" / "pptx" / "2026-10-27-client-139788-step3-deadleg-holdup-1page.pptx",
        ROOT / "139788 Step 3 Dead-leg and Hold-up Retrofit_EN_1page.pptx",
    ]
    # first two paths are the five-slide deck
    for path in OUTS:
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(path)
        print(path)
    for path in outs[2:]:
        path.parent.mkdir(parents=True, exist_ok=True)
        one.save(path)
        print(path)


def one_page():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "139788 Step 3 Dead-leg and Hold-up Retrofit — one page"
    prs.core_properties.author = "Asymchem"
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, WHITE)
    header_bar(s)
    add_text(s, "139788 Step 3 — Dead-leg and hold-up retrofit", 0.38, 0.16, 9.2, 0.36, 20, True, NAVY)
    add_text(s, "Demo: Option 1    Pre-PPQ and PPQ: Option 2    Commercial: Option 3", 0.38, 0.52, 9.4, 0.26, 12, False, MUTED)
    add_text(s, "All four steps handle process material.", 9.6, 0.52, 3.4, 0.26, 12, False, MUTED, PP_ALIGN.RIGHT)

    options = [
        ["Option", "What it covers", "Potential risk", "Mitigation", "After mitigation", "Due"],
        [
            "1",
            "Dead legs outside the 2D/3D criterion, diaphragm valves that are not self-draining, and non-BPE piping stay in place.",
            "Hold-up can degrade material, support microbial growth, or cause cross-contamination. Flush, drain, purge, and disassembly after each use are not defined.",
            "After each system transfer, flush or clean with process solvent, supported by bioburden data. Define those steps in the procedures.",
            "Low",
            "Nov 5, 2026",
        ],
        [
            "2",
            "Rebuild piping and vessels with hygienic T-valves and 3D short-outlet tees. Control dead legs, slope, and drainability. Reserve CIP tie-in valves.",
            "New change control, with a concern for product quality. Pre-use and post-use steps may still be insufficient.",
            "Qualify the modified piping. Define the operating and cleaning sequence and update the records.",
            "Low",
            "Feb 6, 2027",
        ],
        [
            "3",
            "Add clean-in-place on the Option 2 piping, including slope, drainage, dead-leg removal, and pre-use flush, cleaning, disinfection, and purge.",
            "The design may be incomplete, and the CIP sequence and controls may be incomplete.",
            "Qualify the system as a whole. Update the SOP and bring the requirements into routine use.",
            "Low",
            "5 months after Option 2",
        ],
    ]
    table = s.shapes.add_table(4, 6, Inches(0.32), Inches(0.88), Inches(12.7), Inches(4.55)).table
    for i, w in enumerate([0.85, 2.85, 3.15, 3.15, 1.15, 1.55]):
        table.columns[i].width = Inches(w)
    table.rows[0].height = Inches(0.36)
    for r in range(1, 4):
        table.rows[r].height = Inches(1.38)
    for r, row in enumerate(options):
        for c, val in enumerate(row):
            if r == 0:
                fill, color, bold, size = NAVY, WHITE, True, 11
            elif c == 0:
                fill, color, bold, size = TEAL, WHITE, True, 16
            elif c == 4:
                fill, color, bold, size = RGBColor(0xE8, 0xF2, 0xEA), LOW, True, 12
            else:
                fill, color, bold, size = WHITE if r % 2 else CARD, INK, False, 11
            style_cell(table.cell(r, c), val, size, bold, color, fill, PP_ALIGN.CENTER if c in (0, 4, 5) else PP_ALIGN.LEFT)

    phase = [
        ["Process step", "Demo", "Pre-PPQ", "PPQ", "Commercial"],
        ["Preparative buffer preparation", "Option 1", "Option 2", "Option 2", "Option 3"],
        ["Preparative chromatography", "Option 1", "Option 2", "Option 2", "Option 3"],
        ["TFF", "Option 1", "Option 2", "Option 2", "Option 3"],
        ["Precipitation", "Option 1", "Option 2", "Option 2", "Option 3"],
    ]
    t2 = s.shapes.add_table(5, 5, Inches(0.32), Inches(5.48), Inches(12.7), Inches(1.62)).table
    for i, w in enumerate([3.9, 2.2, 2.2, 2.2, 2.2]):
        t2.columns[i].width = Inches(w)
    for r, row in enumerate(phase):
        t2.rows[r].height = Inches(0.32)
        for c, val in enumerate(row):
            if r == 0:
                fill, color, bold = NAVY, WHITE, True
            elif c == 0:
                fill, color, bold = CARD, NAVY, True
            elif val == "Option 1":
                fill, color, bold = RGBColor(0xE7, 0xF3, 0xF3), TEAL, True
            elif val == "Option 3":
                fill, color, bold = RGBColor(0xE8, 0xF2, 0xEA), LOW, True
            else:
                fill, color, bold = WHITE, INK, True
            style_cell(t2.cell(r, c), val, 11, bold, color, fill, PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)
    add_text(s, "139788 Step 3  ·  Dead-leg and hold-up retrofit", 0.32, 7.2, 8, 0.22, 10, color=MUTED)
    return prs


if __name__ == "__main__":
    main()
