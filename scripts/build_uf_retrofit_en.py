#!/usr/bin/env python3
"""English versions of 超滤设备改造.pptx and 超滤设备改造-压.xlsx (139788 step3 TFF gap assessment).

PPTX: all text and tables become English in Times New Roman; pictures are untouched. On slides 4-5 the
embedded P&ID (OLE) is left as is; only its red numbered retrofit notes are covered with a white-filled
English text box.
XLSX: only xl/sharedStrings.xml and the font names in xl/styles.xml are rewritten, so every picture,
merge and column width stays byte-identical. The 09.02 meeting sheet is already English + Chinese; the
Chinese half of each question is dropped.

Usage: python3 scripts/build_uf_retrofit_en.py
"""
import copy
import re
import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_139788_retrofit_v2_en import (  # noqa: E402
    A, FONT, add_label, cell_key, clean_rpr, fill_paragraph, set_fonts, set_text, text_width_pt,
    translate_title,
)

ROOT = Path(__file__).resolve().parents[1]
PPT_SRC = ROOT / "超滤设备改造.pptx"
PPT_OUT = ROOT / "deliverables" / "超滤设备改造_EN.pptx"
XLS_SRC = ROOT / "超滤设备改造-压.xlsx"
XLS_OUT = ROOT / "deliverables" / "超滤设备改造_EN.xlsx"

CJK = re.compile(r"[\u3000-\u303f\u3400-\u9fff\uff00-\uffef]")

# ================================================================ PPTX
TITLES = {
    1: "Gap Identification – TFF",
    2: "Category 3 Piping and Equipment Requirements",
    3: "Category 4 Piping and Equipment Requirements",
    4: "P&ID before Retrofit – Tank Section",
    5: "P&ID after Retrofit – Tank Section",
    6: "Category 4 Tank Requirements",
}

FINISH = "Internal surface finish: NMT 25 µin Ra or 0.63 µm"
TRI_CLAMP = "Tri-clamp connections (solid gasket seal)"
CIP_VALVES = "Diaphragm valves or other clean-in-place valves; ball valves prohibited"
CIP_3D = "Piping can be CIP-cleaned; 3D design to minimize residual dead legs"
SLOPE = "Slope > 1%, fully drainable"
FINISH_OK = "Surface finish per manufacturer standard, compliant;"

# cell text (paragraphs joined by "\n") -> English paragraphs
CELLS = {
    # ---- slides 2-3: Category 3 / Category 4 requirement tables
    "标准项目": ["Standard Item"],
    "描述 / 指导要点": ["Description / Guidance"],
    "Lines, Connections, Components\n管路、连接、部件": ["Lines, Connections, Components"],
    "・三夹快装接头（实体垫片密封）": ["• " + TRI_CLAMP],
    "・ASME BPE 规格管材": ["• ASME BPE tubing"],
    "・球阀：存在风险残留物的系统禁止使用球阀，球体后方积料会成为隐患；球阀可用于溶剂系统隔离工况": [
        "• Ball valves: not permitted in systems at risk of residue, as material trapped behind the ball is a "
        "hazard; ball valves may be used for isolation in solvent systems"],
    "・材质 MOC：不锈钢、高合金": ["• MOC: stainless steel, high alloy"],
    "・执行 ASME BPE 标准；若采用其他卫生级标准，需要评估确认其可接受性": [
        "• Follow ASME BPE; if another sanitary standard is used, its acceptability must be assessed"],
    "Tanks 罐体": ["Tanks"],
    "・材质：不锈钢或高合金": ["• Material: stainless steel or high alloy"],
    "・内表面粗糙度：NMT 25Ra（微英寸），即 0.63μm": ["• Internal surface finish: NMT 25 µin Ra, i.e. 0.63 µm"],
    "・三夹快装连接": ["• Tri-clamp connections"],
    "・尽量缩短罐封头接管伸出长度": ["• Minimize the projection length of nozzles on tank heads"],
    "Cleaning 清洗": ["Cleaning"],
    "・该标准适用于缓冲液系统：介质为溶剂、不滋生微生物、无物料堆积，不需要开展常规清洗的工况": [
        "• Applies to buffer systems: solvent media, not growth-promoting, no product build-up, "
        "no routine cleaning required"],
    "・尽量减小死管段；若存在死管段，必须具备清洗、换产、排空通路": [
        "• Minimize dead legs; where a dead leg exists, it must have a path for cleaning, changeover and draining"],
    "○ 评估死管段执行 BPE 的2L/D 指标，需要能够流通冲洗": [
        "  ○ Assess dead legs against the BPE 2 L/D criterion; they must be flushable"],
    "○ 系统正常运行时始终满管状态，允许存在极小死管段": [
        "  ○ Very small dead legs are acceptable where the system is always liquid-full in normal operation"],
    "・罐体、大型非管路部件需要做喷淋覆盖测试，便于冲洗 / 换产；不需要开展清洗验证": [
        "• Tanks and large non-piping components require a spray coverage test to support flushing / "
        "changeover; cleaning validation not required"],
    "Surface Finish 表面粗糙度": ["Surface Finish"],
    "・内表面粗糙度：NMT 25Ra（微英寸）或 0.63 微米": ["• " + FINISH],
    "・焊缝内部检查，焊缝内壁光滑，无物料滞留死角": [
        "• Internal weld inspection: smooth internal weld surface, no dead spots for product hold-up"],
    "Line Slope 管路坡度": ["Line Slope"],
    "・不强制管路坡度，但系统必须考虑可排空性能": [
        "• Line slope not mandatory, but the system must be designed to be drainable"],
    "・根据需要设置低点排放口，保障设备有效冲洗与换产": [
        "• Provide low-point drains as needed to ensure effective flushing and changeover"],
    "Documentation 文件资料": ["Documentation"],
    "・ASME 标准全套焊接文件": ["• Full set of ASME weld documentation"],
    "・接触产品或直接接触物料（溶剂、缓冲液等）的弹性体，需提供合规文件：FDA 21CFR 或 USP Class VI 材质证明": [
        "• Elastomers in contact with product or process materials (solvents, buffers, etc.) require compliance "
        "documents: FDA 21 CFR or USP Class VI material certificates"],
    "• 隔膜阀或其他可在位清洗阀门，禁止使用球阀": ["• " + CIP_VALVES],
    "・采用最小死区管件（零静态阀、集成式无菌 GMP 阀、仪表 T 型接头等）": [
        "• Minimum dead-leg fittings (zero-static valves, integrated aseptic GMP valves, instrument tees, etc.)"],
    "・材质 MOC：不锈钢、合金": ["• MOC: stainless steel, alloy"],
    "・遵循 ASME BPE 标准；采用其他卫生标准时，必须评估其可接受性": [
        "• Follow ASME BPE; where another sanitary standard is used, its acceptability must be assessed"],
    "・材质：不锈钢或合金": ["• Material: stainless steel or alloy"],
    "・内表面粗糙度：NMT 25Ra（微英寸）/0.63μm": ["• Internal surface finish: NMT 25 µin Ra / 0.63 µm"],
    "・三夹接头、NovAseptic 或同等等级无菌连接": ["• Tri-clamp, NovAseptic or equivalent aseptic connections"],
    "・尽量缩短罐封头接管长度": ["• Minimize nozzle length on tank heads"],
    "・产品接触罐体、管路必须做清洗验证": ["• Product-contact tanks and piping require cleaning validation"],
    "○ 需要提高清洗频次，加强清洗检查力度": ["  ○ Higher cleaning frequency and more rigorous cleaning inspection"],
    "○ 根据指导文件，需要设定清洁 / 污染存放时限，开展微生物负载取样；系统必须可完全排空冲洗，不能留存微生物滋生基质": [
        "  ○ Per guidance, establish clean / dirty hold times and perform bioburden sampling; the system must be "
        "fully drainable and flushable, leaving no substrate for microbial growth"],
    "・需要清洗验证的罐体、大型非管路部件，必须做喷淋覆盖率测试；例：AFD 粉尘过滤器、卸料斜槽、产品手套箱，保证清洗验证过程可复现": [
        "• Tanks and large non-piping components requiring cleaning validation must pass a spray coverage test "
        "(e.g. AFD dust filter, discharge chute, product glovebox) so that cleaning validation is reproducible"],
    "・焊缝内部检查，内壁光滑，无物料滞留死角": [
        "• Internal weld inspection: smooth internal surface, no dead spots for product hold-up"],
    "• 强制设置管路坡度，保证管路系统完全排空": ["• Line slope mandatory to ensure the piping system is fully drainable"],
    "・高级别焊接文件": ["• Enhanced weld documentation"],
    "• 管路坡度图": ["• Line slope drawings"],
    "• 焊缝分布图 (Weld maps)": ["• Weld maps"],
    "・弹性体材质证明（满足 FDA 21CFR 合规）": ["• Elastomer material certificates (FDA 21 CFR compliant)"],
    # ---- slides 6-8: comparison tables
    "项目": ["Item"],
    "材质": ["Material"],
    "阀门": ["Valves"],
    "连接": ["Connection"],
    "连接方式": ["Connection"],
    "管路标准": ["Piping Standard"],
    "清洁": ["Cleaning"],
    "坡度": ["Slope"],
    "抛光": ["Polishing"],
    "照明": ["Lighting"],
    "三通死角": ["Tee Dead Legs"],
    "仪表": ["Instruments"],
    "流量计": ["Flowmeter"],
    "低点排净": ["Low-Point Drain"],
    "备注": ["Remarks"],
    "C3标准": ["C3 Standard"],
    "C4标准": ["C4 Standard"],
    "316L不锈钢或合金": ["316L stainless steel or alloy"],
    "隔膜阀，或其他可在位清洗阀门": ["Diaphragm valve or other clean-in-place valve"],
    "三夹快装接头（实体垫片密封）": [TRI_CLAMP],
    "可拆卸清洗，可排空，合规": ["Removable for cleaning, drainable, compliant"],
    "坡度大于1%，实现完全排空": [SLOPE],
    "内表面粗糙度：NMT 25Ra（微英寸）或 0.63 微米": ["Internal finish: NMT 25 µin Ra or 0.63 µm"],
    "符合3D，最小死角设计": ["3D compliant, minimum dead-leg design"],
    "安装满足3D要求，便于清洁，无死角": ["Installed to meet 3D, easy to clean, no dead legs"],
    "Tank\n储罐": ["Tank"],
    "液位计隔膜阀角度不符合": ["Level gauge diaphragm valve angle non-compliant"],
    "符合要求，均为卡盘连接": ["Compliant, all clamp connections"],
    "不符合要求，液位计位置的隔膜阀无法在线清洁": [
        "Non-compliant: diaphragm valve at level gauge cannot be cleaned in line"],
    "坡度符合要求": ["Slope compliant"],
    "符合要求": ["Compliant"],
    "不符合，无试镜灯": ["Non-compliant, no sight-glass light"],
    "不符合，不满足3D要求": ["Non-compliant, does not meet 3D"],
    "不符合项\n需要将储罐整体更换为满足BPE需求储罐;\n物料单独进;\n储罐增加在线CIP清洗，满足清洁需求。": [
        "Non-compliance:", "Replace the entire tank with a BPE-compliant tank;", "Separate inlet for each material;",
        "Add in-line CIP to the tank to meet cleaning needs."],
    "TFF\n超滤主体": ["TFF", "TFF Skid"],
    "隔膜阀": ["Diaphragm valves"],
    "不锈钢、合金": ["Stainless steel, alloy"],
    "隔膜阀或其他可在位清洗阀门，禁止使用球阀": [CIP_VALVES],
    "BPE规格管材": ["BPE tubing"],
    "管路可进行CIP清洗，使用３D减少残留死区": [CIP_3D],
    "管路可进行CIP清洗，使用３D减少残留该死区": [CIP_3D],
    "设置低点排放": ["Low-point drains provided"],
    "1. 表面抛光度为出厂标准，符合要求；\n2.管路不涉及流量计问题；": [
        "1. " + FINISH_OK, "2. No flowmeter issue on these lines;"],
    "Solvent Drum\n→\nTJ4S-1211-TFF02\n乙腈/水管路": ["Solvent Drum ", "→ ", " TJ4S-1211-TFF02", "ACN / Water Line"],
    "符合，使用卡盘连接": ["Compliant, clamp joints"],
    "不符合要求，管路需更换为BPE管": ["Non-compliant; replace with BPE tubing"],
    "不符合": ["Non-compliant"],
    "改造方案": ["Retrofit Plan"],
    "有低点排净": ["Has low-point drain"],
    "使用符合3D的管路减少残留死角": ["Use 3D-compliant piping to minimize residual dead legs", ""],
    "CNC-Solvent   →\nTJ4S-1211-TFF02-T01\n甲醇管路": ["CNC-Solvent   →    ", "TJ4S-1211-TFF02-T01", "Methanol Line"],
    "TJ4S-1211-R02\n→\nTJ4S-1211-TFF02\n氢氧化钠管路": ["TJ4S-1211-R02", "→ ", "TJ4S-1211-TFF02", "NaOH Line"],
    "有低点排净，符合要求": ["Has low-point drain, compliant"],
    "更换过滤器，满足排净需求": ["Replace filter to meet drainage requirement", ""],
    "不锈钢、合金材质，BPE": ["Stainless steel or alloy, BPE"],
    "无不可排净的流量计": ["No non-drainable flowmeters"],
    "1.管路材质为BPE，符合要求\n2.表面抛光度为出厂标准，符合要求\n\n\n详情见附件": [
        "1. Piping material is BPE, compliant;", "2. " + FINISH_OK, "", "", "See attachment for details"],
    "TJ4S-1211-HPPC03-AT03\n→\nTJ4S-1211-TFF02-T01\n制备组分加入储罐": [
        "TJ4S-1211-HPPC03-AT03   ", "→ ", " TJ4S-1211-TFF02-T01", "Prep Fraction Feed to Tank"],
    "符合要求，BPE": ["Compliant, BPE"],
    "符合，使用隔膜阀": ["Compliant, diaphragm valves"],
    "符合，流量计竖直安装": ["Compliant, flowmeter installed vertically"],
    "符合": ["Compliant"],
    "改造管路满足3D标准，使用CIP进行在线清洁": ["Retrofitted piping meets 3D; cleaned in place by CIP", ""],
    "TJ4S-1211-TFF02-T01   →\nTJ4S-1211-TFF02\n甲醇和制备组分加入超滤": [
        "TJ4S-1211-TFF02-T01   →    ", "TJ4S-1211-TFF02", "Methanol and Prep Fraction Feed to TFF"],
    "不符合，管路无低点排净，过滤器无低点排净": ["Non-compliant: no low-point drain on piping or filter"],
    "增加低点排净口；\n过滤器增加进液排净口": ["Add a low-point drain;", "Add an inlet drain to the filter", ""],
}
# (slide, row, col) -> font size (pt) where the English needs a smaller size to keep the source row heights
CELL_SIZE = {
    (6, 2, 10): 8,
    (8, 4, 9): 8,
}

# ---- slides 4-5: red numbered notes inside the embedded P&ID
NOTES = [
    "1. Replace the tank with a BPE-compliant tank;",
    "2. Replace the DP level transmitter with a radar level transmitter;",
    "3. Feed each material through a separate inlet;",
    "4. Remove the pressure gauge root valve;",
    "5. Add in-line CIP cleaning to the tank;",
    "6. Change the dip-tube thermometer to bottom-outlet measurement;",
    "7. Modify the tank bottom outlet valve.",
]
# slide -> (x0, y0, x1, y1) in pt covering the Chinese notes, font pt, line pitch pt (= source pitch, so the
# CAD leader lines still end at the same note)
NOTE_BOX = {
    4: ((726, 95.5, 956, 224), 8, 17.2),
    5: ((681, 81, 956, 188), 8.5, 14.6),
}
NOTE_RED = RGBColor(0xFF, 0x00, 0x00)


def para_text(p):
    return "".join(t.text or "" for t in p.iter(A + "t"))


def cell_text(cell):
    ps = [para_text(p).strip() for p in cell._tc.txBody.findall(A + "p")]
    while ps and not ps[-1]:
        ps.pop()
    return "\n".join(ps)


def fill_cell(txbody, paras, size=None):
    """Write `paras` into a txBody, each paragraph keeping its own pPr and first-run format (e.g. red)."""
    ps = txbody.findall(A + "p")
    fallback = next(txbody.iter(A + "rPr"))
    for i, text in enumerate(paras):
        if i < len(ps):
            p = ps[i]
        else:
            p = copy.deepcopy(ps[-1])
            ps[-1].addnext(p)
            ps.append(p)
        base = next(p.iter(A + "rPr"), None)
        if base is None:
            base = p.find(A + "endParaRPr")
        base = copy.deepcopy(base if base is not None else fallback)
        base.tag = A + "rPr"
        if text:
            fill_paragraph(p, text, base, size)
        else:
            for child in list(p):
                if child.tag in (A + "r", A + "br", A + "fld"):
                    p.remove(child)
    for p in ps[len(paras):]:
        txbody.remove(p)


def times_new_roman(el):
    for rpr in el.iter(A + "rPr", A + "endParaRPr"):
        set_fonts(rpr)


def translate_table(n, table, missing):
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            if cell.is_spanned:
                continue
            key = cell_text(cell)
            if CJK.search(key):
                if key in CELLS:
                    fill_cell(cell._tc.txBody, CELLS[key], CELL_SIZE.get((n, r, c)))
                else:
                    missing.append((n, r, c, key))
    times_new_roman(table._tbl)


def translate_slide_title(shape, text):
    """Right-aligned top-right title; source runs without sz use the 18 pt text-box default."""
    rpr = next(shape.text_frame._txBody.iter(A + "rPr"))
    if rpr.get("sz") is None:
        rpr.set("sz", "1800")
    translate_title(shape, text)


def add_notes(slide, n):
    (x0, y0, x1, y1), size, pitch = NOTE_BOX[n]
    for note in NOTES:
        assert text_width_pt(note, size, bold=False) < x1 - x0 - 4, note
    shp = add_label(slide, Emu(int(Pt(x0))), Emu(int(Pt(y0))), Emu(int(Pt(x1))), Emu(int(Pt(y1))),
                    NOTES, size, "t")
    shp.name = "EN retrofit notes"
    for p in shp.text_frame.paragraphs:
        p.line_spacing = Pt(pitch)
        p.space_before = p.space_after = Pt(0)
        for run in p.runs:
            run.font.color.rgb = NOTE_RED


def build_pptx():
    prs = Presentation(PPT_SRC)
    missing = []
    for n, slide in enumerate(prs.slides, 1):
        for shape in list(slide.shapes):
            if shape.has_text_frame and CJK.search(shape.text_frame.text):
                if n == 1:
                    set_text(shape.text_frame._txBody, [TITLES[n]])
                else:
                    translate_slide_title(shape, TITLES[n])
            elif shape.has_table:
                translate_table(n, shape.table, missing)
        if n in NOTE_BOX:
            add_notes(slide, n)
    assert not missing, "\n".join(map(str, missing))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert not CJK.search(shape.text_frame.text), shape.text_frame.text
    prs.save(PPT_OUT)
    print(PPT_OUT.relative_to(ROOT))


# ================================================================ XLSX
M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
TBP = "To be provided after retrofit"
CIP_FILTER = "Line filter cleaned offline; filter replaced by a jumper for CIP"
DRAIN_FILTER = "Line filter body: add a drain point;"
N2_PURGE = "Solvent line purged with N2, no in-line CIP;"
DL_3D = "3D design to minimize residual dead legs;"
PUMP = "Pneumatic pump: drains before/after pump, N2 purge;"
GRAVITY = "1. Add a line low point for gravity drainage;\n2. " + DRAIN_FILTER

STRINGS = {
    "项目": "Project",
    "步骤": "Step",
    "设备类型": "Equipment Type",
    "材质": "Material",
    "用途": "Purpose",
    "容量/ L": "Capacity / L",
    "设备编号": "Equipment No.",
    "储罐排查项": "Tank Check",
    "阀门排查项": "Valve Check",
    "清洁排查项": "Cleaning Check",
    "罐口管路连接方式排查项": "Tank Nozzle Piping Connection Check",
    "抛光排查项": "Polishing Check",
    "照明排查项": "Lighting Check",
    "搅拌桨排查项": "Agitator Check",
    "三通死角排查项": "Tee Dead-Leg Check",
    "仪表排查项": "Instrument Check",
    "等级": "Category",
    "储罐": "Tank",
    "是否符合，并描述": "Compliant? (Description)",
    "改造前PID图纸": "P&ID before Retrofit",
    "不符合照片": "Non-compliance Photo",
    "改造后PID图纸": "P&ID after Retrofit",
    "改造照片": "Retrofit Photo",
    "阀门（隔膜阀是否有角度）": "Valves (diaphragm valve installation angle)",
    "清洁": "Cleaning",
    "罐口管路连接方式": "Tank Nozzle Piping Connection",
    "抛光": "Polishing",
    "照明": "Lighting",
    "湿氮气": "Humidified Nitrogen",
    "搅拌桨": "Agitator",
    "三通处造成的死角": "Dead Legs at Tees",
    "位置": "Location",
    "仪表是否符合要求": "Instrument Compliance",
    "文件资料": "Documentation",
    "风险评估": "Risk Assessment",
    "超滤": "TFF",
    "暂存罐": "Hold Tank",
    "不符合\n需要将储罐更换为满足BPE需求储罐;\n物料单独进;\n储罐增加在线CIP清洗;":
        "Non-compliant\nReplace the tank with a BPE-compliant tank;\nSeparate inlet for each material;\n"
        "Add in-line CIP to the tank;",
    "待改造后提供": TBP,
    "隔膜阀": "Diaphragm valve",
    "否，液位计隔膜阀角度不符合，易存料；\n":
        "No, level gauge diaphragm valve angle non-compliant, prone to product hold-up;\n",
    "可拆卸清洗，可排空，合规": "Removable for cleaning, drainable, compliant",
    "符合要求\n": "Compliant\n",
    "现场为卡箍连接，符合要求": "Clamp connections on site, compliant",
    "符合要求": "Compliant",
    "符合要求\n内表面粗糙度：NMT 25Ra（微英寸）或 0.63 微米":
        "Compliant\nInternal surface finish: NMT 25 µin Ra or 0.63 µm",
    "有视镜灯": "Sight-glass light provided",
    "不符，无视镜灯": "Non-compliant, no sight-glass light",
    "不涉及": "Not applicable",
    "符合3D，最小死角设计": "3D compliant, minimum dead-leg design",
    "否，进水管路三通过长，存在清洁死角": "No, tee on the water inlet line is too long, cleaning dead leg present",
    "LT-001差压液位计": "LT-001 differential-pressure level transmitter",
    "否，差压液位计侧口安装手阀": "No, manual valve installed on the side port of the DP level transmitter",
    "储罐根部不满足3D需求": "Root connection at tank does not meet 3D",
    "PT-001远传压力表": "PT-001 remote-reading pressure gauge",
    "否，压力表下口安装手阀": "No, manual valve installed below the pressure gauge",
    "否，进料管线气动隔膜阀角度不符合要求，易存料；其余阀门角度合格":
        "No, pneumatic diaphragm valve angle on the feed line non-compliant, prone to product hold-up; "
        "other valve angles acceptable",
    "否，罐底出料和排液口三通长度不符合，存在清洁死角；":
        "No, tee lengths at tank bottom outlet and drain port non-compliant, cleaning dead legs present;",
    "TT-001温度计": "TT-001 temperature sensor",
    "否，温度计为内伸管": "No, thermometer is of the dip-tube type",
    "BPE不允许内伸温度计": "BPE does not allow dip-tube thermometers",
    "储罐根部不满足2D需求": "Root connection at tank does not meet 2D",
    "符合要求\n废液排放管路的止逆阀后端使用球阀，其他为隔膜阀":
        "Compliant\nBall valve downstream of the check valve on the waste drain line; all others are "
        "diaphragm valves",
    "符合，有视镜灯": "Compliant, sight-glass light provided",
    "储罐有下搅拌\n": "Tank has a bottom-mounted agitator\n",
    "工序": "Process Step",
    "管路起始": "Line From",
    "管路结束": "Line To",
    "材质/管径": "Material / Size",
    "液相": "Fluid",
    "管路标准": "Piping Standard",
    "连接方式排查项": "Connection Check",
    "阀门类型排查项": "Valve Type Check",
    "坡度排查项": "Slope Check",
    "不可排尽流量计排查项": "Non-drainable Flowmeter Check",
    "低点排尽排查项": "Low-Point Drain Check",
    "连接方式": "Connection Type",
    "阀门类型": "Valve Type",
    "坡度": "Slope",
    "不可排尽的流量计": "Non-drainable Flowmeter",
    "管路是否有低点排尽口": "Low-point drain on line?",
    "文件": "Documentation",
    "制备组分加入": "Prep fraction addition",
    "卫生卡箍": "Sanitary clamp",
    "1.管路可进行CIP清洗.可正洗、反洗；\n2.使用３D减少残留死区；":
        "1. Line can be CIP-cleaned, with forward and reverse flushing;\n2. " + DL_3D,
    "不符合": "Non-compliant",
    "坡度大于1%，实现完全排空": "Slope > 1%, fully drainable",
    "流量计竖直安装": "Flowmeter installed vertically",
    "符合": "Compliant",
    "管路低点满足3D需求；": "Line low points meet 3D;",
    "制备组分及甲醇溶剂加入": "Prep fraction and methanol addition",
    "1.管路可实现CIP；\n2.管路过滤器-离线清洗，过滤器短接进行CIP清洗；\n3.管道过滤器本体：增加排净点；\n4.使用３D减少残留死区；":
        f"1. Line is CIP-capable;\n2. {CIP_FILTER};\n3. {DRAIN_FILTER}\n4. {DL_3D}",
    "无低点排尽口\n": "No low-point drain\n",
    "不符合\n管路无低点排净\n过滤器无低点排净": "Non-compliant\nNo low-point drain on line\nNo low-point drain on filter",
    "乙腈/水加入": "ACN / water addition",
    "不符合要求，管路需更换为BPE管路；": "Non-compliant; line to be replaced with BPE tubing;",
    "1.溶剂管路通过N2进行吹扫，不进行在线CIP；\n2.管路过滤器-离线清洗；\n3.使用３D减少残留死区；":
        f"1. {N2_PURGE}\n2. Line filter cleaned offline;\n3. {DL_3D}",
    "1.增加管路低点可实现重力排净点；\n2.管道过滤器本体：增加排净点；": GRAVITY,
    "气动泵：泵前及泵后有排净，氮气吹扫；\n乙腈过滤器进口带排净口":
        PUMP + "\nACN filter inlet has a drain port",
    "无低点排净": "No low-point drain",
    "甲醇加入": "Methanol addition",
    "1.溶剂管路通过N2进行吹扫，不进行在线CIP；\n2.管路过滤器-离线清洗；\n3.桶装溶剂可直接进设备或通过喷淋球进设备;\n4.使用３D减少残留死区；":
        f"1. {N2_PURGE}\n2. Line filter cleaned offline;\n3. Drummed solvent can be fed directly to the equipment "
        f"or via a spray ball;\n4. {DL_3D}",
    "气动泵：泵前及泵后有排净，氮气吹扫；": PUMP,
    "氢氧化钠溶液加入": "NaOH solution addition",
    "1.管路可实现CIP；\n3.管路过滤器-离线清洗，过滤器短接进行CIP清洗；\n6.使用３D减少残留死区；":
        f"1. Line is CIP-capable;\n3. {CIP_FILTER};\n6. {DL_3D}",
    "1.增加管路低点可实现重力排净点；\n2.管道过滤器本体：增加排净点；\n": GRAVITY + "\n",
    "无低点排净\n不符合": "No low-point drain\nNon-compliant",
    "BLA Topic：09.02.2026 Questions - Ciara": "BLA Topic: 09.02.2026 Questions - Ciara",
}
CJK_FONTS = {"宋体", "等线", "汉仪书宋二KW", "Helvetica Neue"}


def si_text(si):
    return "".join(t.text or "" for t in si.iter(M + "t"))


def set_si(si, text):
    runs = si.findall(M + "r")
    rpr = copy.deepcopy(runs[0].find(M + "rPr")) if runs else None
    for child in list(si):
        si.remove(child)
    if rpr is not None:
        rfont = rpr.find(M + "rFont")
        if rfont is not None:
            rfont.set("val", FONT)
        r = etree.SubElement(si, M + "r")
        r.append(rpr)
        t = etree.SubElement(r, M + "t")
    else:
        t = etree.SubElement(si, M + "t")
    t.text = text
    t.set(XML_SPACE, "preserve")


def drop_chinese(si):
    """Remove the Chinese translation block of a bilingual question, keeping run formatting of the rest."""
    text = si_text(si)
    lines = text.split("\n")
    cjk = [i for i, line in enumerate(lines) if CJK.search(line)]
    start = sum(len(line) + 1 for line in lines[:cjk[0]])
    end = sum(len(line) + 1 for line in lines[:cjk[-1] + 1])
    if end < len(text) and text[end] == "\n":
        end += 1
    pos = 0
    for t in si.iter(M + "t"):
        s = t.text or ""
        a, b = pos, pos + len(s)
        t.text = s[:max(0, start - a)] + s[max(0, end - a):]
        t.set(XML_SPACE, "preserve")
        pos = b
    for r in si.findall(M + "r"):
        if not r.find(M + "t").text:
            si.remove(r)
    last = list(si.iter(M + "t"))[-1]
    last.text = last.text.rstrip()


def build_xlsx():
    zin = zipfile.ZipFile(XLS_SRC)
    sst = etree.fromstring(zin.read("xl/sharedStrings.xml"))
    missing = []
    for si in sst.findall(M + "si"):
        text = si_text(si)
        if not CJK.search(text):
            continue
        if text in STRINGS:
            set_si(si, STRINGS[text])
        elif re.match(r"\d+\. [A-Za-z]", text):
            drop_chinese(si)
        else:
            missing.append(text)
        assert not CJK.search(si_text(si)), si_text(si)
    assert not missing, missing
    for rfont in sst.iter(M + "rFont"):
        if rfont.get("val") in CJK_FONTS:
            rfont.set("val", FONT)

    styles = etree.fromstring(zin.read("xl/styles.xml"))
    for font in styles.find(M + "fonts"):
        name = font.find(M + "name")
        if name is not None and name.get("val") in CJK_FONTS:
            name.set("val", FONT)
            scheme = font.find(M + "scheme")
            if scheme is not None:
                font.remove(scheme)

    new = {
        "xl/sharedStrings.xml": etree.tostring(sst, xml_declaration=True, encoding="UTF-8", standalone=True),
        "xl/styles.xml": etree.tostring(styles, xml_declaration=True, encoding="UTF-8", standalone=True),
    }
    with zipfile.ZipFile(XLS_OUT, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, new.get(item.filename, zin.read(item.filename)))
    print(XLS_OUT.relative_to(ROOT))


if __name__ == "__main__":
    build_pptx()
    build_xlsx()
