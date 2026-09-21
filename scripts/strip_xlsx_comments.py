#!/usr/bin/env python3
"""Remove Excel notes/comments and column-entry popups from a workbook."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook


def strip_comments(path: Path) -> int:
    wb = load_workbook(path)
    removed = 0
    for name in list(wb.defined_names.keys()):
        # leftover names from the old multi-sheet workbook
        if name.startswith("Milestone"):
            del wb.defined_names[name]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    cell.comment = None
                    removed += 1
        if hasattr(ws, "_comments"):
            ws._comments = []
        for dv in ws.data_validations.dataValidation:
            dv.prompt = None
            dv.promptTitle = None
            dv.showInputMessage = False
    wb.save(path)
    return removed


def main() -> None:
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "meetings/POA生物制品项目改造-行动项跟踪 (1).xlsx")
    dest = Path(sys.argv[2] if len(sys.argv) > 2 else "meetings/BLA改造行动项跟踪.xlsx")
    removed = strip_comments(src)
    if dest.resolve() != src.resolve():
        shutil.copy2(src, dest)
    print(f"removed {removed} comments; wrote {src} and {dest}")


if __name__ == "__main__":
    main()
