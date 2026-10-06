#!/usr/bin/env python3
"""English versions of 9788项目焊缝统计汇报.pptx and 10-06.xlsx (weld inspection plan and photo details).

PPTX: logos and the "Confidential | 保密文件" / website footers stay as in the source; all content text becomes
English in Times New Roman. The embedded Excel objects on slides 3-5 (Chinese preview images) are replaced by
native tables, with the original JPEG photos placed unchanged (same bytes, aspect ratio and rotation).

XLSX: edited at XML level so the photos keep their original bytes and rotation; strings are translated, fonts
changed to Times New Roman at presentation size, and columns, rows and photo anchors resized to match.

Usage: python3 scripts/build_9788_weld_summary_en.py
"""
import io
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import date, timedelta
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_139788_retrofit_v2_en import A, cell_key, clean_rpr, set_text, text_width_pt, translate_title  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SRC_PPTX = ROOT / "9788项目焊缝统计汇报.pptx"
SRC_XLSX = ROOT / "10-06.xlsx"
OUT_PPTX = ROOT / "deliverables" / "9788项目焊缝统计汇报_EN.pptx"
OUT_XLSX = ROOT / "deliverables" / "10-06_EN.xlsx"

FONT = "Times New Roman"
CJK = re.compile(r"[\u3400-\u9fff（）、，。！]")

# ---------------------------------------------------------------- shared terms (PPT slide 2 and the workbook)
TERMS = {
    # overall plan
    "步骤": "Step",
    "设备编号": "Equipment No.",
    "材质、规格": "Material / Size",
    "管路": "Line",
    "目前进展": "Current Status",
    "计划完成时间": "Planned Completion",
    "step1": "Step 1",
    "step2": "Step 2",
    "搪瓷搅拌罐、500L": "Glass-lined agitated tank, 500 L",
    "搪瓷反应釜、2000L": "Glass-lined reactor, 2000 L",
    "哈氏合金多肽固相合成仪、2000L": "Hastelloy solid-phase peptide synthesizer, 2000 L",
    "316L激活釜、1000L": "316L activation reactor, 1000 L",
    "哈氏合金压滤罐、DN900": "Hastelloy pressure filter, DN900",
    "白钢搅拌罐、3000L": "Stainless steel agitated tank, 3000 L",
    "哈氏合金反应釜、8000L": "Hastelloy reactor, 8000 L",
    "哈氏合金三合一、DN1600": "Hastelloy agitated filter dryer (AFD), DN1600",
    "甲苯管线": "Toluene line",
    "DMF管线": "DMF line",
    "脱保护溶液管线": "Deprotection solution line",
    "1213-AT08管线": "1213-AT08 line",
    "MTBE管线": "MTBE line",
    "IPAC管线": "IPAC line",
    "1213-R23反应釜管线": "1213-R23 reactor line",
    "激活罐转移合成仪管线": "Activation tank to synthesizer transfer line",
    "非专用溶剂泵管线": "Non-dedicated solvent pump line",
    "循环清洗管路": "Recirculation cleaning line",
    "自循环管线": "Self-recirculation line",
    "进行中": "In progress",
    "已完成": "Completed",
    "内衬PTFE材质、不涉及": "PTFE-lined, not applicable",
    "管路为搪瓷、内衬PTFE": "Glass-lined / PTFE-lined piping",
    # photo details
    "设备设施拍摄照片明细及说明": "Weld Inspection Photos of Equipment and Facilities: Details and Rationale",
    "序号": "No.",
    "管线及配件类别": "Piping / Fitting Category",
    "焊接形式": "Weld Type",
    "执行方式": "Inspection Approach",
    "评估": "Assessment",
    "外表面照片": "External Surface Photo",
    "内窥镜照片": "Borescope Photo",
    "验收记录": "Acceptance Record",
    "标准管路\n（基于工艺需求组装的管路）": "Standard piping\n(spools assembled per process requirements)",
    "设备间固定转移管路\n（与设备同步安装）": "Fixed inter-equipment transfer piping\n(installed together with the equipment)",
    "标准配件（含自制）": "Standard fittings\n(incl. in-house fabricated)",
    "设备罐口（Head）": "Equipment nozzles (Head)",
    "溶剂管路": "Solvent piping",
    "3、4类管路": "Category 3 and 4 piping",
    "手动焊": "Manual weld",
    "自动焊": "Automatic weld",
    "自动焊、手动焊": "Automatic and manual welds",
    "不检查": "Not inspected",
    "本次全部进行拆卸检查焊缝": "All dismantled for weld inspection",
    "该部分按20%进行抽检": "20% sampling inspection",
    "本次从Head依次检查到第一道阀门处": "Inspected from the head up to the first valve",
    "从Head向主管路进行内窥检查，达到内窥镜的最远距离":
        "Borescope inspection from the head toward the main line, up to the maximum reach of the borescope",
    "本次不额外检查": "No additional inspection",
    "标准管路先经过清洗、验收、QA放行、然后进行安装，由于该部分每次安装前都会进行检查、放行，因此本次不进行检查。":
        "Standard piping is cleaned, accepted and released by QA before installation. Because it is inspected "
        "and released before every installation, it is not included in this inspection.",
    "标准配件为自动焊，经过抛光处理。": "Standard fittings are automatically welded and polished.",
    "自制配件经过抛光处理，整体焊缝处理的标准较高，出现不合格的几率较低。":
        "In-house fabricated fittings are polished; the overall weld finishing standard is high, so the "
        "likelihood of nonconformance is low.",
    "管路安装前已经进行了焊缝检查，自动焊20%，手动焊100%，相应图片已经存档":
        "Welds were inspected before the piping was installed (automatic welds 20%, manual welds 100%); "
        "the corresponding photos are archived.",
}
TRANSFER_LINE = re.compile(r"^(TJ4S-\d{4}-[A-Z]+\d+)转移管线$")
CN_DATE = re.compile(r"^10\s*月\s*(\d{1,2})\s*日$")
SHEET_NAMES = {"整体计划": "Overall Plan", "反馈明细": "Weld Photo Details"}


def tr(text):
    key = text.strip()
    if not CJK.search(key) and key not in TERMS:
        return key
    if key in TERMS:
        return TERMS[key]
    m = TRANSFER_LINE.match(key)
    if m:
        return f"{m.group(1)} transfer line"
    m = CN_DATE.match(key)
    if m:
        return f"Oct {int(m.group(1))}"
    raise KeyError(key)


# ---------------------------------------------------------------- photo-detail rows (sheet 反馈明细, rows 3-9)
# (No., category, weld type, approach, assessment); None = covered by the cell above (merged)
DETAIL_ROWS = [
    ("1", "标准管路\n（基于工艺需求组装的管路）", "手动焊", "不检查",
     "标准管路先经过清洗、验收、QA放行、然后进行安装，由于该部分每次安装前都会进行检查、放行，因此本次不进行检查。"),
    ("2", "设备间固定转移管路\n（与设备同步安装）", "手动焊", "本次全部进行拆卸检查焊缝", "N/A"),
    ("3", "标准配件（含自制）", "自动焊", "该部分按20%进行抽检", "标准配件为自动焊，经过抛光处理。"),
    ("4", None, "手动焊", "该部分按20%进行抽检", "自制配件经过抛光处理，整体焊缝处理的标准较高，出现不合格的几率较低。"),
    ("5", "设备罐口（Head）", "手动焊", "本次从Head依次检查到第一道阀门处", "N/A"),
    ("6", "溶剂管路", "手动焊", "从Head向主管路进行内窥检查，达到内窥镜的最远距离", "N/A"),
    ("7", "3、4类管路", "自动焊、手动焊", "本次不额外检查",
     "管路安装前已经进行了焊缝检查，自动焊20%，手动焊100%，相应图片已经存档"),
]
# drawing rId in 10-06.xlsx -> (0-based sheet row, 0-based column F/G/H = 5/6/7)
PHOTO_CELLS = {"rId8": (2, 7), "rId3": (3, 5), "rId4": (3, 6), "rId1": (4, 5), "rId2": (4, 6), "rId6": (5, 5),
               "rId7": (5, 6), "rId9": (6, 5), "rId5": (6, 6), "rId11": (7, 5), "rId10": (7, 6)}
DETAIL_HEADER = ["序号", "管线及配件类别", "焊接形式", "执行方式", "评估", "外表面照片", "内窥镜照片", "验收记录"]


def load_photos():
    """{(row, col): (jpeg bytes, rotation in degrees)} from the source workbook drawing."""
    with zipfile.ZipFile(SRC_XLSX) as z:
        drawing = z.read("xl/drawings/drawing1.xml").decode()
        rels = z.read("xl/drawings/_rels/drawing1.xml.rels").decode()
        targets = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="\.\./([^"]+)"', rels))
        targets.update({k: v for v, k in re.findall(r'Target="\.\./([^"]+)"[^>]*Id="(rId\d+)"', rels)})
        photos = {}
        for anchor in re.findall(r"<xdr:twoCellAnchor.*?</xdr:twoCellAnchor>", drawing, re.S):
            rid = re.search(r'r:embed="(rId\d+)"', anchor).group(1)
            rot = re.search(r'<a:xfrm rot="(-?\d+)"', anchor)
            photos[PHOTO_CELLS[rid]] = (z.read("xl/" + targets[rid]), int(rot.group(1)) / 60000 if rot else 0)
    assert len(photos) == len(PHOTO_CELLS)
    return photos


def visual_size(blob, rot):
    w, h = Image.open(io.BytesIO(blob)).size
    return (h, w) if round(rot) % 180 == 90 else (w, h)


# ================================================================ PPTX
TITLES = {
    1: "Project 9788 Weld Inspection Summary",
    2: "Overall Execution Plan",
    3: "Weld Inspection Photo Summary",
    6: "THANK YOU!",
}
PLAN_SZ = 20
PLAN_MIN_COL_W = [1300000, 2900000]

# photo slides: slide number -> (detail row indices, show caption row)
PHOTO_SLIDES = {3: ([0, 1], True), 4: ([2, 3], False), 5: ([4, 5, 6], False)}
PH_LEFT, PH_TOP, PH_BOTTOM = 900000, 1870000, 12760000
PH_COL_W = [900000, 2650000, 1900000, 2700000, 3350000, 3700000, 3700000, 3700000]
PH_HEAD_H = 900000
PH_TWO_ROW_H = 4540000
PH_BODY_SZ, PH_HEAD_SZ, PH_CAPTION_SZ = 22, 20, 24
PH_PAD = 110000
LINE_W = 6350
HEAD_FILL, CAPTION_FILL = "BDD7EE", "9BC2E6"


def fit_plan_columns(table):
    """Rebalance slide-2 column widths so single-row English cells stay on one line."""
    total = sum(c.width for c in table.columns)
    need = [0] * len(table.columns)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            if cell.is_spanned or (cell.is_merge_origin and cell.span_height > 2):
                continue
            text = cell_key(cell)
            need[c] = max(need[c], int(Pt(text_width_pt(text, PLAN_SZ, bold=r == 0) * 1.06 + 8)))
    for i, w in enumerate(PLAN_MIN_COL_W):
        need[i] = max(need[i], w)
    spare = total - sum(need)
    assert spare >= 0, (need, total)
    widths = [n + spare * n // sum(need) for n in need]
    widths[-1] += total - sum(widths)
    for col, w in zip(table.columns, widths):
        col.width = Emu(w)


def translate_plan(table):
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            if cell.is_spanned:
                continue
            key = cell_key(cell)
            if key:
                set_text(cell._tc.txBody, [tr(key)])
    fit_plan_columns(table)


def set_cell_border(cell, color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(LINE_W), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(fill)
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "solid"}))
        tcPr.append(ln)
    # DrawingML requires the borders before the cell fill
    fill = tcPr.find(qn("a:solidFill"))
    if fill is not None:
        tcPr.remove(fill)
        tcPr.append(fill)


def write_cell(cell, text, size, bold=False, fill="FFFFFF", align=PP_ALIGN.CENTER):
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    set_cell_border(cell)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Emu(70000)
    cell.margin_top = cell.margin_bottom = Emu(40000)
    tf = cell.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(0, 0, 0)
        clean_rpr(run._r.get_or_add_rPr())


def remove_ole(slide):
    for shape in list(slide.shapes):
        if shape.shape_type == 7:  # EMBEDDED_OLE_OBJECT
            shape._element.getparent().remove(shape._element)
    xml = slide._element.xml
    for rel in list(slide.part.rels.values()):
        if rel.reltype.endswith(("/oleObject", "/vmlDrawing", "/image")) and f'"{rel.rId}"' not in xml:
            slide.part.drop_rel(rel.rId)


def place_photo(slide, blob, rot, x, y, w, h):
    """Fit the photo (as displayed, i.e. after rotation) into the box, centred, without changing its aspect."""
    vw, vh = visual_size(blob, rot)
    s = min((w - 2 * PH_PAD) / vw, (h - 2 * PH_PAD) / vh)
    dw, dh = int(vw * s), int(vh * s)
    cx, cy = x + w // 2, y + h // 2
    if round(rot) % 180 == 90:
        dw, dh = dh, dw
    pic = slide.shapes.add_picture(io.BytesIO(blob), Emu(cx - dw // 2), Emu(cy - dh // 2), Emu(dw), Emu(dh))
    pic.rotation = rot
    return pic


def build_photo_slide(slide, rows, caption, photos):
    remove_ole(slide)
    n_body = len(rows)
    heights = ([PH_HEAD_H] if caption else []) + [PH_HEAD_H]
    body_h = PH_TWO_ROW_H if n_body == 2 else (PH_BOTTOM - PH_TOP - sum(heights)) // n_body
    heights += [body_h] * n_body
    width = sum(PH_COL_W)
    gf = slide.shapes.add_table(len(heights), len(PH_COL_W), Emu(PH_LEFT), Emu(PH_TOP), Emu(width),
                                Emu(sum(heights)))
    gf.name = "Weld photo table"
    table = gf.table
    tblPr = table._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # No Style, Table Grid
    for i, w in enumerate(PH_COL_W):
        table.columns[i].width = Emu(w)
    for i, h in enumerate(heights):
        table.rows[i].height = Emu(h)

    r = 0
    if caption:
        table.cell(0, 0).merge(table.cell(0, len(PH_COL_W) - 1))
        write_cell(table.cell(0, 0), tr("设备设施拍摄照片明细及说明"), PH_CAPTION_SZ, bold=True, fill=CAPTION_FILL)
        r = 1
    for c, key in enumerate(DETAIL_HEADER):
        write_cell(table.cell(r, c), tr(key), PH_HEAD_SZ, bold=True, fill=HEAD_FILL)
    first_body = r + 1
    for i, d in enumerate(rows):
        tr_i = first_body + i
        no, cat, weld, approach, assess = DETAIL_ROWS[d]
        write_cell(table.cell(tr_i, 0), no, PH_BODY_SZ, bold=True)
        if cat is not None:
            span = 1
            while i + span < len(rows) and DETAIL_ROWS[rows[i + span]][1] is None:
                span += 1
            if span > 1:
                table.cell(tr_i, 1).merge(table.cell(tr_i + span - 1, 1))
            write_cell(table.cell(tr_i, 1), tr(cat), PH_BODY_SZ)
        write_cell(table.cell(tr_i, 2), tr(weld), PH_BODY_SZ)
        write_cell(table.cell(tr_i, 3), tr(approach), PH_BODY_SZ)
        write_cell(table.cell(tr_i, 4), tr(assess), PH_BODY_SZ,
                   align=PP_ALIGN.CENTER if assess == "N/A" else PP_ALIGN.LEFT)
        for c in (5, 6, 7):
            write_cell(table.cell(tr_i, c), "" if (d + 2, c) in photos else "N/A", PH_BODY_SZ)

    xs = [PH_LEFT + sum(PH_COL_W[:c]) for c in range(len(PH_COL_W))]
    for i, d in enumerate(rows):
        y = PH_TOP + sum(heights[:first_body + i])
        for c in (5, 6, 7):
            if (d + 2, c) in photos:
                blob, rot = photos[(d + 2, c)]
                place_photo(slide, blob, rot, xs[c], y, PH_COL_W[c], body_h)


def fit_cover_title(shape, text):
    tb = shape.text_frame._txBody
    size = int(next(tb.iter(A + "rPr")).get("sz")) / 100
    usable = Emu(shape.width).pt * 0.94
    size = min(size, int(usable / text_width_pt(text, 1)))
    set_text(tb, [text], size)


def build_pptx(photos):
    prs = Presentation(SRC_PPTX)
    for n, slide in enumerate(prs.slides, 1):
        for shape in list(slide.shapes):
            if shape.has_text_frame and CJK.search(shape.text_frame.text) and "保密文件" not in shape.text_frame.text:
                if n == 1:
                    fit_cover_title(shape, TITLES[1])
                elif n == 6:
                    set_text(shape.text_frame._txBody, [TITLES[6]])
                else:
                    translate_title(shape, TITLES[min(n, 3)])
            elif getattr(shape, "has_table", False) and shape.has_table:
                translate_plan(shape.table)
        if n in PHOTO_SLIDES:
            rows, caption = PHOTO_SLIDES[n]
            build_photo_slide(slide, rows, caption, photos)
    prs.save(OUT_PPTX)


# ================================================================ XLSX
X_BODY_SZ, X_TITLE_SZ = 18, 22
DIGIT_PX = 7  # max digit width of the 11 pt default font; column widths are in these units
X_PAD_PX = 10
X_PHOTO_BOX = (380, 330)  # max displayed photo size in px at 100% zoom; sources are ~430-540 px
X_LINE_PT = 1.28
DATE_FMT_ID = 164
PLAN_COL_MIN = {0: 12}
DETAIL_TEXT_COLS = {0: 8, 1: 32, 2: 22, 3: 32, 4: 46}


def col_px(width):
    return int(((256 * width + int(128 / DIGIT_PX)) / 256) * DIGIT_PX)


def px_to_width(px):
    return round((px - 5) / DIGIT_PX + 5 / DIGIT_PX, 2)


_FONTS = {}


def text_px(text, size, bold):
    key = (size, bold)
    if key not in _FONTS:
        style = "Bold" if bold else "Regular"
        _FONTS[key] = ImageFont.truetype(f"/usr/share/fonts/truetype/liberation/LiberationSerif-{style}.ttf", 1000)
    return _FONTS[key].getlength(text) * size / 1000 * 96 / 72


def wrapped_lines(text, size, bold, avail_px):
    n = 0
    for para in text.split("\n"):
        line = ""
        n += 1
        for word in para.split(" "):
            trial = (line + " " + word).strip()
            if line and text_px(trial, size, bold) * 1.08 > avail_px:
                n += 1
                line = word
            else:
                line = trial
    return n


def row_pt(lines, size):
    return round(lines * size * X_LINE_PT + 10, 1)


def parse_sheet(xml):
    """{(row, col): (style, shared-string index or None)} for cells with a value or style."""
    cells = {}
    for ref, attrs, body in re.findall(r'<c r="([A-Z]+\d+)"([^>]*?)(?:/>|>(.*?)</c>)', xml):
        col = ord(ref[0]) - 65
        row = int(ref[1:]) - 1
        s = re.search(r's="(\d+)"', attrs)
        v = re.search(r"<v>(.*?)</v>", body or "")
        is_str = 't="s"' in attrs
        cells[(row, col)] = (int(s.group(1)) if s else 0, int(v.group(1)) if (v and is_str) else None,
                             v.group(1) if v else None)
    return cells


def merged(xml):
    out = []
    for a, b in re.findall(r'<mergeCell ref="([A-Z]+\d+):([A-Z]+\d+)"/>', xml):
        out.append(((int(a[1:]) - 1, ord(a[0]) - 65), (int(b[1:]) - 1, ord(b[0]) - 65)))
    return out


def set_cols(xml, widths):
    cols = "".join(f'<col min="{i + 1}" max="{i + 1}" width="{w}" style="4" customWidth="1"/>'
                   for i, w in enumerate(widths))
    cols += f'<col min="{len(widths) + 1}" max="16384" width="8.6640625" style="4"/>'
    return re.sub(r"<cols>.*?</cols>", f"<cols>{cols}</cols>", xml, flags=re.S)


def set_row_heights(xml, heights):
    def repl(m):
        r = int(m.group(1))
        tag = re.sub(r' ht="[^"]*"', f' ht="{heights[r - 1]}"', m.group(0))
        return tag
    return re.sub(r'<row r="(\d+)"[^>]*>', repl, xml)


def add_row_breaks(xml, after_rows):
    brks = "".join(f'<brk id="{r}" max="16383" man="1"/>' for r in after_rows)
    tag = f'<rowBreaks count="{len(after_rows)}" manualBreakCount="{len(after_rows)}">{brks}</rowBreaks>'
    return re.sub(r"(<pageSetup [^>]*/>)", r"\1" + tag, xml, count=1)


def set_view(xml, zoom):
    xml = re.sub(r"<sheetView [^>]*>", lambda m: re.sub(
        r' (topLeftCell|zoomScale|zoomScaleNormal)="[^"]*"', "", m.group(0)).replace(
        "<sheetView ", f'<sheetView zoomScale="{zoom}" zoomScaleNormal="{zoom}" '), xml)
    xml = re.sub(r'<selection [^>]*/>', '<selection activeCell="A1" sqref="A1"/>', xml)
    xml = re.sub(r"<pageSetup ([^>]*)/>", lambda m: "<pageSetup " + re.sub(
        r' ?orientation="[^"]*"', "", m.group(1)) + ' orientation="landscape" fitToHeight="0"/>', xml)
    if "<sheetPr" not in xml:
        xml = xml.replace("<dimension ", '<sheetPr><pageSetUpPr fitToPage="1"/></sheetPr><dimension ', 1)
    return xml


def layout_plan(xml, strings, cell_text):
    cells = parse_sheet(xml)
    spans = merged(xml)
    tall = {s for s in spans if s[1][0] - s[0][0] >= 2}
    widths_px = [0] * 6
    for (r, c), (style, si, raw) in cells.items():
        text = cell_text(si, raw, style)
        # tall merged Step / Material cells may wrap; equipment numbers stay on one line
        if not text or (c != 1 and any(a == (r, c) for a, _ in tall)):
            continue
        widths_px[c] = max(widths_px[c], text_px(text, X_BODY_SZ, bold=r == 0) * 1.08 + 2 * X_PAD_PX + 8)
    widths = [max(PLAN_COL_MIN.get(i, 0), px_to_width(px)) for i, px in enumerate(widths_px)]
    n_rows = max(r for r, _ in cells) + 1
    heights = [round(X_BODY_SZ * X_LINE_PT + 14, 1)] * n_rows
    heights[0] = round(X_BODY_SZ * X_LINE_PT + 22, 1)
    for (a, b) in tall:  # merged equipment cells: make sure wrapped text fits the merged height
        style, si, raw = cells[a]
        text = cell_text(si, raw, style)
        lines = wrapped_lines(text, X_BODY_SZ, False, col_px(widths[a[1]]) - 2 * X_PAD_PX)
        need = row_pt(lines, X_BODY_SZ)
        have = sum(heights[a[0]:b[0] + 1])
        if need > have:
            extra = round((need - have) / (b[0] - a[0] + 1) + 0.5, 1)
            for r in range(a[0], b[0] + 1):
                heights[r] += extra
    for (a, b) in spans:
        if (a, b) not in tall and a[1] > 0:
            style, si, raw = cells[a]
            text = cell_text(si, raw, style)
            lines = wrapped_lines(text, X_BODY_SZ, False, col_px(widths[a[1]]) - 2 * X_PAD_PX)
            have = sum(heights[a[0]:b[0] + 1])
            need = row_pt(lines, X_BODY_SZ)
            if need > have:
                for r in range(a[0], b[0] + 1):
                    heights[r] = round(heights[r] + (need - have) / (b[0] - a[0] + 1) + 0.5, 1)
    xml = set_cols(xml, widths)
    xml = set_row_heights(xml, heights)
    return add_row_breaks(set_view(xml, 100), [18])


def layout_details(xml, strings, cell_text, photos):
    cells = parse_sheet(xml)
    photo_col_px = X_PHOTO_BOX[0] + 2 * X_PAD_PX
    widths = [DETAIL_TEXT_COLS[c] for c in range(5)] + [px_to_width(photo_col_px)] * 3
    n_rows = 9
    heights = [0.0] * n_rows
    heights[0] = round(X_TITLE_SZ * X_LINE_PT + 18, 1)
    for r in range(1, n_rows):
        lines = 1
        for c in range(8):
            if (r, c) not in cells:
                continue
            style, si, raw = cells[(r, c)]
            text = cell_text(si, raw, style)
            if not text:
                continue
            if r == 4 and c == 1:  # merged B5:B6, wrapped over two rows
                continue
            lines = max(lines, wrapped_lines(text, X_BODY_SZ, r == 1, col_px(widths[c]) - 2 * X_PAD_PX))
        heights[r] = row_pt(lines, X_BODY_SZ)
        if any((r, c) in photos for c in (5, 6, 7)):
            heights[r] = max(heights[r], round((X_PHOTO_BOX[1] + 2 * X_PAD_PX) * 72 / 96, 1))
    xml = set_cols(xml, widths)
    xml = set_row_heights(xml, heights)
    return add_row_breaks(set_view(xml, 85), [4, 6]), widths, heights


def relayout_drawing(drawing, photos, widths, heights):
    """Re-anchor every photo centred in its cell at the largest size that fits X_PHOTO_BOX, keeping aspect."""
    emu = 9525
    col_x = [0]
    for w in widths:
        col_x.append(col_x[-1] + col_px(w) * emu)
    row_y = [0]
    for h in heights:
        row_y.append(row_y[-1] + int(h * 12700))

    def locate(pos, edges):
        i = max(k for k in range(len(edges) - 1) if edges[k] <= pos)
        return i, pos - edges[i]

    def repl(m):
        anchor = m.group(0)
        rid = re.search(r'r:embed="(rId\d+)"', anchor).group(1)
        r, c = PHOTO_CELLS[rid]
        blob, rot = photos[(r, c)]
        vw, vh = visual_size(blob, rot)
        s = min(X_PHOTO_BOX[0] / vw, X_PHOTO_BOX[1] / vh)
        vw_e, vh_e = int(vw * s) * emu, int(vh * s) * emu
        cx = (col_x[c] + col_x[c + 1]) // 2
        cy = (row_y[r] + row_y[r + 1]) // 2
        x0, y0, x1, y1 = cx - vw_e // 2, cy - vh_e // 2, cx + vw_e // 2, cy + vh_e // 2
        fc, fco = locate(x0, col_x)
        fr, fro = locate(y0, row_y)
        tc, tco = locate(x1, col_x)
        tr_, tro = locate(y1, row_y)
        anchor = re.sub(r"<xdr:from>.*?</xdr:from>",
                        f"<xdr:from><xdr:col>{fc}</xdr:col><xdr:colOff>{fco}</xdr:colOff><xdr:row>{fr}</xdr:row>"
                        f"<xdr:rowOff>{fro}</xdr:rowOff></xdr:from>", anchor, flags=re.S)
        anchor = re.sub(r"<xdr:to>.*?</xdr:to>",
                        f"<xdr:to><xdr:col>{tc}</xdr:col><xdr:colOff>{tco}</xdr:colOff><xdr:row>{tr_}</xdr:row>"
                        f"<xdr:rowOff>{tro}</xdr:rowOff></xdr:to>", anchor, flags=re.S)
        # the anchor is the displayed box; a:xfrm holds the unrotated frame
        ew, eh = (vh_e, vw_e) if round(rot) % 180 == 90 else (vw_e, vh_e)
        anchor = re.sub(r'<a:off x="-?\d+" y="-?\d+"/><a:ext cx="\d+" cy="\d+"/>',
                        f'<a:off x="{cx - ew // 2}" y="{cy - eh // 2}"/><a:ext cx="{ew}" cy="{eh}"/>', anchor)
        return anchor

    return re.sub(r"<xdr:twoCellAnchor.*?</xdr:twoCellAnchor>", repl, drawing, flags=re.S)


def build_styles(xml):
    xml = xml.replace("<fonts ", f'<numFmts count="1"><numFmt numFmtId="{DATE_FMT_ID}" formatCode="mmm d"/>'
                                 "</numFmts><fonts ", 1) if "<numFmts" not in xml else xml
    fonts = (
        '<fonts count="6"><font><sz val="11"/><color theme="1"/><name val="Times New Roman"/><family val="1"/></font>'
        '<font><sz val="9"/><name val="Times New Roman"/><family val="1"/></font>'
        '<font><b/><sz val="11"/><color theme="1"/><name val="Times New Roman"/><family val="1"/></font>'
        f'<font><sz val="{X_BODY_SZ}"/><color theme="1"/><name val="Times New Roman"/><family val="1"/></font>'
        f'<font><b/><sz val="{X_BODY_SZ}"/><color theme="1"/><name val="Times New Roman"/><family val="1"/></font>'
        f'<font><b/><sz val="{X_TITLE_SZ}"/><color theme="1"/><name val="Times New Roman"/><family val="1"/></font>'
        "</fonts>")
    xml = re.sub(r"<fonts .*?</fonts>", fonts, xml, flags=re.S)

    def fix_xf(xf):
        xf = re.sub(r'fontId="0"', 'fontId="3"', xf)
        xf = re.sub(r'fontId="2"', 'fontId="4"', xf)
        xf = xf.replace('numFmtId="58"', f'numFmtId="{DATE_FMT_ID}"')
        xf = re.sub(r"<alignment ([^/]*?)( wrapText=\"1\")?/>", r'<alignment \1 wrapText="1"/>', xf)
        xf = re.sub(r"<alignment (?![^>]*horizontal=)", '<alignment horizontal="left" indent="1" ', xf)
        if "applyFont" not in xf:
            xf = xf.replace("<xf ", '<xf applyFont="1" ', 1)
        return xf

    cellxfs = re.search(r"<cellXfs .*?</cellXfs>", xml, re.S).group(0)
    xfs = re.findall(r"<xf [^>]*/>|<xf [^>]*>.*?</xf>", cellxfs, re.S)
    new = [fix_xf(x) for x in xfs]
    title_xf = re.sub(r'fontId="\d+"', 'fontId="5"', new[21])
    new.append(title_xf)
    xml = xml.replace(cellxfs, f'<cellXfs count="{len(new)}">' + "".join(new) + "</cellXfs>")
    return xml, len(new) - 1


def build_xlsx(photos):
    tmp = Path(tempfile.mkdtemp())
    with zipfile.ZipFile(SRC_XLSX) as z:
        z.extractall(tmp)
        names = z.namelist()

    sst_path = tmp / "xl" / "sharedStrings.xml"
    sst = sst_path.read_text(encoding="utf-8")
    items = re.findall(r"<si>(.*?)</si>", sst, re.S)
    raw = ["".join(re.findall(r"<t[^>]*>(.*?)</t>", it, re.S)) for it in items]
    strings = [tr(s.strip("\n")) for s in [re.sub(r"&amp;", "&", s) for s in raw]]
    body = "".join(f'<si><t xml:space="preserve">{escape(s)}</t></si>' for s in strings)
    sst = re.sub(r"(<sst [^>]*>).*(</sst>)", lambda m: m.group(1) + body + m.group(2), sst, flags=re.S)
    sst_path.write_text(sst, encoding="utf-8")

    def cell_text(si, rawv, style):
        if si is not None:
            return strings[si]
        if rawv is None:
            return ""
        if style == 8 and re.fullmatch(r"\d+(\.0)?", rawv):
            day = date(1899, 12, 30) + timedelta(days=int(float(rawv)))
            return f"{day:%b} {day.day}"
        return rawv

    styles_path = tmp / "xl" / "styles.xml"
    styles, title_xf = build_styles(styles_path.read_text(encoding="utf-8"))
    styles_path.write_text(styles, encoding="utf-8")

    s1 = tmp / "xl" / "worksheets" / "sheet1.xml"
    s1.write_text(layout_plan(s1.read_text(encoding="utf-8"), strings, cell_text), encoding="utf-8")

    s2 = tmp / "xl" / "worksheets" / "sheet2.xml"
    xml = s2.read_text(encoding="utf-8")
    xml = re.sub(r'(<c r="[A-H]1" )s="21"', rf'\1s="{title_xf}"', xml)
    xml, widths, heights = layout_details(xml, strings, cell_text, photos)
    s2.write_text(xml, encoding="utf-8")

    d = tmp / "xl" / "drawings" / "drawing1.xml"
    d.write_text(relayout_drawing(d.read_text(encoding="utf-8"), photos, widths, heights), encoding="utf-8")

    for rel in ("xl/workbook.xml", "docProps/app.xml"):
        p = tmp / rel
        text = p.read_text(encoding="utf-8")
        for cn, en in SHEET_NAMES.items():
            text = text.replace(f'"{cn}"', f'"{en}"').replace(f">{cn}<", f">{en}<")
        text = text.replace(">工作表<", ">Worksheets<")
        if rel == "xl/workbook.xml" and "<definedNames>" not in text:
            names_xml = "".join(
                f"<definedName name=\"_xlnm.Print_Titles\" localSheetId=\"{i}\">'{SHEET_NAMES[cn]}'!{rows}</definedName>"
                for i, (cn, rows) in enumerate((("整体计划", "$1:$1"), ("反馈明细", "$2:$2"))))
            text = text.replace("</sheets>", f"</sheets><definedNames>{names_xml}</definedNames>", 1)
        p.write_text(text, encoding="utf-8")

    with zipfile.ZipFile(OUT_XLSX, "w", zipfile.ZIP_DEFLATED) as out:
        for name in names:
            out.write(tmp / name, name)
    shutil.rmtree(tmp)


def main():
    OUT_PPTX.parent.mkdir(exist_ok=True)
    photos = load_photos()
    build_pptx(photos)
    build_xlsx(photos)
    for p in (OUT_PPTX, OUT_XLSX):
        print(p.relative_to(ROOT))


if __name__ == "__main__":
    main()
