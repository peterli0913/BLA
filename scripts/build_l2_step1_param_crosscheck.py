#!/usr/bin/env python3
"""Cross-check L2 Step 1 parameters across the PVP, the parameter workbook, and the MOR; write findings to Excel."""

import re
import zipfile
from pathlib import Path

import openpyxl
from lxml import etree
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path("/workspace")
PVP = ROOT / "L2 Step1 PVP-1400200301.01_-10-28-2026.1.docx"
MOR = ROOT / "MOR_10282026.docx"
OUT = ROOT / "L2-Step1-PVP-参数表-MOR-三方核对_10-28-2026.xlsx"

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def accepted_text(path):
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    parts = []

    def walk(node):
        if node.tag == W + "del":
            return
        if node.tag == W + "t" and node.text:
            parts.append(node.text)
        for child in node:
            walk(child)

    walk(root)
    return re.sub(r"\s+", " ", "".join(parts))


THREE_WAY = [
    (
        "脱保护后 IPAc 洗涤量（CPP，每个标准氨基酸循环，各 4 次）",
        "第 38–41 行等：10.0 v，目标 8.73 kg/kg（419.04 kg），NOR 8.381–9.079 kg/kg（402.3–435.8 kg），PAR NLT 7.93",
        "表 6：Add Isopropyl Acetate (10.0 V, 8.73 kg/kg)\n表 7 / 附录 4：NOR 8.38 - 9.08",
        "3.5.5 等：Add 445.92 kg (Acceptable weighing range：432.54 ~ 459.30 kg) (9.29 kg/kg, 10.64 V)\n★NOR 【432.54 kg~459.30 kg】\n同一步下一行又写 Add 419.04 kg IPAc (10 V)，称量范围有的写 406.47~431.61 kg，有的写 440.34~467.58 kg",
        "三份文件给了两套洗涤量，MOR 自己也有两套。MOR 的 NOR 下限 432.54 kg 已接近参数表 NOR 上限 435.8 kg，两套范围基本不重叠。“419.04 kg”配“440.34~467.58 kg”这一句连自己的称量范围都不包含。按现在的 MOR 执行，批记录数会落在 PVP 表 7 的 NOR 之外。",
        "先定一个数（见“需要你确认”第 1 条），再同步三份文件。若维持 10.0 V / 8.73 kg/kg：MOR 3.5.5、3.8.5、3.11.5、3.16.8、3.19.5、3.22.5、3.25.5、3.28.5、3.31.5、3.34.5、3.37.5、3.40.5、3.43.5 改为 419.04 kg，NOR 改为 402.3~435.8 kg，并删除 445.92 kg 那一句。",
    ),
    (
        "偶联温度 NOR（CPP）",
        "标准 AA 20–27 °C（如第 158 行）；Phe23 19–28 °C（第 58 行）；NMeAsn22、αMePhe15 19–27 °C（第 84、287 行）；Asn21 19–28 °C（第 111 行）；Gly24 20–30 °C（第 21 行）",
        "表 6：Heat the mixture to 20 - 27 °C 等\n表 7：Coupling temperature 四行，NOR 同参数表",
        "所有偶联模块：The operator shall confirm that the reaction mixture temperature is within the Normal Operating Range (NOR) :20~30℃\n正文：The mixture is allowed to react at【25±5 ℃ (target: 25 ℃)】",
        "MOR 仍是旧的 25 ± 5 °C。除 Gly24 外，MOR 允许的 27–30 °C（标准 AA）或 28–30 °C（Phe23、Asn21）都超出 PVP 的 NOR，标准 AA 的 28.4–30 °C 还超出 PAR。",
        "MOR 3.7、3.10、3.13、3.18、3.21、3.24、3.27、3.30、3.33、3.36、3.39、3.42、3.47 的偶联温度和 ★NOR 句按参数表逐个改：标准 AA 20–27 °C，Phe23 和 Asn21 19–28 °C，NMeAsn22 和 αMePhe15 19–27 °C。超温预警值也要跟着改。",
    ),
    (
        "溶解夹套温度（CPP）",
        "第 51、75、102、151、278 行等：目标 20 °C，NOR 15–25 °C",
        "表 6：Adjust the jacket temperature to 20 ± 5 °C … Stir at least 15 minutes at 20 ± 5 °C\n表 7：Dissolution jacket temperature 四行",
        "3.6.3 / 3.9.3 / 3.17.3 等：Maintain stirring for NLT 15 min（溶解时没有温度要求）\n之后 3.9.4 等才 Adjust the activation tank … mixture to 【20±5 ℃】；Phe23 在 3.6.4 直接降到 5 ± 5 °C",
        "PVP 把溶解夹套温度列为 CPP，MOR 在溶解阶段既不控夹套温度，也没有记录点和 ★ 标识。PPQ 时附录 4 的这几行将无数据可填。",
        "MOR 每个活化模块在加 DMF 后、加氨基酸前加一步“夹套调至 20 ± 5 °C（目标 20 °C）”，溶解期间记录夹套温度最高/最低值，并标 ★。Phe23 的溶解也要先在 20 ± 5 °C 完成，再降到 5 ± 5 °C。",
    ),
    (
        "标准 De-Fmoc 时间",
        "第 37、66、93 行等：目标 120 min，NOR 110–130 min，PAR 68–180 min",
        "表 6：stir for 110 -130 min (target: 2 h)（12 处）",
        "3.5.3、3.8.3、3.11.3、3.19.3、3.22.3、3.25.3、3.28.3、3.31.3、3.34.3、3.37.3、3.40.3、3.43.3：After stirring for 2 h±15 min (target: 2 h)",
        "MOR 是 105–135 min，比 NOR 宽，两端各超出 5 min。",
        "MOR 12 处改为 110–130 min（目标 120 min）。",
    ),
    (
        "偶联总时间上限",
        "标准 AA PAR 240–600 min，NOR 4.5–9.5 h；Glu14 PAR 480–1083 min，NOR 8.5–17.5 h；Asn21 总时间 PAR 780–1680 min（第 118 行）；Gly24 PAR 180–481 min；Phe23 360–841；NMeAsn22、αMePhe15 360–840",
        "表 6：4.5 h ≤ Coupling time ≤ 9.5 h；Glu14 8.5 h ≤ … ≤ 17.5 h；Asn21 13.5 h ≤ total time ≤ 27.5 h",
        "标准 AA、Glu14、Asn21 模块都写：The total reaction … coupling time：NMT 840 min\nGly24：★The total reaction Fmoc-Gly-OH coupling time：NLT 600min（中文写“不超过600min”）\n不合格处置注：Continue hourly sampling up to the NMT 841 min hold limit",
        "标准 AA 的 840 min 超出 PAR 600 min；Glu14 的 840 min 低于 PVP NOR 上限 17.5 h（1050 min）；Asn21 的 840 min（14 h）几乎卡在 PVP 下限 13.5 h，PVP 允许到 27.5 h。Gly24 的 600 min 超出 PAR 481 min，英文 NLT 与中文“不超过”意思相反。不合格处置注统一写 841 min，也是照抄 Phe23。",
        "MOR 按氨基酸写各自上限：标准 AA NMT 600 min，Glu14 NMT 1083 min，Asn21 NMT 1680 min（或按 PVP 27.5 h），Gly24 NMT 481 min，Phe23 841 min，NMeAsn22/αMePhe15 840 min。不合格处置注引用本氨基酸的上限。",
    ),
    (
        "Gly24 偶联时间",
        "第 22 行：NOR 3.5 h ± 15 min",
        "表 6：react for 3.5 h ± 15 min (target: 3.5 h)",
        "3.3.5：The mixture is allowed to react fo 【3 h±15 min (target: 3 h)】",
        "MOR 仍是 3 h。",
        "MOR 3.3.5 改为 3.5 h ± 15 min（目标 3.5 h），顺手改 “fo” 为 “for”。",
    ),
    (
        "Pre-Ser20 第二次脱保护的哌啶/DMF 量",
        "第 137 行：10.39 v，9.68 kg/kg（464.6 kg），NOR 9.196–10.16",
        "表 6：Add 20% piperidine in DMF (10.39 V, 9.68 kg)",
        "正文 3.16.4：Add 464.64 kg【20% (V/V) piperidine/DMF(9.68 kg/kg, 10.40 V)】\n原料信息表 3.16.4：仍是 8 V、7.448 kg/kg、357.50 kg",
        "MOR 正文已改，原料信息表没改。体积三处写法不同（10.39 / 10.40 / final1 版的 10.4）。PVP 单位少了 /kg。",
        "统一为 10.4 V、9.68 kg/kg（464.64 kg）。改 MOR 原料信息表 3.16.4；PVP 表 6 改单位。",
    ),
    (
        "Ser20 后再溶胀（Pre-Ser20 脱保护后）",
        "第 144–145 行：DMF charge (13.0 v)，目标 NLT 12.32",
        "表 6：Add DMF (13.0 V, 12.32 kg/kg)，2 次",
        "3.16.9、3.16.11：Add 546.05 kg (Acceptable weighing range: 529.67 ~ 562.43 kg) (12 V)；原料信息表同",
        "MOR 仍是 12 V。",
        "MOR 3.16.9、3.16.11 及原料信息表改为 13 V、591.36 kg（12.32 kg/kg × 48 kg）。",
    ),
    (
        "Val17 前、Leu12 前再溶胀",
        "第 222–223 行（Arg18 块末）、第 352–353 行（Ala13 块末）：目标 NLT 9.48 kg/kg（10 V）",
        "表 6：Add DMF (10.0 V, 9.48 kg/kg)，2 次（两处）",
        "3.25.6、3.25.8、3.40.6、3.40.8：Add 546.05 kg … (12 V)；原料信息表同",
        "MOR 仍是 12 V。Johan（礼来）在 PVP 批注里也确认 PAR 为 2 × NLT 9.48 kg/kg。",
        "MOR 四处及原料信息表改为 10 V、455.04 kg。",
    ),
    (
        "Orn11 偶联后的 IPAc、甲苯、MTBE 量（CPP）",
        "第 401–411 行：IPAc 11.0 v / 9.60；甲苯 9.59；MTBE 8.15（均为 11 V）",
        "表 6：IPAC (11.0 V, 9.60 kg/kg) ×5；Toluene (11.0 V, 9.59 kg/kg)；MTBE (11.0 V, 8.15 kg/kg)\n表 7 NOR：9.22–9.98 / 9.21–9.97 / 7.82–8.48",
        "3.47.6：Add 419.04 kg … (10 V) isopropyl acetate\n3.47.7、3.47.9、3.48.1：418.56 kg (10 V) toluene\n3.47.8、3.47.10、3.47.11、3.48.2、3.48.3：355.68 kg (10 V) MTBE；原料信息表同",
        "MOR 仍是 10 V，换算后（8.73 / 8.72 / 7.41 kg/kg）全部低于 PVP 表 7 的 NOR 下限。",
        "MOR 改为 IPAc 460.80 kg、甲苯 460.32 kg、MTBE 391.20 kg（11 V），NOR 按参数表 kg 列。",
    ),
    (
        "Orn11 后甲苯 / MTBE 洗涤次数",
        "第 406–411 行：甲苯、MTBE、MTBE、甲苯、MTBE、MTBE（2 次甲苯、4 次 MTBE）",
        "表 6 两行：Toluene、MTBE、MTBE；再 Toluene、MTBE、MTBE + 氯冉酸检测",
        "3.47.7–3.47.11：甲苯、MTBE、甲苯、MTBE、MTBE，然后 IPC；3.48（甲苯、MTBE、MTBE）只在 IPC 不合格时调用",
        "MOR 常规路径只有 3 次 MTBE，比 PVP 和参数表少 1 次。",
        "确认以哪份为准（见“需要你确认”第 3 条）。若以 PVP 为准，在 MOR 3.47.8 后补一次 MTBE。",
    ),
    (
        "干燥",
        "第 412 行：夹套 NMT 30 °C",
        "表 6：Sweep with nitrogen, and vacuum underneath at T ≤ 30 oC by jacket … until LOD ≤ 9% (w/w)\n表 9：LOD ≤ 9 %",
        "3.49.1：nitrogen flow（11~15 m3/h target: 14 m3/h）… sample the solid for LOD analysis every 6±4 h until LOD ≤10.0%\nMOR IPC 表：LOD ≤9%",
        "MOR 正文 LOD 仍是 10.0%，与 MOR 自己的 IPC 表和 PVP 不一致。PVP 没写氮气流量和复测间隔，MOR 有。",
        "MOR 3.49.1 改为 LOD ≤ 9% (w/w)。PVP 表 6 干燥步补氮气流量 11–15 m³/h（目标 14 m³/h）和每 6 ± 4 h 复测。",
    ),
    (
        "Gly24 DIPEA 当量（CPP）",
        "第 16 行：NOR 5.70–6.30 eq",
        "表 7：5.70 - 6.30",
        "3.2：★The operator shall confirm that DIPEA quantity is within the Normal Operating Range (NOR) 【5.82~6.18 eq.】",
        "MOR 把 ±3% 称量容差当成了 NOR。PVP 表 7 表头写的是“NOR (MOR)”，两边数字应相同。",
        "二选一：MOR 改为 5.70~6.30 eq.；或 PVP 表头改为 NOR，不再写 (MOR)。见“需要你确认”第 2 条。",
    ),
    (
        "MeOH 封端后 DMF 洗涤量（CPP）",
        "第 31–33 行：9.0 v，8.53 kg/kg，NOR 8.10–8.96 kg/kg（388.8–430.1 kg）",
        "表 7：Resin loading – Capping DMF wash quantity，NOR 8.10 – 8.96",
        "3.4：★… DMF quantity is within the Normal Operating Range (NOR) 【397.25~421.83 kg】",
        "同上，MOR 用的是 ±3% 称量范围（8.28–8.79 kg/kg）。",
        "同上一条的处理方式。",
    ),
    (
        "Pre-Ser20 两次脱保护时间（CPP）",
        "第 136、139 行：处理 1 NOR 11–36 min；处理 2 NOR 20–51 min",
        "表 6：After stirring for 11 - 36 min … After stirring for 20 - 51 min；表 7 同",
        "3.16.3：After stirring for【15 ± 2 min】；3.16.6：After stirring for【30 ± 2 min】",
        "MOR 比 NOR 窄，执行上不越界，但 PVP 表 7 标为“NOR (MOR)”的数与 MOR 不同。",
        "同 DIPEA 一条的处理方式。",
    ),
    (
        "预处理 DMF 量",
        "第 3–4 行：5.21 kg/kg（250.08 kg），NOR ±5%（237.58–262.58 kg）",
        "表 6：Add DMF (5.50 V, 5.21 kg/kg)",
        "3.0.1 / 3.0 (2)：Add 200.00 - 300.00 kg DMF",
        "MOR 范围比参数表 NOR 宽。该步是测含水前的冲洗，不是 CPP，执行风险低，但三份写法不同。",
        "二选一：MOR 改为 250.08 kg（237.58–262.58 kg）；或参数表注明该步按 200–300 kg 执行，不设 NOR。",
    ),
    (
        "Gly24 溶解搅拌时间",
        "参数表没有这一行",
        "表 6：Stir for at least 15 minutes, Confirm no visible solids and later use.",
        "3.2.3：Maintain stirring for 1~45 min (target: 5 min). Confirm no visible solids and later use.",
        "MOR 最短 1 min，PVP 要求至少 15 min。",
        "以 PFD 为准统一，并在参数表补一行。",
    ),
    (
        "Phe23 活化时间",
        "第 54 行：目标 20 min，NOR 5–35 min",
        "表 6：Activate for 20 ± 5 min at 5 ± 5 °C",
        "3.6.9：Maintain stirring for 20±5 min (target: 20 min)",
        "参数表 NOR 与 PVP、MOR 不同，其他氨基酸的活化时间 NOR 都是 15–25 min。",
        "参数表第 54 行改为 15–25 min，或确认 5–35 min 是否为有意放宽。",
    ),
    (
        "偶联计时起点",
        "Phe23 NOR 19–28 °C；αMePhe15 19–27 °C",
        "表 6 Phe23：Start timing when temperature reaches 20 °C",
        "3.7.6、3.33.5、3.36.5、3.47.5：Start timing when temperature reaches 20 ℃",
        "NOR 下限是 19 °C，若批次控制在 19 °C，就永远等不到“20 °C 开始计时”。",
        "改为“温度进入 NOR 时开始计时”，或把计时温度写成各自 NOR 下限。",
    ),
    (
        "Asn21 偶联后的氯冉酸 IPC",
        "参数表没有",
        "表 9 和附录 3 没有",
        "MOR IPC 表没有",
        "Johan（礼来）10/25、10/27 两次未关闭：PAR 报告把补加 Oxyma/DIC 之后的氯冉酸检测列为 IPC，之前那次可作为 FIO。",
        "三份同时补：PVP 表 9、附录 3；MOR 3.13 及 IPC 表；参数表加 IPC 行。方法 AM-140020-030。",
    ),
    (
        "1.0 V / 2.0 V DMF 冲洗的 kg/kg 折算",
        "标准 AA：DMF charge (1.0 v, 0.94 kg/kg)；NMeAsn22、αMePhe15、Asn21：0.95；Phe23 2.0 v 写 1.89，NMeAsn22 2.0 v 写 1.90",
        "表 6 跟参数表：标准 AA 1.0 V，0.94 kg/kg；Phe23 2.0 V，1.89 kg/kg",
        "原料信息表：1 V = 0.948 kg/kg（45.50 kg，±2 kg）；2 V = 1.896 kg/kg（91.01 kg）",
        "同一体积写成两个数。0.948 四舍五入是 0.95，1.896 是 1.90。kg 数仍在 MOR 称量范围内，不影响执行，但会被客户问。",
        "统一为 0.95 和 1.90（见“需要你确认”第 4 条）。",
    ),
    (
        "DIC 加料温度、哌啶溶液温度是否为 CPP",
        "有 Temperature of charging DIC、Temperature of charging 20% piperidine/DMF 行",
        "表 7 只有 Activation temperature 和 Deprotection temperature，没有这两项",
        "多处：★Note:The DIC addition temperature is a CPP.；★Confirm and record that the 20%Piperidine solution temperature is within the NOR range 20 oC ~30 oC",
        "MOR 标为 CPP，PVP 未列入，附录 4 也不会汇总。",
        "确认这两项是否为 CPP（见“需要你确认”第 6 条），三份文件保持一致。",
    ),
]

PVP_ROWS = [
    ("第 1 节首段", "Sciin-procesence", "公司名被打坏，应为 Asymchem Life Science (Tianjin) Co., Ltd.。"),
    ("第 2 节范围", "against CPPs and CQAs", "第 9 节写明 Step 1 不制定 CQA。改为 against CPPs and IPCs。"),
    ("第 3 节策略", "release specifications", "ASYM-141366 没有放行标准。改为 meets its IPC acceptance criteria and is suitable for forward processing。"),
    ("第 5 节职责最后一句", "The personals who are executing", "personals 改为 personnel；“shall be attached report”改为“shall be attached to the report”。"),
    ("第 6 节 Table 4 下的注", "calculated based on the input range and the expected yield range", "产量按 MOR 公式计算：理论产量 = 树脂量 + 树脂量 × 取代度 × 2787.03 ÷ 1000，再乘 60%–130%。MOR 写明 2-CTC 取代度标准为 1.0–1.2 mmol/g，据此产量应为 108.84–271.66 kg；105.63–278.62 kg 对应取代度约 0.96–1.24，超出标准。请按 MOR 一起改，并在注里写明取代度这一计算基准。句末还有一个多余的“.”。"),
    ("第 6 节批次策略段", "Fragment 2 PPQ batch strategy", "Fragment 2 应为 Fragment L2；commercialscale 应为 commercial-scale；“three … batches is based” 主谓不一致。同段用 master batch records (MBRs)，全文其他地方用 MOR，请统一。"),
    ("第 7 节首段", "fit for process requirement", "“<RA-140020030101.01> Equipment selection” 之间缺句号，同一个 RA 号在一句话里引用两次，删一处即可。"),
    ("第 8 节首段", "RPT- 1194035", "引用残缺：“<Refer to RPT- 1194035 (Lilly document number) -and form”。改为 (Lilly PAR report RPT-1194035) and forms the basis。"),
    ("表 7 表下注", "Non-bold font format", "表头 PAR* 的星号指向这条注，但注讲的是“非粗体不是 CPP”，而表中所有行都标为 CPP。注与星号对不上，请改写或删去。"),
    ("表 7 和附录 4 的 Target 列（洗涤量各行）", "Post Deprotection IPAc wash quantity", "Target 栏写的是 PAR 下限，不是目标：IPAc wash quantity 写 NLT 8.73，工艺实际加 9.60 kg/kg（NOR 9.22–9.98）；Toluene、MTBE 同理（9.59、8.15）；Capping DMF wash 写 NLT 7.58，实际 8.53；Post Asn21 capping DMF wash 实际 7.58。脱保护后 IPAc 的 Target 写 NLT 8.73，NOR 却从 8.38 开始，低于“NLT”目标。建议 Target 栏写实际加入量（8.73、9.60、9.59、8.15、8.53、7.58），NLT 只放在 PAR 栏。"),
    ("表 6 Pre-Ser20 第二次 De-Fmoc", "10.39 V, 9.68 kg", "final1 版已改为 10.4 V，这一版又回到 10.39 V；单位仍少 /kg。改为 10.4 V, 9.68 kg/kg，与 MOR 正文的 10.40 V、464.64 kg 一致。Johan 10/27 的意见是只保留一位小数。"),
    ("表 9 第 2 行和附录 3 第 16–17 行", "Ser20~aMePhe15", "final1 版已改成 αMePhe，这一版又回到 aMePhe（表 9 两处，附录 3 两处）。改回 αMePhe15。"),
    ("表 9 IPC 计划", "The testing plan of IPCs", "缺 Asn21 补加后的氯冉酸 IPC，Johan（礼来）未关闭。按 Leu12、Orn11 的格式补一行，并在附录 3 补结果行。"),
    ("表 6 Phe23 偶联", "Start timing when temperature reaches 20 °C", "NOR 已是 19–28 °C，“到 20 °C 才计时”与 NOR 下限冲突。改为进入 NOR 即开始计时。同句 “(target: 25 °C).. ” 双句号。"),
    ("表 6 Asn21 偶联", "13.5 h ≤ total time ≤ 27.5 h", "两段固定时间（8 h ± 15 min + 5 h ± 15 min，合计约 12.5–13.5 h）与“总时间可到 27.5 h”说的不是同一件事，且 MOR 写 NMT 840 min。请写清总时间的定义和到点后的动作。Johan 关于这一步 DMF 体积（7.0 V / 8.0 V 与 PAR 6.65–8.55 kg/kg）的提问也还没在正文里回答。"),
    ("表 6 两处再溶胀", "Stirring and turn off the nitrogen", "缺谓语和搅拌时间，改为与其他再溶胀行一致：Bubble with nitrogen, after stirring, turn off the nitrogen。"),
    ("表 6 标准氨基酸活化（8 处）", "target: 20 °C, activate", "括号没关：“(target: 20 °C, activate for 20±5 min”。另有 “TJ4S-1213-Sy15-R01),” 多一个右括号，“15minutes” 缺空格（5 处），“1.0 V，0.94” 用了全角逗号。"),
    ("表 6 Leu12 后 10 次 IPAc", "lighter or consistent with the reference solution and then", "句子停在 and then。补完：…and then proceed to re-swelling。"),
    ("表 6 Orn11 活化", "Orn(Boc)-OH (2.00 eq.) coupling", "投料句末多了 coupling；物料名 “Fmoc-L- Orn(Boc)-OH” 中间多空格；行标题 “coupling activation and transfer” 多了 coupling。"),
    ("表 6 Orn11 偶联", "react at 20 - 27 (target", "温度少了单位，改为 20 - 27 °C。"),
    ("表 6 Orn11 后洗涤", "repeat Step 93", "表 6 没有 93 号步骤。改为“repeat this toluene/MTBE wash sequence”。句末 “If the color is lighter or consistent with the reference solution,” 以逗号结束，补 “proceed to drying.”。"),
    ("表 6 干燥", "Sweep with nitrogen, and vacuum underneath", "MOR 有氮气流量 11–15 m³/h（目标 14 m³/h）和每 6 ± 4 h 复测，PVP 没写。补上，避免 PVP 与 MOR 不一致。"),
    ("表 6 配封端液", "activation tank TJ4S-1213- R22", "R22 是搪玻璃釜，不是活化釜；设备号 “1213- R22” 多空格（4 处）。"),
    ("表 5 设备", "-14~57", "Sy15 同一标签下写了 -14~57 和 -4~53 两组温度；前者是夹套、后者是批内，请分别标注 jacket / batch。同表 AT02 搅拌 “0-119.84 rpm” 像计算残留，ESA 写的是确认最大 125 rpm。"),
    ("表 5 设备", "Glass-line Reactor", "Glass-line 应为 Glass-lined；同表 “Mass flow mete” 少 r（2 处）。"),
    ("表 9 第 5 行和附录 3", "Orn11 post coupling wash piperidine residue.", "括号没有关上，两处都要补 “)”。"),
    ("第 11 节 LOD 注", "simple technique that require", "语法错，且与“所有 IPC 方法已验证、引用 AM-140020-016-VP-01”矛盾。改为引用方法确认结论。"),
    ("第 11 节方法表", "ethod Number", "表头少 M；数据行为空。按正文已列出的 L2 方法和验证号填表。"),
    ("表 8 Sunresin 地址", "No.135 Jinye Road, High-Tech Zone)", "Johan 10/27 未关闭：地址被改短，缺城市、省和国家，与表 2、VMP 不一致。补全。"),
    ("第 10 节", "registration starting materials", "改为 regulatory starting materials，与第 1、2 节一致。"),
    ("表 2 及供应商地址（15 处）", "Fu Mongolia county", "会被读成“蒙古”。改为 Fuxin Mongol Autonomous County (Fumeng County)。"),
    ("第 14 节包装", "after release Packaging", "缺句号；本物料不放行，只“released for further processing”。下一段干燥剂规则混用全角标点，且 “NLT desiccant ≥” 重复；外包装 HPDE 应为 HDPE（2 处）。"),
    ("第 15 节", "PSP-140020-2026-003.01>. .", "多一个句号。"),
    ("第 17 节清洁验证", "CVP-1400200301,", "缺版本号 .01；段末 “cleaning validation report” 缺句号。"),
    ("第 18 节", "Quality argeement", "argeement 拼写错误。"),
    ("第 19 节接受标准", "Validation Master Plan (VP-14002001.01)", "同一编号在第 1 节、附录 2 叫 Process Validation Plan，这里叫 Validation Master Plan，统一为前者。"),
    ("第 19 节接受标准", "approved Master Operation Record (MOR-TJ414002001.04)", "接受标准要求按 MOR .04 生产，但 MOR_10282026（即 .04）还有多处参数与本方案表 6、表 7 不一致（见“三方参数核对”）。MOR 改完前，这一条无法同时满足。"),
    ("第 9 节", "meeting their specifications", "主语是 Step 2 product，应为 meeting its specification。"),
    ("表 6 预处理", "proceed to next operation if Water content", "合格和不合格两句粘在一起，改为两句，并写明重复的是 DMF 冲洗和取样。"),
]

EXCEL_ROWS = [
    ("第 144–145 行（De-Fmoc 块）", "DMF charge (13.0 v, kg/kg)", "目标 NLT 12.32，NOR 却是 10.81–11.95（12 V 的 NOR），目标在 NOR 之外。13 V 的 ±5% 应为 11.70–12.94 kg/kg（561.6–620.9 kg）。"),
    ("第 222–223 行（Arg18 块）、第 352–353 行（Ala13 块）", "DMF charge (12.0 v, kg/kg)", "标签写 12.0 v，目标是 NLT 9.48（10 V）。两处 NOR 用了不同容差：9.006–9.954（±5%）和 9.101–9.859（±4%）。标签改 10.0 v，容差统一。"),
    ("第 42–43 行（Capping 块，Phe23 前再溶胀）", "DMF charge (12.0 v, kg/kg)", "NOR 10.92–11.84（±4%），其他 12 V 再溶胀行都是 10.81–11.95（±5%）。统一容差。"),
    ("第 38 行起所有脱保护后 IPAc 行", "8.73（NLT8.73 kg）", "目标写“8.73（NLT 8.73）”，NOR 下限 8.381 低于 NLT。若 8.73 是目标，去掉 NLT；若 8.73 是下限，NOR 和目标都要上移（MOR 的 9.29 kg/kg 可能就是这样来的）。Orn11 后的 “9.60 NLT8.73”、甲苯、MTBE 同理，目标和 PAR 下限挤在一格。"),
    ("第 282 行（α-Me-Phe15）", "DMF charge for  Fmoc-L-α-Me-Phe-OH dissolution(5.0 v, kg/kg)", "目标 5.69 是 6.0 V（4 V + 2 V 冲洗）。标签应为 6.0 v，参照第 79 行 NMeAsn22。"),
    ("第 283 行（α-Me-Phe15）", "Temperature of charging DIC", "PAR 写 12.6~26.5，同块第 284 行活化温度和 NMeAsn22 第 80 行都是 12.6~26.9，像是抄了标准 AA。"),
    ("第 288 行（α-Me-Phe15）", "Coupling Time – NMeAsn22", "标签抄错，应为 Coupling Time – αMePhe15。"),
    ("第 112、118 行（Asn21）", "Coupling time 1", "两行同名。第 118 行是总时间（PAR 780–1680 min），应改名 Total coupling time，并补 NOR（PVP 写 13.5–27.5 h）。"),
    ("第 114 行（Asn21）", "DMF charge (1.0 v, 0.95 kg/kg)", "标签写 0.95，目标填 0.94。"),
    ("第 54 行（Phe23）", "Activation time", "NOR 5–35 min，与 PVP、MOR 的 20 ± 5 min 不同。"),
    ("第 28 行（Capping）", "Temperature of charging DIPEA", "PAR 两格为空。请补，或注明 PAR 报告未研究。"),
    ("第 73 行", "DMF charge for Fmoc-L-N-Me-Asn(Trt)-OH dissolution (4.0 v, kg/kg)", "这是 NMeAsn22 的溶解 DMF，AA 列却归在 Phe23 下，按块筛选时会漏。"),
    ("K–O 列（换算 kg 的列）", "#VALUE!", "目标或 PAR 写成“NLT …”文字的行，kg 列显示 #VALUE!，无法直接和 MOR 的 kg 比对。建议把 NLT 拆到单独一列。"),
]

MOR_ROWS = [
    ("原料信息表 3.16.4", "3.16.4", "仍是 8 V、7.448 kg/kg、357.50 kg；正文 3.16.4 已是 10.40 V、464.64 kg。"),
    ("3.5.5 等 13 处脱保护后 IPAc", "Add 445.92 kg", "同一步写了 445.92 kg（9.29 kg/kg、10.64 V）和 419.04 kg（10 V）两套；“419.04 kg”配的称量范围有 440.34~467.58 kg（不含 419.04）和 406.47~431.61 kg 两种。3.11.5 还写成 “432.54 ~ 4459.30 kg”。"),
    ("原料信息表 3.37.1、3.37.5、3.38.1", "3.37.1", "体积比（8、10、5）错填到了“摩尔比”列。"),
    ("3.3 Gly24 偶联", "coupling time：NLT 600min", "英文 NLT（不少于）与中文“不超过”相反；且 600 min 超出 PAR 481 min。"),
    ("3.24、3.32、3.36、3.39 的偶联总时间句", "The total reaction Fmoc-L-Leu-OH coupling time", "αMePhe15、Glu14、Ala13 模块写成 Fmoc-L-Leu-OH；Arg18 模块写成 Fmoc-L-Ser(tBu)-OH。"),
    ("3.49.1 干燥", "until LOD ≤10.0%", "与本 MOR IPC 表的 ≤9% 不一致。"),
    ("分析方法 / IPC 表", "the resin is yellow to orange", "茚三酮（AM-140020-029）的判据写成了溴酚蓝的“树脂黄至橙”；茚三酮应为“溶液黄至浅粉”。"),
    ("3.0 预处理", "搅拌5 min±5 min", "英文 5 min ± 2 min，中文 5 min ± 5 min。"),
    ("3.17.8、3.26.8 活化", "Maintain for NLT 15 min", "英文 NLT 15 min（目标 20 min），中文 20 ± 5 min，参数表 NOR 15–25 min。"),
    ("3.17.1", "Start the agitator to 95±5 rpm", "其余都是 95 ± 10 rpm。"),
    ("3.2.5", "TJ4S-1213-Sy15-RT01", "设备号多了 T。"),
    ("多处笔误", "Shirnk with MTBE", "Shirnk（2 处）、react fo、soulution、Ddeprotection、Adust、activation tan、Pyriine（代码表）。3.40.3 写 “Deprotection of Fmoc-L-Leu-OH activation”；3.45 只有再溶胀，标题却是 “(14) Deprotection”；3.43.5 编号用了两次。"),
    ("原料信息表 3.2.4 与 3.4.3", "129.24", "同一物料 DIPEA 的分子量写 129.25 和 129.24。"),
]

MAPPING = [
    ("第 3–4 行 洗釜测含水", "3.0.1 和 3.0 (2)", "第 2 行对应不合格时的重复，常规只执行一次", "中"),
    ("第 5 行 树脂", "3.1.1", "", "高"),
    ("第 6–8 行 Gly 溶解 DMF、Gly、冲洗", "3.2.1、3.2.2、3.2.3", "", "高"),
    ("第 9 行 dissolution (5.25 v)", "3.2.1 + 3.2.3 之和", "MOR 无单独一行", "中"),
    ("第 10–12 行 溶胀 7 V、1 V、合计 8 V", "3.1.2、3.1.4、二者之和", "", "高"),
    ("第 13–15 行 溶胀温度、时间、5 V 洗涤", "3.1.5–3.1.7", "", "高"),
    ("第 16–20 行 DIPEA、活化", "3.2.4、3.2.5、3.3.2", "第 20 行 1.0 v 判断为 3.3.2 罐壁冲洗", "中"),
    ("第 21–24 行 Gly 偶联", "3.3.4、3.3.5、3.3.6", "第 23 行 6.25 v = 3.2.1 + 3.2.3 + 3.3.2", "中"),
    ("第 25–30 行 MeOH 封端", "3.4.1–3.4.5", "", "高"),
    ("第 31–33 行 DMF post-coupling wash (9.0 v)", "3.4.6（×3）", "名称写 post-coupling，实为 MeOH 封端后洗涤", "中"),
    ("第 34–43 行（AA 列为 Capping）", "3.5（Phe23 前脱保护）3.5.1–3.5.9", "", "高"),
    ("第 44–59 行 Phe23", "3.6、3.7", "第 55 行 2.0 v = 3.7.1；第 56 行 1.0 v 可能是 3.6.7 或 3.7.3", "低（第 56 行）"),
    ("第 60–62 行 DMF activation flask charge", "3.7.7（×3）", "", "高"),
    ("第 63–72 行", "3.8（NMeAsn22 前脱保护）", "", "高"),
    ("第 73–89 行 NMeAsn22", "3.9、3.10", "第 73 行 AA 列标为 Phe23，实为 3.9.1", "高"),
    ("第 90–99 行", "3.11（Asn21 前脱保护）", "", "高"),
    ("第 100–122 行 Asn21", "3.12、3.13", "第 113–115 行 = 3.13.6、3.13.5、3.13.7；第 120–122 行 = 3.13.8", "高"),
    ("第 123–130 行 Ac2O 封端", "3.14、3.15", "第 128 行 7 v = 3.14.1（6 V）+ 3.15.2（1 V）", "中"),
    ("第 131–133 行 Post Asn21 capping wash", "3.15.6（×3）", "", "高"),
    ("第 134–145 行 Pre-Ser20", "3.16.1、3.16.4、3.16.8、3.16.9/3.16.11", "", "高"),
    ("每个标准 AA 块（Ser20 起）", "3.x 活化 / 3.x+1 偶联 / 3.x+2 下一个脱保护", "块末的脱保护、IPAc、再溶胀属于下一个氨基酸的 3.x+2", "高"),
    ("Dissolution Jacket Temp 各行", "MOR 无对应控制点", "见三方核对第 3 条", "—"),
    ("第 222–223 行", "3.25.6、3.25.8（Val17 前）", "", "高"),
    ("第 352–353 行", "3.40.6、3.40.8（Leu12 前）", "", "高"),
    ("第 375–384 行 IPAc ×10", "3.43.5 + 3.44.1", "第 10 次起取样", "高"),
    ("第 385–386 行", "3.45.1、3.45.3（Orn11 前再溶胀）", "", "高"),
    ("第 401–405 行 IPAc ×5", "3.47.6（×5）", "", "高"),
    ("第 406–411 行 甲苯/MTBE", "前三行 = 3.47.7、3.47.8、(缺一次 MTBE)；后三行 = 3.47.9–3.47.11 还是 3.48.1–3.48.3？", "3.48 只在 IPC 不合格时调用", "低"),
    ("第 412 行 干燥温度", "3.49.1", "", "高"),
]

QUESTIONS = [
    "脱保护后 IPAc 最终按哪一套：参数表和 PVP 的 10.0 V / 8.73 kg/kg（NOR 8.38–9.08），还是 MOR 的 10.64 V / 9.29 kg/kg（NOR 432.54–459.30 kg）？",
    "MOR 的 ★NOR 是否必须与 PVP 表 7 完全一致？现在 DIPEA（5.82–6.18 vs 5.70–6.30）、封端 DMF（397.25–421.83 kg vs 8.10–8.96 kg/kg）、Pre-Ser20 时间（15 ± 2 / 30 ± 2 min vs 11–36 / 20–51 min）都是 MOR 更窄。是改 MOR，还是把 PVP 表头的“(MOR)”去掉？",
    "Orn11 后的甲苯/MTBE：以 PVP 和参数表的“甲苯、MTBE、MTBE”×2 为准，还是以 MOR 3.47 的“甲苯、MTBE、甲苯、MTBE、MTBE”为准？",
    "1.0 V 和 2.0 V DMF 冲洗统一写 0.95 / 1.90 kg/kg，还是保留参数表的 0.94 / 1.89？",
    "预处理 DMF：保留 MOR 的 200–300 kg，还是按参数表 250.08 kg ± 5%？",
    "DIC 加料温度和 20% 哌啶溶液温度：MOR 标为 ★CPP，PVP 表 7 没有。这两项是不是 CPP？",
    "我把参数表第 31–33 行“DMF post-coupling wash (9.0 v)”对到 MOR 3.4.6（MeOH 封端后洗涤），对吗？",
    "参数表第 56 行 Phe23 的 “DMF charge (1.0 v)” 对应 MOR 3.6.7（冲洗加料口）还是 3.7.3（转移前冲洗罐壁）？两处都是 1 V。",
    "参数表第 406–411 行后三行（甲苯、MTBE、MTBE）对应 MOR 3.47.9–3.47.11，还是 IPC 不合格时的 3.48.1–3.48.3？",
]


def style_sheet(ws, headers, widths, rows):
    head_fill = PatternFill("solid", fgColor="1B4F72")
    head_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    side = Side(style="thin", color="D0D7DE")
    border = Border(left=side, right=side, top=side, bottom=side)
    for col, text in enumerate(headers, 1):
        cell = ws.cell(1, col, text)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, values in enumerate(rows, 2):
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.alignment = wrap
            cell.font = Font(name="Calibri", size=11)
            cell.border = border
            if r % 2 == 1:
                cell.fill = PatternFill("solid", fgColor="F7F9FB")
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def main():
    pvp_text = accepted_text(PVP)
    mor_text = accepted_text(MOR)
    missing = [q for _, q, _ in PVP_ROWS if q not in pvp_text]
    missing += [q for _, q, _ in MOR_ROWS if q not in mor_text]
    if missing:
        raise SystemExit(f"Search strings not found: {missing}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "三方参数核对"
    style_sheet(
        ws,
        ["序号", "参数", "参数表 10-28（行号 / 值）", "PVP 10-28（位置 / 写法）", "MOR 10-28（操作号 / 写法）", "判断", "建议改哪里"],
        [6, 22, 38, 40, 46, 46, 46],
        [(i, *row) for i, row in enumerate(THREE_WAY, 1)],
    )
    for name, headers, widths, rows in [
        ("PVP 硬伤", ["序号", "位置", "内容（Word 中搜索）", "批注内容（问题与改法）"], [6, 30, 40, 90], PVP_ROWS),
        ("参数表内部问题", ["序号", "位置", "内容（Excel 中搜索）", "问题与改法"], [6, 34, 44, 86], EXCEL_ROWS),
        ("MOR 内部问题", ["序号", "位置", "内容（Word 中搜索）", "问题与改法"], [6, 30, 36, 90], MOR_ROWS),
        ("参数表-MOR 对应关系", ["序号", "参数表行", "我判断的 MOR 操作号", "说明", "把握"], [6, 40, 46, 52, 10], MAPPING),
    ]:
        sheet = wb.create_sheet(name)
        style_sheet(sheet, headers, widths, [(i, *row) for i, row in enumerate(rows, 1)])
    qs = wb.create_sheet("需要你确认")
    style_sheet(qs, ["序号", "问题"], [6, 140], [(i, q) for i, q in enumerate(QUESTIONS, 1)])
    wb.save(OUT)
    print(f"wrote {OUT.name}: three-way {len(THREE_WAY)}, PVP {len(PVP_ROWS)}, Excel {len(EXCEL_ROWS)}, MOR {len(MOR_ROWS)}, mapping {len(MAPPING)}, questions {len(QUESTIONS)}")


if __name__ == "__main__":
    main()
