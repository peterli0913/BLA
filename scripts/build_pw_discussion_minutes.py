#!/usr/bin/env python3
"""Minutes of the 09/25/2026 Lilly call, purified-water discussion (second half only)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from build_lilly_call_script import (
    MUTED, NAVY, add_bullet, add_h1, add_h2, add_page_number, add_table, set_run_font,
)

OUT = "/workspace/meetings/2026/2026-09-25-礼来纯水讨论会议纪要.docx"

SUMMARY = [
    "纯水排查范围要扩大：礼来提出，除了直接加入工艺的纯化水，清洁用的纯化水（如 SPPS、裂解设备清洁）也要一并排查，并提出改造前的过渡管理方案。礼来提醒这可能更直接地影响早期步骤的时间。",
    "改造期间的供水：施工、钝化和验证合计约 7 周。礼来关心这段时间工艺罐怎么用水，要求我方说明过渡方案；Andrew 会通过邮件把问题写清楚发来。",
    "我方将持续更新 Excel 跟踪表，最晚下周初发给礼来。",
    "湿氮气（会后与 Mark 简短沟通）：WFI 无菌袋在哪里连接，礼来初步倾向在 C 级区，请我方先提出建议方案，下周讨论。",
]

NEEDS_PW = [
    ["清洁用纯化水",
     "我方表示已评估所有工艺用纯化水管线。礼来追问清洁是否也用纯化水（如 SPPS 设备）；如果分配系统有同类缺陷，清洁用水同样有微生物污染风险。需要先弄清缺陷范围，再定改造前怎么管理，可以考虑用后续的溶剂淋洗和干燥来降低风险并论证。之前的讨论集中在纯化（水直接加入工艺，风险最高），但要整体审视所有用水系统。",
     "确认 SPPS 和裂解设备清洁用纯化水。会把清洁用纯化水纳入排查，跟进后再讨论。"],
    ["3 周验证的内容",
     "改造后的 3 周验证是新取样点的验证，还是设备确认？",
     "属于设备确认，包括 2 周取样和 1 周微生物检测（转写不清，需核实）。"],
    ["改造期间是否停用",
     "施工 3 周 + 钝化 1 周 + 验证 3 周，约 7 周。这期间是否不能用水，需要 7 周停产窗口？",
     "纯化水主循环不受影响：改造的是从各分配点到设备的支管，改造后的水不回主水箱。"],
    ["工艺罐怎么供水",
     "礼来理解主循环不用停，但这些支管连的都是工艺罐。支管施工和再验证期间，怎么给工艺供水？管线上除了硬管还有流量计等部件，改用软管还需要考虑便携式流量计等。请说明 7 周期间的过渡供水方案。Andrew 会通过邮件把问题写清楚。",
     "可能用软管临时连接；收到邮件后回复过渡方案。"],
    ["最后一项",
     "我方原计划再讲两点：第一点与上面的再验证讨论相关，已覆盖；最后一点礼来认为比较明确，更新 SMP 即可，不需要再展开（具体内容需核实）。",
     "—"],
]

NEEDS_N2 = [
    ["WFI 无菌袋的连接位置",
     "上次讨论决定把湿氮气加湿器移到洁净区。我方已采购 WFI 无菌袋，问无菌袋需要放到 C 级区，还是可以放在 CNC 区、通过无菌接头接入 C 级区。Mark 初步倾向在 C 级区连接，但需要和更大范围的团队讨论，并了解我方有哪些限制。请我方从物流角度提出建议方案。",
     "下周提出问题和建议方案，与礼来讨论。"],
]

ACTIONS = [
    ["1", "排查清洁用纯化水（如 SPPS、裂解设备清洁）的管线是否有同类缺陷，明确范围，并提出改造前的过渡管理方案（如用后续溶剂淋洗、干燥降低风险的论证）",
     "未定", "可能更直接影响早期步骤的时间"],
    ["2", "说明纯化水支管施工、钝化和再验证（约 7 周）期间如何给工艺罐供水（软管临时连接、流量计等）",
     "收到礼来邮件后回复", "Andrew 会发邮件细化问题"],
    ["3", "确认 3 周验证的内容（2 周取样 + 1 周微生物检测），与第 2 项一并说明",
     "随第 2 项", "转写不清，需核实"],
    ["4", "更新 SMP（会上最后一项）", "未定", "具体内容需核实"],
    ["5", "持续更新 Excel 跟踪表并发给礼来", "最晚下周初", "与清洁分组会纪要第 7 项相同"],
    ["6", "湿氮气 WFI 无菌袋连接位置（C 级区，或 CNC 区经无菌接头接入）：提出建议方案，与礼来讨论",
     "下周", "Mark 初步倾向在 C 级区连接"],
]


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.7)
    sec.top_margin = sec.bottom_margin = Cm(1.5)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(hp.add_run("Project 139788 BLA Retrofit｜内部纪要 Internal"), size=8.5, color=MUTED)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run("礼来纯水讨论会议纪要"), size=20, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("2026-09-25｜依据会议后半段录音整理，前半段未录音"), size=10, color=MUTED)

    add_h1(doc, "一、会议要点")
    for s in SUMMARY:
        add_bullet(doc, s)

    add_h1(doc, "二、客户需求与关注点")
    add_h2(doc, "1. 纯水")
    add_table(doc, ["主题", "礼来的问题 / 需求", "凯莱英答复"], NEEDS_PW, [2.8, 8.8, 6.0], size=9.5, first_col_bold=True)
    add_h2(doc, "2. 湿氮气（会后简短沟通）")
    add_table(doc, ["主题", "礼来的问题 / 需求", "凯莱英答复"], NEEDS_N2, [2.8, 8.8, 6.0], size=9.5, first_col_bold=True)

    add_h1(doc, "三、后续 Action Items")
    add_table(doc, ["#", "事项", "期限", "备注"], ACTIONS, [1.0, 9.4, 3.2, 4.0], size=9.5)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
