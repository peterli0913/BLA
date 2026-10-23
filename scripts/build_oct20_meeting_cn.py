#!/usr/bin/env python3
"""Chinese translation and internal minutes for the 20 Oct 2026 GG917 JPT touchpoint."""

import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from build_lilly_call_script import (
    GOLD, MUTED, NAVY, add_body, add_bullet, add_h1, add_h2, add_page_number,
    add_table, set_run_font,
)

ROOT = Path("/workspace")
MEET = ROOT / "meetings" / "2026"

TRANS_NAME = "2026-10-20-GG917沟通会-中文译文.docx"
MIN_NAME = "2026-10-20-礼来客户会会议纪要.docx"


def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.7)
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)
    return doc


def title(doc, main, sub):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run(main), size=18, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run(sub), size=12, bold=True, color=GOLD)


def save_pair(doc, name):
    MEET.mkdir(parents=True, exist_ok=True)
    out = MEET / name
    root = ROOT / name
    doc.save(out)
    shutil.copy(out, root)
    print(out)
    print(root)


def build_translation():
    doc = new_doc()
    title(
        doc,
        "GG917 SPPS 与纯化启动准备",
        "每日 JPT 沟通会 · 中文译文 · 2026-10-20",
    )
    add_body(
        doc,
        "本文是 20_Oct.docx 的中文译文。原文说明：会议有转写，讨论摘要依据转写整理，对发言做了转述，不是逐字引用。人名、数字、批号和未展开的缩写按原文保留。",
    )

    add_h1(doc, "会议信息")
    for line in [
        "会议：GG917 SPPS 与纯化启动准备 — 每日 JPT 沟通会",
        "日期：2026 年 10 月 20 日",
        "预定时间：12:00–13:00",
        "组织者：Steven Thomas Ewing",
        "形式 / 地点：Microsoft Teams / #358-2-B",
        "调取时状态：会议进行中",
        "转写：有",
    ]:
        add_bullet(doc, line)

    add_h1(doc, "1. 概要")
    add_body(doc, "会议集中讨论纯化准备、DIPEA / 乙醛控制、溶剂试车准备、尚未完成的文件，以及 M+53 调查实验。")
    add_body(doc, "最重要的几点是：")
    for line in [
        "原则上已对齐：DIPEA 加严为 2–8°C 储存，并按桶检测和放行。长期稳定性，以及检测应距离使用多近，仍需要支持数据。",
        "若干纯化文件已退回凯莱英，并带有 comment。团队希望尽快解决，以支持溶剂试车。",
        "礼来实验室按拟定的缓冲液配制做了核对，稀释至 50 mM 后 pH 为 7.65，被认为接近预期。一套按质量配制的方法正在转给团队使用。",
        "流动相 / 进料 pH 和温度被识别为很可能的关键工艺参数（CPP）。最终参数清单要等 PAR 统计分析完成。",
        "对于把 pH 监测和过滤改造推迟到 demo 之后，会上没有达成最终一致。礼来担心：在 pH 属于 CPP 时，GMP 批次使用可能不可靠的 pH 测量。",
        "溶剂试车被视为确认实际温控表现的一次重要机会。如果设备改动推迟到其后，改造完成之后、工艺验证（PV）之前，可能还需要再做一次溶剂试车。",
        "M+53 调查仍在继续。当前树脂供应紧张。会上讨论留出一部分剩余树脂，并有可能把后续实验统一为 5 g，前提是确认混合和搅拌仍然有代表性。",
        "又一批 39mer-on-resin 的合成已经开始，凯莱英将分享时间线。",
    ]:
        add_bullet(doc, line)

    add_h1(doc, "2. DIPEA / 乙醛控制策略")
    add_body(
        doc,
        "Richard Morgan 说明：DIPEA 质量标准的变更控制讨论，引出了在 PV 之前增加操作控制的建议。其中包括把 DIPEA 转入 2–8°C 温控储存；在长期稳定性更清楚之前，考虑把检测时间明显移到更接近批次使用的时候。",
    )
    add_body(
        doc,
        "Runqian Yang 表示凯莱英同意 2–8°C 控制。她还说明了一种做法：按桶跟踪乙醛，分别检测、分别放行；达不到要求的物料不使用。",
    )
    add_body(
        doc,
        "Richard Morgan 指出，DIPEA 稳定性仍有不确定性。乙醛水平在不同供应商之间有差异；按会上的讨论，demo 所用批次的不同桶之间也有差异。因此团队需要能够支持 DIPEA 货架期和使用期限决定的数据。",
    )
    add_body(
        doc,
        "立场：原则上支持加严储存和按桶控制。长期稳定性及检测策略仍需要证据，并进一步对齐。",
        bold=True,
    )

    add_h1(doc, "3. 纯化文件与溶剂试车准备")
    add_body(
        doc,
        "Ezekiel Lee 过了纯化文件的时间安排，并强调必须定稿支持溶剂试车的文件。Simon Cushen（Network）说明：讨论中的四份文件已经退回凯莱英，comment 需要回复。凯莱英表示打算会后处理这些 comment。",
    )
    add_h2(doc, "缓冲液配制")
    add_body(
        doc,
        "Steven Benner 报告：礼来对缓冲液配制做了实验室试验，包括先配 pH 7.4 的 1 M Tris 浓储液，再稀释到 50 mM。稀释后测得 pH 7.65，被认为接近预期。已经做出一套完全按质量配制的方法，正通过 Andrew James Mouat 转交。",
    )
    add_body(doc, "Steven Benner 同时强调，稀释后仍要核实 pH，因为这个 pH 是关键工艺参数。")

    add_h1(doc, "4. 纯化过滤策略")
    add_body(
        doc,
        "Ezekiel Lee 说明了现有过滤布置：纯化水和乙腈进料使用 0.22 µm 过滤；最终流动相罐下游为 0.45 µm，他说可以改为 0.22 µm。他提出一个问题：较小浓缩液罐附近另要求增加的过滤器，是否必须在 demo 之前安装。",
    )
    add_body(
        doc,
        "Steven Benner 指出，这项具体的过滤要求来自 Andrew James Mouat，他不愿代为解释理由。Robert Moylan 建议直接向 Andrew James Mouat 确认这项要求，尤其因为这些过滤器对 demo 批可能是有用的微生物风险控制。",
    )
    add_body(
        doc,
        "未决：额外过滤器是在 demo 执行前必须安装，还是可以放到计划中的 pre-PPQ 改造窗口再装。",
        bold=True,
    )

    add_h1(doc, "5. 温度控制策略")
    add_body(
        doc,
        "纯化团队说明了整个工艺的温控结构。较大的浓盐罐有主动温控，因为 Tris-HCl 体系对温度敏感。其他流动相罐并不都有同等控制。层析系统进口周围设有加热。",
    )
    add_body(
        doc,
        "Steven Benner 问：进口换热器在要求的流量和体积下，是否已经证明能够有效控制。凯莱英说，新配置下这一点尚未确立，并建议用溶剂试车产生性能数据。",
    )
    add_body(
        doc,
        "Robert Moylan 还问到层析柱周围的温度监测，包括可能的夹套 / 出口监测，以及房间温度控制。讨论表明，进入的流动相有温度测量和加热，但层析柱 / 夹套的具体布置，以及房间控制的实际表现，仍需确认。",
    )
    add_body(
        doc,
        "已对齐的方向：用溶剂试车数据评估现有温控安排是否有效；同时凯莱英核对相关的历史或可比设备信息，以及设施控制。",
        bold=True,
    )

    add_h1(doc, "6. 在线 pH 监测及拟议的设备改造")
    add_body(
        doc,
        "一项关键的准备问题，是第一台在线混合器下游的 pH 监测。凯莱英说明：测量点是有的，但在当前高流速下，读数可能不够稳定。改这个布置据说可能大约增加 5 天，因此提出 demo 之后、pre-PPQ 之前再做改造。",
    )
    add_body(
        doc,
        "礼来对这个做法有明显担心。Robert Moylan 指出，demo 是 GMP 批次；如果 pH 是 CPP，接受不可靠的测量有问题。由于管线的实际布置，在该点方便地取离线样也比较困难。",
    )
    add_body(
        doc,
        "Richard Morgan 随后指出：把改造往后放，可能变成 demo 之后还要再做一次溶剂试车，以便在 PV 之前证明改造后的设备。也就是说，无论改造放在 demo 之前还是之后，进度都可能受影响。",
    )
    add_body(
        doc,
        "决定状态：未决。双方同意会下讨论，再决定这些改造必须在 demo 之前完成，还是可以后置并随后做确认。",
        bold=True,
    )

    add_h1(doc, "7. 关键工艺参数")
    add_body(doc, "Steven Benner 指出，目前预计为关键参数的有：")
    for line in ["温度", "进料 pH", "流动相 pH"]:
        add_bullet(doc, line)
    add_body(doc, "完整清单尚未最终确定，要等 PAR 研究的统计分析完成。")

    add_h1(doc, "8. PPQ / 准备文件")
    add_body(
        doc,
        "Runqian Yang 问了 PPQ 文件更整体的批准状态。Richard Morgan 表示：除 M+53 调查和 CAPA 所要求的更新之外，大部分文件大体可以接受；仍有具体文件带有 comment，或还没有审完，其中包括 FMEA 的讨论。",
    )
    add_body(
        doc,
        "Runqian Yang 请礼来审阅正在更新的 M+53 杂质调查报告。Richard Morgan 表示打算在 10 月 20 日或 21 日审。",
    )
    add_body(
        doc,
        "她还索要工艺流程图（PFD）。Richard Morgan 说，密度已于 10 月 19 日更新，文件正在审批，预计很快可以发出。",
    )
    add_body(
        doc,
        "会上也讨论了分析方法验证方案。凯莱英正在按礼来的 comment 更新；David Quirke 计划把礼来的 comment 收口，以便推进批准。",
    )

    add_h1(doc, "9. PPQ 留样")
    add_body(
        doc,
        "Runqian Yang 请礼来明确：PPQ 的哪些原料和中间体需要留样。Sarah Downing 说会带回礼来团队，并提到 Lara 用过的做法有可能作为起点。",
    )
    add_body(
        doc,
        "Richard Morgan 建议凯莱英先起草完整的拟留样清单，然后礼来审阅并补充。Sarah Downing 也建议把 Lara 的做法作为可能的基础文件分享出来。",
    )

    add_h1(doc, "10. M+53 调查与树脂策略")
    add_body(doc, "凯莱英说明，还有两项调查需要新脱保护的树脂：")
    add_bullet(doc, "用 5% 哌啶/DMF 替代现有 DeMtt 后洗涤条件的实验。")
    add_bullet(doc, "5,000 ppm 乙醛加铁加标实验。")
    add_body(
        doc,
        "随后讨论了树脂有限。Runqian Yang 表示大约还剩 200 g，建议脱保护 100 g，另外 100 g 先不动。Subha Mukherjee 支持节约物料，因为前面的实验消耗树脂很快。",
    )
    add_h2(doc, "5 g 实验规模")
    add_body(
        doc,
        "Subha Mukherjee 建议后续实验考虑以 5 g 而不是 10 g 作为常规规模，这样剩余树脂有可能多做几项实验。Shivani Choudhary 认为可以接受，前提是较小规模下混合和搅拌仍然有代表性。",
    )
    add_body(
        doc,
        "Shivani Choudhary 要求提供计划中实验的搅拌视频，以便礼来判断 5 g 规模看起来是否合适。Runqian Yang 说视频已经录了，可以分享。",
    )
    add_body(doc, "额外的 39mer-on-resin 合成已经开始，凯莱英将提供时间线。")

    add_h1(doc, "11. E13 / E14 后续")
    add_body(
        doc,
        "E13：礼来要求用 NMR 和 DNPH 衍生化，复测这些实验所用的 DIPEA。凯莱英说样品在，结果还没出来。",
    )
    add_body(
        doc,
        "Matthew Jacob 还追问了 E13 实验要求做的回顾性氮气流量确定。Runqian Yang 说尚未完成，将补进行动项。",
    )
    add_body(
        doc,
        "E14：六个氮气鼓泡流量实验已经开始。结果预计在 10 月 22–23 日左右，测量点为 4、18 和 36 小时。Matthew Jacob 要求 4 小时的结果一出来就发，不要等全套数据齐了再发。",
    )

    add_h1(doc, "后续行动")
    add_h2(doc, "凯莱英 / Runqian Yang / 团队")
    for line in [
        "回复已退回的纯化 / MOR 文件 comment，推进关闭和溶剂试车准备。",
        "继续 M+53 调查报告，并把正在更新的草稿提供给礼来审阅。",
        "如有可比的换热器性能数据，予以确认；并用溶剂试车确立当前温控安排的表现。",
        "确认层析柱和房间的温度控制、监测安排及其表现。",
        "起草完整的 PPQ 留样清单，覆盖计划中的物料和中间体，供礼来审阅。",
        "分享已经开始的额外 39mer-on-resin 合成的时间线。",
        "评估后续调查实验能否统一使用 5 g 规模，并确保混合和搅拌仍然有代表性。",
        "分享 5 g 反应搅拌 / 混合的视频，供礼来审阅。",
        "E13 所用 DIPEA 的 NMR 和 DNPH 衍生化复测结果出来后提供。",
        "完成并记录 E13 实验的回顾性氮气流量评估。",
        "分享 E14 氮气研究结果；4 小时数据一有就发，不要等全套数据。",
        "提供正在进行的调查实验完整清单。这是 Richard Morgan 具体要求的。",
        "分享 Shivani Choudhary 所要的工程 / 氮气对比幻灯或链接。",
    ]:
        add_bullet(doc, line)

    add_h2(doc, "礼来团队")
    for line in [
        "决定 pH 监测改造必须在 demo 之前完成，还是可以后置。需要同时考虑 GMP / CPP 风险，以及改造后可能还要再做一次溶剂试车。",
        "与 Andrew James Mouat 澄清额外过滤点的理由和时间要求，特别是 demo 之前是否必须。",
        "PAR 统计分析完成后，确定 / 收口 CPP 清单。",
        "Richard Morgan 审阅 M+53 调查报告草稿。会上所说时间为 10 月 20–21 日。",
        "推进 PFD 的最终批准和发出。会上说密度更新后，文件正在审批。",
        "David Quirke 收口分析方法验证方案上的 comment，以便推进批准。",
        "Sarah Downing / 礼来团队审阅 PPQ 留样要求；合适的话，提供 Lara 的留样做法作为起始参考。",
        "Andrew James Mouat / 相关礼来同事把实验室试验形成的按质量配制缓冲液方法提供给凯莱英。",
    ]:
        add_bullet(doc, line)

    add_h1(doc, "需继续跟踪的关键未决事项 / 风险")
    add_body(doc, "优先级最高的准备问题：")
    for line in [
        "pH 测量：demo 能否带着现有在线 pH 测量进行，还是必须先改造？",
        "额外过滤器：拟定的过滤变更对 demo 是否必需，尤其考虑到微生物标准；还是可以等到 pre-PPQ 改造窗口？",
        "温度控制：现有进口加热 / 控制，在实际操作条件下能否充分控制流动相 / 层析柱温度？溶剂试车就是用来提供证据的。",
        "DIPEA 控制：2–8°C 储存和按桶检测正在推进，但仍需要足够的稳定性数据，才能确定检测与使用之间的适当关系。",
        "M+53 关闭：尚未完成的调查和 CAPA 内容，仍关系 PPQ 文件能否往下走。",
        "树脂节约：剩余调查用树脂有限，因此要确认 5 g 实验规模是否成立，以及新的 39mer-on-resin 何时能供应。",
    ]:
        add_bullet(doc, line)

    add_h1(doc, "术语 / 待核实")
    add_table(
        doc,
        ["原文", "本文用法", "备注"],
        [
            ["JPT", "JPT", "全称原文未给出。待核实。"],
            ["MOR", "MOR，生产操作记录", "全称在本项目材料中未展开。"],
            ["solvent run", "溶剂试车", ""],
            ["demo", "demo 批", "原文写明该批是 GMP 批次。正式批次名称待核实。"],
            ["PV", "工艺验证（PV）", "与 pre-PPQ 同时出现，二者关系按原文保留。"],
            ["pre-PPQ", "pre-PPQ", "PPQ 前的改造窗口。"],
            ["PAR", "PAR", "按 proven acceptable range（已证明可接受范围）理解。"],
            ["CPP", "关键工艺参数（CPP）", ""],
            ["DeMtt", "DeMtt", "原文写法。是否即 Mtt 脱保护，待核实。"],
            ["Lara", "Lara", "是否即此前文件中的 LARRA，待核实。"],
            ["DNPH", "DNPH 衍生化", "通常指 2,4-二硝基苯肼。原文未展开。"],
            ["PFD", "工艺流程图（PFD）", ""],
            ["39mer-on-resin", "39mer-on-resin", "树脂上的 39mer。"],
            ["Runqian Yang18280", "Runqian Yang", "18280 为原文署名后缀，不作为姓名的一部分。"],
        ],
        [4.2, 5.4, 8.0],
        size=9,
    )
    save_pair(doc, TRANS_NAME)


SUMMARY = [
    "DIPEA 在 PV 前加严控制。原则上对齐：转入 2–8°C 储存，按桶跟踪乙醛、分别检测和放行，达不到要求的不使用。乙醛水平在供应商之间有差异，demo 所用批次的桶与桶之间也有差异。长期稳定性，以及检测应距离使用多近，还要数据支持。",
    "讨论中的四份纯化文件已退回，comment 需要回复。我方打算会后处理，以支持溶剂试车。礼来实验室试配缓冲液：1 M Tris 浓储液 pH 7.4，稀释到 50 mM 后 pH 7.65，接近预期。按质量配制的方法正经由 Andrew James Mouat 转出。稀释后的 pH 仍要核实，因为该 pH 是 CPP。",
    "在线 pH 改造和额外过滤都没有最终定下来。第一台在线混合器下游已有 pH 监测，当前高流速下读数可能不够稳定；改造大约增加 5 天，我方提议 demo 之后、pre-PPQ 之前再改。礼来认为 demo 是 GMP 批次，pH 若是 CPP，不宜接受不可靠的测量，该位置取离线样也不方便。改造如果后置，demo 之后、PV 之前可能还要再做一次溶剂试车。额外过滤是否必须在 demo 前安装，直接问 Andrew James Mouat。",
    "温控用溶剂试车看实际表现。较大浓盐罐因 Tris-HCl 对温度敏感已有主动温控，其余流动相罐并不都有同等控制，层析系统进口有加热。新配置下，进口换热器在要求的流量和体积下是否有效，尚未确立。层析柱夹套 / 出口以及房间温控的具体安排也要再确认。目前预计为 CPP 的有温度、进料 pH、流动相 pH；完整清单等 PAR 统计分析完成。",
    "PPQ 文件除 M+53 调查和 CAPA 需要更新的部分外，大部分大体可接受；FMEA 等仍有 comment 或尚未审完。M+53 报告草稿，Richard Morgan 打算 10/20–10/21 审。PFD 已于 10/19 更新密度，正在审批。分析方法验证方案我方按 comment 更新，David Quirke 收口礼来意见。PPQ 留样由我方先起草完整清单，礼来审阅补充，Lara 的做法可作为起点。树脂大约剩余 200 g，建议脱保护 100 g、另 100 g 先留着。后续实验可考虑统一 5 g，前提是搅拌仍有代表性。39mer-on-resin 已开始补合成。E14 已开始，4 小时数据一有就发。",
]

TOPICS = [
    {
        "title": "1. DIPEA / 乙醛控制",
        "rows": [
            ("礼来关注 / 要求",
             "• PV 之前增加 DIPEA 的操作控制：转入 2–8°C 温控储存；在长期稳定性更清楚之前，考虑把检测时间明显移到更接近批次使用的时候。\n"
             "• 乙醛水平在不同供应商之间有差异，demo 所用批次的不同桶之间也有差异。需要能支持货架期和使用期限决定的数据。"),
            ("我方会上回应",
             "• 同意 2–8°C 控制。\n"
             "• 按桶跟踪乙醛，分别检测、分别放行；达不到要求的不使用。"),
            ("会上共识",
             "原则上支持加严储存和按桶控制。长期稳定性及检测与使用的时间关系，仍要证据并进一步对齐。"),
        ],
    },
    {
        "title": "2. 纯化文件与溶剂试车",
        "rows": [
            ("礼来关注 / 要求",
             "• 支持溶剂试车的纯化文件需要定稿。\n"
             "• 讨论中的四份文件已经退回，comment 需要回复。"),
            ("我方会上回应",
             "• 打算会后处理这些 comment。"),
            ("会上共识",
             "会后回复 comment，推进文件关闭和溶剂试车准备。"),
        ],
    },
    {
        "title": "3. 缓冲液配制",
        "rows": [
            ("礼来关注 / 要求",
             "• 实验室已试配：pH 7.4 的 1 M Tris 浓储液，稀释到 50 mM 后 pH 7.65，接近预期。\n"
             "• 稀释后仍要核实 pH，因为这个 pH 是 CPP。\n"
             "• 完全按质量配制的方法，正通过 Andrew James Mouat 转交。"),
            ("会上共识",
             "稀释结果接近预期。按质量配制的方法转给我方使用。稀释后 pH 仍需核实。"),
        ],
    },
    {
        "title": "4. 过滤",
        "rows": [
            ("礼来关注 / 要求",
             "• 较小浓缩液罐附近另要求增加的过滤器，是否必须在 demo 之前安装。该项要求来自 Andrew James Mouat，理由直接向他确认。\n"
             "• 这些过滤器对 demo 批可能是有用的微生物风险控制。"),
            ("会上共识",
             "会上说明的现有布置：纯化水和乙腈进料为 0.22 µm；最终流动相罐下游为 0.45 µm，可以改为 0.22 µm。额外过滤器是 demo 前安装，还是放到计划中的 pre-PPQ 改造窗口，尚未决定。"),
        ],
    },
    {
        "title": "5. 温度控制",
        "rows": [
            ("礼来关注 / 要求",
             "• 进口换热器在要求的流量和体积下，是否已经证明能有效控制。\n"
             "• 层析柱周围的温度监测，包括可能的夹套 / 出口监测，以及房间温度控制，需要确认实际表现。"),
            ("我方会上回应",
             "• 新配置下，进口换热器的有效性尚未确立。用溶剂试车产生性能数据。\n"
             "• 同时核对历史或可比设备信息，以及设施控制。"),
            ("会上共识",
             "较大浓盐罐因 Tris-HCl 对温度敏感，已有主动温控；其他流动相罐并不都有同等控制；层析系统进口有加热，进入的流动相有温度测量。用溶剂试车评估这套安排。层析柱 / 夹套布置和房间控制的表现仍要确认。"),
        ],
    },
    {
        "title": "6. 在线 pH 与改造时机",
        "rows": [
            ("礼来关注 / 要求",
             "• demo 是 GMP 批次。pH 若是 CPP，不宜接受不可靠的 pH 测量。\n"
             "• 该点因管线布置，不方便取离线样。\n"
             "• 改造如果放到 demo 之后，PV 之前可能还要再做一次溶剂试车，以证明改造后的设备。无论改造放在 demo 前还是后，进度都可能受影响。"),
            ("我方会上回应",
             "• 第一台在线混合器下游已有 pH 监测，当前高流速下读数可能不够稳定。\n"
             "• 改造大约增加 5 天。提议 demo 之后、pre-PPQ 之前再改。"),
            ("会上共识",
             "未决。会下决定：改造必须在 demo 之前完成，还是可以后置并随后做确认。"),
        ],
    },
    {
        "title": "7. CPP",
        "rows": [
            ("礼来关注 / 要求",
             "• 目前预计为关键参数的有：温度、进料 pH、流动相 pH。\n"
             "• 完整清单等 PAR 统计分析完成后再定。"),
            ("会上共识",
             "上述三项按目前预计的 CPP 对待。最终清单未定。"),
        ],
    },
    {
        "title": "8. PPQ 文件、M+53 报告、PFD 与分析方法验证方案",
        "rows": [
            ("礼来关注 / 要求",
             "• 除 M+53 调查和 CAPA 需要更新的内容外，大部分 PPQ 文件大体可接受。具体文件仍有 comment 或尚未审完，包括 FMEA。\n"
             "• Richard Morgan 审阅正在更新的 M+53 杂质调查报告，时间为 10/20 或 10/21。\n"
             "• PFD 的密度已于 10/19 更新，文件正在审批，预计很快发出。\n"
             "• David Quirke 收口分析方法验证方案上的礼来 comment，以便推进批准。"),
            ("我方会上回应",
             "• 请礼来给出 PPQ 文件的整体批准状态，并审阅 M+53 调查报告草稿。\n"
             "• 索要 PFD。分析方法验证方案按礼来 comment 更新。"),
            ("会上共识",
             "除 M+53 调查和 CAPA 需要更新的部分外，大部分文件大体可接受。仍有 comment 或未审完的文件继续处理。报告、PFD 和验证方案按上面的安排推进。"),
        ],
    },
    {
        "title": "9. PPQ 留样",
        "rows": [
            ("礼来关注 / 要求",
             "• Sarah Downing 将留样范围带回礼来团队。Lara 的做法有可能作为起点，也可作为基础文件分享。\n"
             "• Richard Morgan 建议我方先出完整的拟留样清单，礼来再审阅并补充。"),
            ("我方会上回应",
             "• 请礼来明确 PPQ 哪些原料和中间体需要留样。"),
            ("会上共识",
             "我方起草完整清单，礼来审阅补充。Lara 做法可作为参考。"),
        ],
    },
    {
        "title": "10. M+53 树脂与 5 g 规模",
        "rows": [
            ("礼来关注 / 要求",
             "• 剩余树脂有限，支持节约使用。前面的实验消耗很快。\n"
             "• 后续实验可考虑以 5 g 而不是 10 g 作为常规规模，以便剩余树脂多做几项。较小规模下混合和搅拌必须仍然有代表性。\n"
             "• 提供计划中实验的搅拌视频，供判断 5 g 是否合适。Shivani Choudhary 同时要工程 / 氮气对比材料。Richard Morgan 要正在进行的调查实验完整清单。"),
            ("我方会上回应",
             "• 还有两项调查需要新脱保护的树脂：用 5% 哌啶/DMF 替代现有 DeMtt 后洗涤；5,000 ppm 乙醛加铁加标。\n"
             "• 树脂大约剩余 200 g。建议脱保护 100 g，另外 100 g 先不动。\n"
             "• 5 g 搅拌视频已经录了，可以分享。额外的 39mer-on-resin 合成已经开始，时间线另提供。"),
            ("会上共识",
             "节约剩余树脂。5 g 可以作为后续实验规模的考虑方向，先用视频确认搅拌是否有代表性。39mer-on-resin 的时间线另发。"),
        ],
    },
    {
        "title": "11. E13 / E14",
        "rows": [
            ("礼来关注 / 要求",
             "• E13：用 NMR 和 DNPH 衍生化，复测实验所用 DIPEA。\n"
             "• E13 的回顾性氮气流量确定需要完成。\n"
             "• E14：4 小时结果一出来就发，不要等 18 小时和 36 小时的全套数据。"),
            ("我方会上回应",
             "• E13 的样品在，复测结果还没出来。氮气流量回顾尚未完成，补进行动项。\n"
             "• E14 六个氮气鼓泡流量实验已经开始。结果预计 10/22–10/23，测量点为 4、18 和 36 小时。"),
            ("会上共识",
             "E13 复测结果出来后提供，氮气流量回顾补做。E14 按 4 小时数据先发、全套约 10/22–10/23 的方式分享。"),
        ],
    },
]

ACTIONS = [
    ["1", "回复已退回的四份纯化 / MOR 文件 comment，推进关闭，支持溶剂试车。", "纯化文件", "会后，未定"],
    ["2", "继续更新 M+53 杂质调查报告，并把草稿提供给礼来审阅。", "M+53", "未定"],
    ["3", "审阅 M+53 调查报告草稿。", "M+53", "10/20–10/21"],
    ["4", "确认有无可比换热器数据，并用溶剂试车确认当前温控在要求流量和体积下的表现。", "温度控制", "溶剂试车"],
    ["5", "确认层析柱夹套 / 出口以及房间的温度控制和监测安排，并说明实际表现。", "温度控制", "未定"],
    ["6", "决定在线 pH 改造必须在 demo 之前完成，还是可以后置并在其后做确认。同时考虑 GMP 批次、CPP，以及是否需要改造后再做一次溶剂试车。", "在线 pH", "未定"],
    ["7", "向 Andrew James Mouat 确认较小浓缩液罐额外过滤的理由，以及 demo 前是否必须安装。", "过滤", "未定"],
    ["8", "把实验室试验形成的按质量配制缓冲液方法提供给我方。稀释后 pH 仍要核实。", "缓冲液", "未定"],
    ["9", "PAR 统计分析完成后，确定 CPP 清单。目前预计包括温度、进料 pH、流动相 pH。", "CPP", "PAR 统计分析完成后"],
    ["10", "完成 PFD 审批并发出。密度已于 10/19 更新。", "PFD", "审批中，尽快"],
    ["11", "收口分析方法验证方案的 comment，我方按 comment 更新，推进批准。", "分析方法验证方案", "未定"],
    ["12", "起草 PPQ 原料和中间体的完整留样清单，供礼来审阅和补充。", "PPQ 留样", "未定"],
    ["13", "审阅 PPQ 留样要求；合适的话，提供 Lara 的留样做法作为起点。", "PPQ 留样", "未定"],
    ["14", "提供已开始的 39mer-on-resin 补合成时间线。", "M+53 / 树脂", "未定"],
    ["15", "评估后续调查实验能否统一按 5 g 进行，并确认混合和搅拌仍有代表性。", "M+53 / 树脂", "未定"],
    ["16", "提供 5 g 反应的搅拌视频，以及工程 / 氮气对比材料。", "M+53 / 树脂", "未定"],
    ["17", "提供正在进行的调查实验完整清单。", "M+53", "未定"],
    ["18", "提供 E13 所用 DIPEA 的 NMR 和 DNPH 衍生化复测结果。", "E13", "结果出来后"],
    ["19", "完成并记录 E13 实验的回顾性氮气流量评估。", "E13", "未定"],
    ["20", "分享 E14 氮气鼓泡流量结果。4 小时数据一有即发，不等 18 小时和 36 小时的全套结果。", "E14", "4 小时数据一有即发；全套约 10/22–10/23"],
]


def build_minutes():
    doc = new_doc()
    title(
        doc,
        "礼来客户会会议纪要",
        "GG917 SPPS and Purification Start Readiness – 10/20/2026",
    )
    add_h1(doc, "一、会议要点")
    for s in SUMMARY:
        add_bullet(doc, s)
    add_h1(doc, "二、分议题纪要")
    for topic in TOPICS:
        add_h2(doc, topic["title"])
        add_table(doc, ["项目", "内容"], topic["rows"], [3.2, 14.4], size=10, first_col_bold=True)
    add_h1(doc, "三、后续 Action Items")
    add_table(doc, ["#", "行动", "议题", "时间"], ACTIONS, [1.1, 10.3, 3.2, 3.0], size=9.5)
    save_pair(doc, MIN_NAME)


def to_pdf(name):
    for i, folder in enumerate((MEET, ROOT)):
        src = folder / name
        profile = f"file:///tmp/lo-oct20-{i}"
        subprocess.run(
            [
                "soffice", "--headless", "--norestore",
                f"-env:UserInstallation={profile}",
                "--convert-to", "pdf", "--outdir", str(folder), str(src),
            ],
            check=True,
        )
        pdf = folder / (Path(name).stem + ".pdf")
        if not pdf.exists() or pdf.stat().st_size < 1000:
            raise SystemExit(f"PDF was not created: {pdf}")
        print(pdf)


if __name__ == "__main__":
    build_translation()
    build_minutes()
    to_pdf(TRANS_NAME)
    to_pdf(MIN_NAME)
