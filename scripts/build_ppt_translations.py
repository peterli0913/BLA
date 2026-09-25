#!/usr/bin/env python3
"""Build bilingual (EN/CN) and English-only versions of the humidified-N2 and PW decks.

Usage: python3 scripts/build_ppt_translations.py
"""
import copy
import re
import subprocess
import tempfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
ADD_SLIDE = Path.home() / ".cursor/skills/pptx/scripts/add_slide.py"

N2_SRC = ROOT / "湿氮气装置改造使用策略2026.09.24.pptx"
PW_SRC = ROOT / "纯水使用方案.pptx"
N2_OUT = {"bi": OUT / "湿氮气装置改造使用策略2026.09.24_中英对照.pptx",
          "en": OUT / "Humidified_Nitrogen_System_Operating_Strategy_2026-09-24_EN.pptx"}
PW_OUT = {"bi": OUT / "纯水使用方案_中英对照.pptx",
          "en": OUT / "Purified_Water_Use_Plan_EN.pptx"}

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
CJK = re.compile(r"[\u3000-\u303f\u3400-\u9fff\uff00-\uffef]")
EN_FONT = "Arial"
CN_GRAY = "595959"
RED = "FF0000"
RPR_ORDER = ["ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
             "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill",
             "latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"]
PPR_ORDER = ["lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts",
             "buFontTx", "buFont", "buNone", "buAutoNum", "buChar", "buBlip", "tabLst",
             "defRPr", "extLst"]


def qn(tag):
    return "{%s}%s" % (A, tag)


def insert_ordered(parent, el, order):
    name = etree.QName(el).localname
    for old in parent.findall(qn(name)):
        parent.remove(old)
    idx = order.index(name)
    for i, child in enumerate(parent):
        cname = etree.QName(child).localname
        if cname in order and order.index(cname) > idx:
            parent.insert(i, el)
            return
    parent.append(el)


def first_rpr(txbody):
    for r in txbody.iter(qn("r")):
        rpr = r.find(qn("rPr"))
        if rpr is not None:
            return copy.deepcopy(rpr)
    return etree.Element(qn("rPr"))


def first_ppr(txbody):
    p = txbody.find(qn("p"))
    if p is not None and p.find(qn("pPr")) is not None:
        return copy.deepcopy(p.find(qn("pPr")))
    return etree.Element(qn("pPr"))


def make_run(text, base, size=None, bold=None, color=None, cn=False):
    r = etree.Element(qn("r"))
    rpr = copy.deepcopy(base)
    for k in ("dirty", "smtClean", "err"):
        rpr.attrib.pop(k, None)
    rpr.set("lang", "zh-CN" if cn else "en-US")
    rpr.set("altLang", "en-US" if cn else "zh-CN")
    if size:
        rpr.set("sz", str(int(round(size * 100))))
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if color:
        sf = etree.Element(qn("solidFill"))
        etree.SubElement(sf, qn("srgbClr")).set("val", color)
        insert_ordered(rpr, sf, RPR_ORDER)
    if not cn:
        latin = etree.Element(qn("latin"))
        latin.set("typeface", EN_FONT)
        insert_ordered(rpr, latin, RPR_ORDER)
    r.append(rpr)
    t = etree.SubElement(r, qn("t"))
    t.text = text
    if text != text.strip():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return r


def make_para(runs, base_ppr, algn=None, line=None, before=None, after=None, end_size=None):
    p = etree.Element(qn("p"))
    ppr = copy.deepcopy(base_ppr)
    if algn:
        ppr.set("algn", algn)
    if line:
        ln = etree.Element(qn("lnSpc"))
        etree.SubElement(ln, qn("spcPct")).set("val", str(int(line * 1000)))
        insert_ordered(ppr, ln, PPR_ORDER)
    for tag, val in (("spcBef", before), ("spcAft", after)):
        if val is not None:
            el = etree.Element(qn(tag))
            etree.SubElement(el, qn("spcPts")).set("val", str(int(val * 100)))
            insert_ordered(ppr, el, PPR_ORDER)
    p.append(ppr)
    for r in runs:
        p.append(r)
    if end_size:
        etree.SubElement(p, qn("endParaRPr")).set("sz", str(int(end_size * 100)))
    return p


def segs(text):
    """Text may be a plain string or a list of (text, opts) segments."""
    return [(text, {})] if isinstance(text, str) else text


def arrows(text):
    """Split on → so arrows stay bold, as in the source deck."""
    out = []
    for i, part in enumerate(text.split("→")):
        if i:
            out.append(("→", {"bold": True}))
        if part:
            out.append((part, {}))
    return out


def render_blocks(txbody, blocks, lang, st):
    """Replace all paragraphs of txbody.

    blocks: list of (kind, en, cn); kind 'h' = bilingual inline header, 'p' = EN then CN line,
    'in' = EN and CN inline on one line (labels/titles).
    st: style dict with en, cn sizes, bold, cn_color, line, algn, before, after.
    """
    base = first_rpr(txbody)
    ppr = first_ppr(txbody)
    for p in txbody.findall(qn("p")):
        txbody.remove(p)
    en_sz, cn_sz = st["en"], st.get("cn", st["en"])
    bold = st.get("bold")
    cn_color = st.get("cn_color", CN_GRAY)
    en_color = st.get("en_color")
    kw = dict(algn=st.get("algn"), line=st.get("line"))
    for kind, en, cn in blocks:
        hsz = st.get("h_en", en_sz) if kind == "h" else en_sz
        hcn = st.get("h_cn", cn_sz) if kind == "h" else cn_sz
        hb = True if kind == "h" else bold

        def en_runs(size=hsz):
            return [make_run(t, base, size, o.get("bold", hb), o.get("color", en_color))
                    for t, o in segs(en)]

        def cn_runs(size=hcn):
            return [make_run(t, base, size, o.get("bold", hb), o.get("color", cn_color), cn=True)
                    for t, o in segs(cn)]

        before = st.get("h_before") if kind == "h" else st.get("before")
        if lang == "en" or not cn:
            txbody.append(make_para(en_runs(), ppr, before=before, after=st.get("after"), **kw))
        elif kind in ("h", "in"):
            runs = en_runs() + [make_run("   ", base, hcn, hb, cn=True)] + cn_runs()
            txbody.append(make_para(runs, ppr, before=before, after=st.get("after"), **kw))
        else:
            txbody.append(make_para(en_runs(), ppr, before=before, after=0, **kw))
            txbody.append(make_para(cn_runs(), ppr, before=0, after=st.get("after"), **kw))


def shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise KeyError(name)


def place(sh, x=None, y=None, w=None, h=None):
    if x is not None:
        sh.left = Inches(x)
    if y is not None:
        sh.top = Inches(y)
    if w is not None:
        sh.width = Inches(w)
    if h is not None:
        sh.height = Inches(h)


def line_to(conn, begin=None, end=None):
    if begin:
        conn.begin_x, conn.begin_y = Inches(begin[0]), Inches(begin[1])
    if end:
        conn.end_x, conn.end_y = Inches(end[0]), Inches(end[1])


def fill(slide, name, blocks, lang, st_bi, st_en, geo_bi=None, geo_en=None):
    sh = shape(slide, name)
    render_blocks(sh.text_frame._txBody, blocks, lang, st_bi if lang == "bi" else st_en)
    geo = geo_bi if lang == "bi" else geo_en
    if geo:
        place(sh, **geo)
    return sh


# ---------------------------------------------------------------- humidified N2 deck

N2_S1_TITLE = [("p", "Humidified Nitrogen – Operating Strategy", "湿氮气-使用策略")]
N2_S1_BODY = [
    ("h", "Overall strategy for the humidified nitrogen system:", "湿氮气装置整体策略："),
    ("h", "I. First Use", "一、首次使用时："),
    ("p", "1. Alkaline cleaning: Recirculate alkaline solution and then purified water through "
          "the unit; finally, purge the unit with nitrogen and drain at the low points.",
     "1. 碱液清洗：依次使用碱液、纯化水进行循环冲洗，最后使用氮气对设备进行吹扫低点排净。"),
    ("p", "2. WFI rinse: Using a sterile bag and an aseptic connector, recirculate and "
          "spray-rinse the unit with WFI, then drain at the low points.",
     "2. 注射水淋洗：使用无菌软袋，通过无菌链接器，使用注射水对装置进行循环喷淋淋洗后，低点排净。"),
    ("p", "3. WFI charge: Charge WFI again through a sterile bag for use.",
     "3. 注射水进料：再次通过无菌袋加入注射水进行使用。"),
    ("h", "II. Periodic WFI Replacement", "二、注射水周期性更换"),
    ("p", "1. Drain at the low points.", "1. 低点排净。"),
    ("p", "2. Add fresh WFI, recirculate to rinse, then drain at the low points.",
     "2. 补充新注射水循环淋洗，低点排净。"),
    ("p", "3. Recharge WFI for use.", "3. 再次补入注射水进行使用。"),
    ("h", "III. Alternate Use of Two Humidifiers", "三、两套氮气加湿装置互补"),
    ("p", "Within the WFI replacement interval, the two units are periodically cleaned by "
          "recirculation per Section II and used alternately.",
     "在注射水周期范围内按照方式二，周期性循环清洗交替使用。"),
]
N2_DRAIN_TITLE = [("in", "Low-Point Drain and High-Point Purge Locations", "设备低点排净和高位吹扫点位")]
N2_S2_LABELS = {
    "文本框 105": [("p", "1. Material line into the clean area: fitted with a cleaning/purge "
                         "port, allowing cleaning and purging in both directions.",
                    "1. 进洁净区物料管路：设置有清洗吹扫口，能够往两个方向进行清洗吹扫。")],
    "文本框 14": [("p", "2. Filter strategy: filters are cleaned off-line.",
                   "2. 过滤器策略：过滤器采用离线清洗的方式。")],
    "文本框 29": [("p", "3. WFI pre-flush drain line.", "3. 注射水预冲洗排净管路。")],
    "文本框 44": [("p", "4. Equipment low-point drain line.", "4. 设备低点排净管路。")],
}
N2_S3_L1 = [("p", "1. Humidified N₂ line to the AFD: fitted with a cleaning/purge port, "
                  "allowing cleaning and purging in both directions.",
             "1. 进三合一的湿氮气管路：设置有清洗吹扫口，能够往两个方向进行清洗吹扫。")]
N2_S3_L2 = [("p", "2. Low-point drain of the humidified N₂ line to the AFD.",
             "2. 进三合一的湿氮气管路低点排净。")]
N2_S3_ASK = [
    ("h", "For Lilly's confirmation:", "需客户确认："),
    ("p", "Should the valves at the drain and purge points be installed to a 2D or a 3D "
          "dead-leg requirement?",
     "排净和吹扫点位阀门安装要求2D还是3D？"),
    ("p", "1. 2D — longer valve procurement lead time;", "1. 2D——阀门采购周期较长；"),
    ("p", "2. 3D — shorter valve procurement lead time; domestic equipment suppliers have "
          "confirmed that a 3D design can be cleaned effectively.",
     "2. 3D——阀门采购周期短，国内设备厂家反馈3D能够清洗干净。"),
]

# source picture on N2 slide 1: (5.40, 0.98) 7.02 x 6.09 in
N2_PIC = dict(x=6.17, y=1.05, w=6.25)


def build_n2(lang):
    prs = Presentation(N2_SRC)
    s1, s2, s3 = prs.slides

    pic = shape(s1, "图片 1")
    k = N2_PIC["w"] / (pic.width / 914400)
    old = (pic.left / 914400, pic.top / 914400)
    place(pic, N2_PIC["x"], N2_PIC["y"], N2_PIC["w"], pic.height / 914400 * k)

    def mapped(pt):
        return (N2_PIC["x"] + (pt[0] - old[0]) * k, N2_PIC["y"] + (pt[1] - old[1]) * k)

    fill(s1, "文本框 95", N2_S1_TITLE, lang,
         dict(en=18, cn=14, bold=True, cn_color=CN_GRAY, line=95),
         dict(en=18, bold=True),
         geo_bi=dict(x=0.55, y=1.0, w=5.55), geo_en=dict(x=0.55, y=1.05, w=5.55))
    fill(s1, "文本框 105", N2_S1_BODY, lang,
         dict(en=11.5, cn=10.5, h_en=12, h_cn=11, bold=False, line=97, after=2, h_before=4),
         dict(en=14, h_en=14.5, bold=False, line=100, after=3, h_before=6),
         geo_bi=dict(x=0.55, y=1.62, w=5.55, h=5.4), geo_en=dict(x=0.55, y=1.5, w=5.55, h=5.4))
    # arrows: tail on the diagram, head at section I items 1 and 2
    heads = {"bi": (2.72, 3.55), "en": (2.62, 3.45)}[lang]
    line_to(shape(s1, "直接箭头连接符 9"), begin=mapped((5.80, 2.37)), end=(6.02, heads[0]))
    line_to(shape(s1, "直接箭头连接符 11"), begin=mapped((5.75, 4.20)), end=(6.02, heads[1]))

    fill(s2, "文本框 15", N2_DRAIN_TITLE, lang,
         dict(en=16, cn=15, bold=True, algn="ctr"), dict(en=20, bold=True, algn="ctr"),
         geo_bi=dict(x=3.3, y=0.98, w=9.6), geo_en=dict(x=3.3, y=0.95, w=9.6))
    lab_bi = dict(en=12, cn=11, bold=True, line=100)
    lab_en = dict(en=14, bold=True, line=100)
    widths = {"文本框 105": 2.36, "文本框 14": 2.45, "文本框 29": 2.28, "文本框 44": 2.36}
    for name, blocks in N2_S2_LABELS.items():
        fill(s2, name, blocks, lang, lab_bi, lab_en,
             geo_bi=dict(w=widths[name]), geo_en=dict(w=widths[name]))

    fill(s3, "文本框 5", N2_DRAIN_TITLE, lang,
         dict(en=16, cn=15, bold=True, algn="ctr"), dict(en=20, bold=True, algn="ctr"),
         geo_bi=dict(x=3.3, y=0.92, w=9.6), geo_en=dict(x=3.3, y=0.9, w=9.6))
    fill(s3, "文本框 3", N2_S3_L1, lang, lab_bi, lab_en,
         geo_bi=dict(x=0.8, y=1.45, w=3.0), geo_en=dict(x=0.8, y=1.45, w=3.0))
    l2_y = {"bi": 2.9, "en": 2.8}[lang]
    fill(s3, "文本框 8", N2_S3_L2, lang, lab_bi, lab_en,
         geo_bi=dict(x=0.84, y=l2_y, w=2.95), geo_en=dict(x=0.84, y=l2_y, w=2.95))
    line_to(shape(s3, "直接箭头连接符 28"), end=(3.86, l2_y + 0.2))
    line_to(shape(s3, "直接箭头连接符 9"), end=(3.86, l2_y + 0.3))
    fill(s3, "文本框 26", N2_S3_ASK, lang,
         dict(en=13, cn=12, h_en=16, h_cn=15, bold=True, line=100, after=3),
         dict(en=15, h_en=17, bold=True, line=100, after=5),
         geo_bi=dict(x=0.76, y=4.0, w=4.75), geo_en=dict(x=0.76, y=3.9, w=4.75))

    prs.save(N2_OUT[lang])


# ---------------------------------------------------------------- purified water deck

PW_TITLES = {0: ("Drain and Water-Fill Flow", "排废及加水流程"),
             1: ("Sampling Flow", "取样流程"),
             2: ("Sampling Flow", "取样流程"),
             3: ("Purified Water Use Procedure", "纯化水使用流程")}
PW_S3_LABEL_EN = [("Purge ", {}), ("①", {"color": "5B9BD5"}), (" first → then purge ", {}),
                  ("②", {"color": RED})]
PW_S3_LABEL_CN = [("先吹扫", {}), ("①", {"color": "5B9BD5"}), ("→吹扫", {}), ("②", {"color": RED})]
PW_S4_BODY = [
    ("h", "1. Purified Water Use", "一、纯化水使用流程："),
    ("p", arrows("Connect the hose from the PW outlet to the fixed line → Open the water valve → "
                 "Flush PW to waste → After flushing, supply water to the point of use"),
     arrows("软管连接纯化水点至固定管线→打开水阀→纯化水开始排废→排废结束后进入使用点位")),
    ("p", arrows("Open the nitrogen purge valve → Purge toward the equipment side → Then purge "
                 "back toward the PW outlet → Close all valves when finished"),
     arrows("开启氮气吹扫阀门→向设备侧进行吹扫→完成后→向纯水点位方向吹扫→结束后关闭所有阀门")),
    ("h", "2. Sampling", "二、取样："),
    ("p", arrows("Connect the hose from the PW outlet to the fixed line → Open the water valve → "
                 "Flush PW to waste → After flushing, take a manual sample"),
     arrows("软管连接纯化水点至固定管线→打开水阀→纯化水开始排废→排废结束后进行手动取样")),
    ("h", "3. Filter Management Principle", "三、过滤器管理原则："),
    ("p", "Install the filter before each use; the filter's time in use shall not exceed 24 hours.",
     "每次使用前安装滤器，保证滤器使用时间不超过24小时"),
]
PW_HEAD = [("Action Item", "行动项"), ("Implementation Plan", "实施方案"), ("Target Date", "执行时间")]
# (slide index in source, row) -> (plan EN lines, plan CN lines, date lines)
PW_ROWS = {
    (4, 1): (["Practice: weekly cleaning; replacement every six months"],
             ["执行标准：每周清洁，每半年更换"], ["10/20/2026"]),
    (4, 2): (["Replace with PTFE hoses:", "1. Domestic brand: lead time ~3 weeks",
              "2. Imported brand: lead time ~8 weeks"],
             ["更换为PTFE材质软管", "1、国内品牌货期约3周", "2、国外品牌货期约8周"],
             ["1. 10/30/2026", "2. 11/30/2026"]),
    (4, 3): (["Update the procedure documents."], ["更新文件"], ["10/20/2026"]),
    (4, 4): (["Submit a change control and carry out the modification per the scheme shown on "
              "Slide 1 (valves designed to the 3D principle). Estimated durations:",
              "Procurement: ~6 weeks", "Installation: ~3 weeks",
              "Pickling and passivation: ~2 weeks",
              "Continuous sampling validation and incubation: ~3 weeks",
              "Line drying validation: ~2 weeks"],
             ["如首页展示方案提交变更并进行改造阀（满足3D原则），经过评估时间如下：", "采购周期约6周",
              "施工约3周", "酸洗钝化约2周", "连续取样验证并培养约3周", "管路干燥验证约2周"],
             ["01/30/2027"]),
    (5, 1): (["Permeable film has been procured (see photo at left)."],
             ["已经采购透气膜，如左图"], ["10/30/2026"]),
    (5, 2): (["Immediate corrective action"], ["立即整改"], ["09/30/2026"]),
    (5, 3): (["Assessment of sampling points is in progress."], ["取样点位评估正在进行"],
             ["10/30/2026"]),
    (5, 4): (["1. Short term: train on-site personnel",
              "2. Long term: install custom stainless-steel hose racks to keep hose ends more "
              "than 6 inches off the floor"],
             ["1、短期方案，进行现场人员培训", "2、长期方案，定制不锈钢水管支架，保证离地高度大于6 inches"],
             ["1. 09/30/2026", "2. 11/30/2026"]),
}
# (slide index, source slide index, source rows kept, column widths, body pt, header pt,
#  height of the photo row or "drop" to delete the photo). In the bilingual deck each
#  table slide is split in two so the text can stay large.
PW_TABLES = {
    "bi": [(4, 4, [1, 2, 3], (6.4, 4.7, 1.78), 18, 15, None),
           (5, 4, [4], (7.0, 4.3, 1.58), 17, 15, None),
           (6, 5, [1, 2], (7.0, 4.0, 1.73), 18, 15, 3.1),
           (7, 5, [3, 4], (7.0, 4.0, 1.73), 18, 15, "drop")],
    "en": [(4, 4, [1, 2, 3, 4], None, 16, 16, None),
           (5, 5, [1, 2, 3, 4], (7.5, 3.5, 1.73), 16, 16, 2.25)],
}
PW_HEAD_H = {"bi": 0.75, "en": 0.55}


def set_cell(cell, paras, size, bold=False, lang_cn=None, colors=None):
    """paras: list of (text, is_cn)."""
    tx = cell.text_frame._txBody
    base = first_rpr(tx)
    ppr = first_ppr(tx)
    for p in tx.findall(qn("p")):
        tx.remove(p)
    for text, is_cn in paras:
        color = (colors or {}).get(is_cn)
        tx.append(make_para([make_run(text, base, size, bold, color, cn=is_cn)], ppr,
                            line=100, after=2))
    tcpr = cell._tc.get_or_add_tcPr()
    tcpr.set("marL", str(Inches(0.08)))
    tcpr.set("marR", str(Inches(0.08)))
    tcpr.set("marT", str(Inches(0.05)))
    tcpr.set("marB", str(Inches(0.05)))


def cell_lines(cell):
    return [ln.strip() for ln in cell.text_frame.text.replace("\x0b", "\n").split("\n") if ln.strip()]


def fill_table(slide, lang, cfg):
    _, src_idx, keep_rows, widths, body_sz, head_sz, photo = cfg
    tbl_shape = [sh for sh in slide.shapes if sh.has_table][0]
    table = tbl_shape.table
    rows = list(table.rows)
    for ci, (en, cn) in enumerate(PW_HEAD):
        paras = [(en, False)] + ([(cn, True)] if lang == "bi" else [])
        set_cell(table.cell(0, ci), paras, head_sz, bold=True)
    for ri in keep_rows:
        lilly = cell_lines(table.cell(ri, 0))
        paras = [(t, bool(CJK.search(t))) for t in lilly]
        if lang == "en":
            paras = [p for p in paras if not p[1]]
        set_cell(table.cell(ri, 0), paras, body_sz)
        plan_en, plan_cn, dates = PW_ROWS[(src_idx, ri)]
        paras = [(t, False) for t in plan_en]
        if lang == "bi":
            paras += [(t, True) for t in plan_cn]
        set_cell(table.cell(ri, 1), paras, body_sz, colors={True: CN_GRAY})
        set_cell(table.cell(ri, 2), [(t, False) for t in dates], body_sz)
    for ri in sorted(set(range(1, len(rows))) - set(keep_rows), reverse=True):
        table._tbl.remove(rows[ri]._tr)
    if widths:
        for col, w in zip(table.columns, widths):
            col.width = Inches(w)
    head_h = PW_HEAD_H[lang]
    for i, r in enumerate(table.rows):
        r.height = Inches(head_h if i == 0 else 0.45)

    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    if photo == "drop":
        for pic in pics:
            pic._element.getparent().remove(pic._element)
    elif photo:
        # the photo sits under the text of the first data row, inside the Action Item column
        table.rows[1].height = Inches(photo)
        table.cell(1, 0)._tc.get_or_add_tcPr().set("anchor", "t")
        col0 = table.columns[0].width
        for pic in pics:
            pic.left = tbl_shape.left + col0 - pic.width - Inches(0.15)
            pic.top = tbl_shape.top + Inches(head_h + photo - 0.12) - pic.height
    return tbl_shape


def build_pw(lang):
    src = PW_SRC
    if lang == "bi":
        tmp = Path(tempfile.mkdtemp()) / "pw_split.pptx"
        subprocess.run(["python3", str(ADD_SLIDE), str(PW_SRC), "slide5.xml", "--after",
                        "slide5.xml", "-o", str(tmp)], check=True, capture_output=True)
        subprocess.run(["python3", str(ADD_SLIDE), str(tmp), "slide6.xml", "--after",
                        "slide6.xml", "-o", str(tmp)], check=True, capture_output=True)
        src = tmp
    prs = Presentation(src)
    slides = list(prs.slides)

    for i, (en, cn) in PW_TITLES.items():
        fill(slides[i], "文本框 1", [("in", en, cn)], lang,
             dict(en=18, cn=16, algn="r", cn_color=CN_GRAY), dict(en=20, algn="r"),
             geo_bi=dict(x=5.4, y=0.22, w=7.7), geo_en=dict(x=5.4, y=0.2, w=7.7))

    s1, s2, s3, s4 = slides[:4]
    fill(s1, "文本框 15", [("p", "② Water-fill path", "②加水流路")], lang,
         dict(en=15, cn=14, line=100, cn_color=None), dict(en=16),
         geo_bi=dict(x=11.25, y=2.28, w=2.05), geo_en=dict(x=11.2, y=2.5, w=2.1))
    fill(s1, "文本框 4", [("p", "① Drain path", "①排废流路")], lang,
         dict(en=15, cn=14, line=100, cn_color=None), dict(en=16),
         geo_bi=dict(x=11.4, y=3.4, w=1.9), geo_en=dict(x=11.4, y=3.62, w=1.9))
    fill(s2, "文本框 13", [("p", "① Drain → ② Sample", "①排废→②取样")], lang,
         dict(en=16, cn=15, line=100, cn_color=None), dict(en=18),
         geo_bi=dict(x=10.0, y=1.38, w=2.8), geo_en=dict(x=10.0, y=1.55, w=2.8))
    fill(s3, "文本框 9", [("in", PW_S3_LABEL_EN, PW_S3_LABEL_CN)], lang,
         dict(en=18, cn=18, cn_color=None), dict(en=20),
         geo_bi=dict(x=2.9, y=1.0, w=8.0), geo_en=dict(x=3.4, y=1.0, w=6.5))
    fill(s4, "文本框 3", PW_S4_BODY, lang,
         dict(en=20, cn=16, h_en=22, h_cn=20, bold=False, line=110, after=5, h_before=12),
         dict(en=22, h_en=24, bold=False, line=112, after=8, h_before=14),
         geo_bi=dict(x=0.65, y=1.05, w=12.0), geo_en=dict(x=0.65, y=1.15, w=12.0))

    for cfg in PW_TABLES[lang]:
        fill_table(slides[cfg[0]], lang, cfg)

    prs.save(PW_OUT[lang])


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    for lang in ("bi", "en"):
        build_n2(lang)
        build_pw(lang)
    for p in list(N2_OUT.values()) + list(PW_OUT.values()):
        print(p.relative_to(ROOT))
