#!/usr/bin/env python3
"""English-only version of 139788BLA.pptx (mobile-phase preparation control logic).

Usage: python3 scripts/build_139788_bla_en.py
"""
import copy
import io
from pathlib import Path

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "139788BLA.pptx"
OUT = ROOT / "deliverables" / "139788BLA_EN.pptx"

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
TIGHT_INS = "18288"  # 0.02 in, for labels that sit on arrows and must stay one line per paragraph

# shape_id -> English text, one string per source paragraph ("\n" = line break inside the paragraph)
TEXT = {
    49: ["Flow\ncheck"],
    50: ["Mobile phase routed to Tank A"],
    51: ["Mobile phase routed to waste tank"],
    52: ["Program stop"],
    55: ["Set conductivity HH/LL alarm limits per process requirements"],
    56: ["Not cleared in 10 min", "or any flow > ±5%"],
    58: ["All flows within ±3%\n(pre-drain approx. 2–5 min)"],
    59: ["5 s after conductivity HH/LL alarm,", "or any flow > ±3%: divert to waste"],
    60: ["Conductivity alarm\nauto-enabled"],
    62: ["Stable for 10 s"],
    64: ["Conductivity alarm / flow deviation cleared within 10 min"],
    65: ["Mobile Phase Preparation Control Logic"],
    66: ["Mobile phase routed to waste tank"],
    67: ["Program start"],
    68: ["Conc. salt flow SV\n(setpoint, kg/min)"],
    69: ["Set blending parameters per process"],
    72: ["Conc. salt PID control"],
    74: ["PW PID control"],
    75: ["PW flow SV\n(setpoint, kg/min)"],
    77: ["Acetonitrile PID control"],
    79: ["Acetonitrile flow SV\n(setpoint, kg/min)"],
}
# Two-line setpoint labels sit bottom-anchored just above their arrow: label id -> (connector id, arrow at connector bottom?)
ABOVE_ARROW = {75: (76, False), 68: (73, False), 79: (78, True)}
ABOVE_ARROW_H = 330000
# The connector through this box's midline falls between two lines of text, as in the source.
CENTER = {49}
TIGHT = {49, 50, 51, 52, 55, 56, 58, 59, 60, 66, 67, 69}
TITLE_ID, TITLE_WIDTH_IN = 65, 3.4

# Rows (y ranges, in pixels) holding the Chinese line of each column of the tracker screenshot.
IMG_BANDS = [((1, 78), (15, 33)), ((80, 740), (33, 53)), ((742, 845), (33, 51)), ((847, 972), (15, 33))]
IMG_INK = 45


def set_paragraph(p, text):
    runs = p.findall(A + "r")
    base = copy.deepcopy(runs[0].find(A + "rPr"))
    for child in list(p):
        if child.tag in (A + "r", A + "br", A + "fld"):
            p.remove(child)
    base.set("lang", "en-US")
    base.attrib.pop("altLang", None)
    for k in ("dirty", "err", "smtClean"):
        base.attrib.pop(k, None)
    end = p.find(A + "endParaRPr")
    for i, line in enumerate(text.split("\n")):
        if i:
            br = p.makeelement(A + "br", {})
            br.append(copy.deepcopy(base))
            _insert_before(p, br, end)
        r = p.makeelement(A + "r", {})
        r.append(copy.deepcopy(base))
        t = r.makeelement(A + "t", {})
        t.text = line
        r.append(t)
        _insert_before(p, r, end)
    ppr = p.find(A + "pPr")
    if ppr is not None:
        if ppr.get("algn") == "just":
            ppr.set("algn", "l")
        ppr.set("indent", "0")


def _insert_before(parent, el, ref):
    if ref is None:
        parent.append(el)
    else:
        ref.addprevious(el)


def translate_shape(shape, texts):
    paras = [p for p in shape._element.iter(A + "p") if "".join(t.text or "" for t in p.iter(A + "t")).strip()]
    assert len(paras) == len(texts), (shape.shape_id, len(paras), len(texts))
    for p, text in zip(paras, texts):
        set_paragraph(p, text)
        if shape.shape_id in CENTER:
            p.find(A + "pPr").set("algn", "ctr")
    if shape.shape_id in TIGHT:
        bp = shape._element.find(".//" + A + "bodyPr")
        bp.set("lIns", TIGHT_INS)
        bp.set("rIns", TIGHT_INS)


def widen_title(group, shape):
    xfrm = group._element.find(".//" + A + "xfrm")
    scale = int(xfrm.find(A + "chExt").get("cx")) / int(xfrm.find(A + "ext").get("cx"))
    new_w = int(TITLE_WIDTH_IN * 914400 * scale)
    shape.left = shape.left - (new_w - shape.width) // 2
    shape.width = new_w


def place_above_arrow(shape, connector, at_bottom):
    arrow_y = connector.top + connector.height if at_bottom else connector.top
    shape.height = ABOVE_ARROW_H
    shape.top = arrow_y - ABOVE_ARROW_H
    bp = shape._element.find(".//" + A + "bodyPr")
    bp.set("anchor", "b")
    bp.set("bIns", TIGHT_INS)
    bp.set("lIns", TIGHT_INS)
    bp.set("rIns", TIGHT_INS)


def clean_screenshot(pic):
    part = pic.part.related_part(pic._element.blipFill.blip.rEmbed)
    im = Image.open(io.BytesIO(part.blob)).convert("RGB")
    bg = im.getpixel((300, 60))
    px = im.load()
    for (x0, x1), (y0, y1) in IMG_BANDS:
        for y in range(y0, y1 + 1):
            for x in range(x0, x1 + 1):
                if sum(abs(a - b) for a, b in zip(px[x, y], bg)) > IMG_INK:
                    px[x, y] = bg
    buf = io.BytesIO()
    im.save(buf, "PNG")
    part._blob = buf.getvalue()


def main():
    prs = Presentation(SRC)
    slide = prs.slides[0]
    for shape in slide.shapes:
        if shape.shape_type == 13:
            clean_screenshot(shape)
        elif shape.shape_type == 6:
            seen = set()
            for child in shape.shapes:
                if child.shape_id in TEXT:
                    translate_shape(child, TEXT[child.shape_id])
                    seen.add(child.shape_id)
                    if child.shape_id == TITLE_ID:
                        widen_title(shape, child)
            assert seen == set(TEXT), set(TEXT) - seen
            children = {c.shape_id: c for c in shape.shapes}
            for label_id, (conn_id, at_bottom) in ABOVE_ARROW.items():
                place_above_arrow(children[label_id], children[conn_id], at_bottom)
    for layout_part in [prs.slide_master] + list(prs.slide_layouts):
        for shape in layout_part.shapes:
            if shape.has_text_frame and "保密文件" in shape.text_frame.text:
                for p in shape.text_frame.paragraphs:
                    if p.runs:
                        set_paragraph(p._p, "Confidential")
    prs.save(OUT)
    print(OUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
