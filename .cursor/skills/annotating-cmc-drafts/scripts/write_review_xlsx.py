#!/usr/bin/env python3
"""Write a review workbook (硬伤清单 / 三方核对) from a JSON spec, after checking every search string.

    python write_review_xlsx.py spec.json            # validate, then write
    python write_review_xlsx.py spec.json --check    # validate only
    python write_review_xlsx.py --example            # print an example spec

Spec:
{
  "output": "L2-Step1-MOR-硬伤清单_10-28-2026.xlsx",
  "sheets": [
    {
      "name": "MOR 硬伤",
      "source": "MOR_10282026.docx",          # optional; .docx or .xlsx; enables search-string checks
      "columns": ["类别", "严重程度", "位置", "内容（Word 中搜索）", "问题与改法"],
      "rows": [["内部数值/逻辑", "高", "3.46.7", "Add DIC(2.20 eq.)", "英文 2.20 eq，中文 2.00 eq ……"]]
    }
  ]
}

序号 is added automatically. Rows may also be objects keyed by column name.
The search column is the first column whose name starts with "内容". Paths are relative to the spec.
Sheets that have both 类别 and 严重程度 are sorted by category order, then 高 > 中 > 低 ("sort": false keeps input order).
Exits non-zero, and writes nothing, if a search string is missing from the source, spans lines, or contains a
non-breaking space; or if a severity is not 高/中/低.
"""

import argparse
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_lines import lines  # noqa: E402

CATEGORY_ORDER = ["与 PVP/参数表不一致", "与 PVP 不一致", "与参数表不一致", "内部数值/逻辑", "中英翻译", "笔误/格式"]
SEVERITY_ORDER = {"高": 0, "中": 1, "低": 2}
WIDTHS = {
    "序号": 6, "类别": 16, "严重程度": 9, "位置": 34, "问题与改法": 92, "批注内容": 90, "批注内容（问题与改法）": 90,
    "参数": 22, "判断": 46, "建议改哪里": 46, "说明": 52, "把握": 10, "问题": 140, "检查项": 18,
    "参数表行": 40, "我判断的 MOR 操作号": 46,
}
PREFIX_WIDTHS = [("内容", 36), ("参数表", 38), ("PVP", 40), ("MOR", 46)]

EXAMPLE = {
    "output": "L2-Step1-MOR-硬伤清单_10-28-2026.xlsx",
    "sheets": [
        {
            "name": "MOR 硬伤",
            "source": "MOR_10282026.docx",
            "columns": ["类别", "严重程度", "位置", "内容（Word 中搜索）", "问题与改法"],
            "rows": [
                [
                    "内部数值/逻辑", "高", "3.46.7 Orn11 DIC 加料（同句在 3.29.7 也有一处，那处中英一致）",
                    "Add DIC(2.20 eq.)",
                    "英文 2.20 eq，中文“(2.00 eq.)”。MOR 以中文为准，执行会按 2.00 eq，低于 PVP 表 7 DIC NOR 2.09–2.31 eq。中文改为 2.20 eq。",
                ]
            ],
        },
        {"name": "需要你确认", "columns": ["问题"], "rows": [["脱保护后 IPAc 按 10.0 V / 8.73 kg/kg，还是按 MOR 的 10.64 V / 9.29 kg/kg？"]]},
    ],
}


def width_for(column):
    if column in WIDTHS:
        return WIDTHS[column]
    for prefix, width in PREFIX_WIDTHS:
        if column.startswith(prefix):
            return width
    return 30


def style_sheet(ws, headers, widths, rows):
    head_fill = PatternFill("solid", fgColor="1B4F72")
    head_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    body_font = Font(name="Calibri", size=11)
    zebra = PatternFill("solid", fgColor="F7F9FB")
    wrap = Alignment(wrap_text=True, vertical="top")
    side = Side(style="thin", color="D0D7DE")
    border = Border(left=side, right=side, top=side, bottom=side)
    for col, text in enumerate(headers, 1):
        cell = ws.cell(1, col, text)
        cell.fill = head_fill
        cell.font = head_font
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, values in enumerate(rows, 2):
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            cell.alignment = wrap
            cell.font = body_font
            cell.border = border
            if r % 2 == 1:
                cell.fill = zebra
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def normalise_rows(sheet):
    columns = [c for c in sheet["columns"] if c != "序号"]
    rows = []
    for row in sheet["rows"]:
        if isinstance(row, dict):
            row = [row.get(c, "") for c in columns]
        if len(row) != len(columns):
            raise SystemExit(f"[{sheet['name']}] row has {len(row)} values, expected {len(columns)}: {row}")
        rows.append(["" if v is None else v for v in row])
    return columns, rows


def sort_rows(columns, rows):
    cat, sev = columns.index("类别"), columns.index("严重程度")

    def key(row):
        c = row[cat]
        return (CATEGORY_ORDER.index(c) if c in CATEGORY_ORDER else len(CATEGORY_ORDER), SEVERITY_ORDER.get(row[sev], 9))

    return sorted(rows, key=key)


def validate(sheet, columns, rows, base):
    errors, warnings = [], []
    if "严重程度" in columns:
        i = columns.index("严重程度")
        for n, row in enumerate(rows, 1):
            if row[i] not in SEVERITY_ORDER:
                errors.append(f"row {n}: 严重程度 must be 高/中/低, got {row[i]!r}")
    source = sheet.get("source")
    search_cols = [i for i, c in enumerate(columns) if c.startswith("内容")]
    if not source or not search_cols:
        return errors, warnings
    texts = [ln.text for ln in lines(base / source)]
    q_i = search_cols[0]
    loc_i = columns.index("位置") if "位置" in columns else None
    for n, row in enumerate(rows, 1):
        q = str(row[q_i])
        if not q.strip():
            errors.append(f"row {n}: empty search string")
            continue
        if "\xa0" in q:
            errors.append(f"row {n}: search string contains a non-breaking space: {q!r}")
            continue
        if "\n" in q:
            errors.append(f"row {n}: search string spans lines: {q!r}")
            continue
        hits = sum(t.count(q) for t in texts)
        if hits == 0:
            errors.append(f"row {n}: not found in {source}: {q!r}")
        elif hits > 1 and (loc_i is None or not str(row[loc_i]).strip()):
            warnings.append(f"row {n}: {q!r} occurs {hits} times; say which one in 位置")
    return errors, warnings


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", nargs="?")
    ap.add_argument("--check", action="store_true", help="validate only")
    ap.add_argument("--example", action="store_true", help="print an example spec and exit")
    args = ap.parse_args()
    if args.example:
        print(json.dumps(EXAMPLE, ensure_ascii=False, indent=2))
        return
    if not args.spec:
        ap.error("spec is required")
    spec_path = Path(args.spec).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    base = spec_path.parent

    prepared, failed = [], False
    for sheet in spec["sheets"]:
        columns, rows = normalise_rows(sheet)
        errors, warnings = validate(sheet, columns, rows, base)
        if sheet.get("sort", True) and {"类别", "严重程度"} <= set(columns):
            rows = sort_rows(columns, rows)
        for w in warnings:
            print(f"[{sheet['name']}] warning: {w}")
        for e in errors:
            print(f"[{sheet['name']}] ERROR: {e}")
        failed |= bool(errors)
        prepared.append((sheet, columns, rows))
    if failed:
        raise SystemExit("validation failed; nothing written")
    if args.check:
        print("all search strings found")
        return

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, columns, rows in prepared:
        ws = wb.create_sheet(sheet["name"][:31])
        headers = ["序号"] + columns
        widths = [6] + (sheet.get("widths") or [width_for(c) for c in columns])
        style_sheet(ws, headers, widths, [[i, *row] for i, row in enumerate(rows, 1)])
    out = base / spec["output"]
    wb.save(out)
    summary = []
    for sheet, columns, rows in prepared:
        if "严重程度" in columns:
            i = columns.index("严重程度")
            counts = {s: sum(r[i] == s for r in rows) for s in SEVERITY_ORDER}
            summary.append(f"{sheet['name']} {len(rows)} ({counts['高']} 高 / {counts['中']} 中 / {counts['低']} 低)")
        else:
            summary.append(f"{sheet['name']} {len(rows)}")
    print(f"wrote {out.name}: " + "; ".join(summary))


if __name__ == "__main__":
    main()
