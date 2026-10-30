#!/usr/bin/env python3
"""Mechanical checks on a bilingual MOR (English/中文 on one line). Output is a candidate list to screen by hand.

    python mor_auto_checks.py MOR.docx
    python mor_auto_checks.py MOR.docx --resin 48 --xlsx MOR-自动检查候选.xlsx

Checks
  encn        numbers differ between the English and Chinese halves of one line (Chinese prevails in an MOR)
  ref         "documented in Ins X" vs "记录于 Y" point to different steps
  placeholder X / xx / TBD left where a value belongs
  limits      ★ limits wider than the step's setpoint range; warning limits outside the ★ limits or setpoint
  time        coupling-time window upper bound above the step's NMT hold time; several different NMT values
  step        step number whose prefix does not match the section it sits in (e.g. 3.25.8 inside 3.28)
  docno       MOR numbers in body/headers/footers that differ from the majority
  weigh       standard charge outside its weighing range, or a range that is not about ±3 % / ±2 kg;
              kg/kg × resin ≠ kg (needs --resin); material-table kg/kg ≠ V × density
  confirm     English "confirm" translated as 执行 (should be 确认)
  spelling    English words not in a dictionary (only if pyspellchecker is installed)
"""

import argparse
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_lines import docx_lines  # noqa: E402

CJK = re.compile(r"[\u4e00-\u9fff]")
NUM = re.compile(r"(?<![A-Za-z\d.])\d+(?:\.\d+)?(?!\.?\d)(?!st|nd|rd|th)")
TEMP_UNIT = r"\s*(?:℃|°C|oC|ºC)"
ALLOWED_WORDS = {
    "fmoc", "dmf", "ipac", "mtbe", "dic", "dipea", "oxymapure", "pybop", "ctc", "trt", "tbu", "boc", "pbf", "orn", "asn",
    "phe", "gly", "ser", "arg", "val", "leu", "glu", "ala", "otbu", "nmt", "nlt", "nor", "par", "cpp", "cpps", "ipc",
    "ipcs", "cipcs", "kf", "lod", "ppq", "mor", "ins", "rpm", "eq", "qa", "pm", "oeb", "sy", "hplc", "sop", "dcs",
    "piperidine", "ninhydrin", "chloranil", "bromophenol", "reswelling", "diisopropylcarbodiimide", "hastelloy",
    "deprotection", "peptidyl", "mmol", "activator", "dispense", "dispensing",
}


def split_en_cn(text):
    """Return (english, chinese) for a single 'English/中文' line, else None."""
    if re.search(r"\t\d+\s*$", text):
        return None
    if len(re.findall(r"/\s*[\u4e00-\u9fff]", text)) != 1:
        return None
    m = CJK.search(text)
    i = text.rfind("/", 0, m.start())
    if i < 0:
        return None
    en, cn = text[:i], text[i + 1:]
    if CJK.search(en) or not re.search(r"[A-Za-z]{3}", en):
        return None
    if not re.match(r"[\s:：.]*(?:[\u4e00-\u9fff（(\d]|.?[\u4e00-\u9fff])", cn):
        return None
    return en, cn


def numbers(text):
    text = text.replace("\xa0", " ")
    text = re.sub(r"^\s*\d+(?:\.\d+)*\.?\s*", "", text)
    text = re.sub(r"[（(]\d+[)）]", "", text)
    text = re.sub(r"\b[A-Za-z][A-Za-z0-9]*\d[A-Za-z0-9\-]*", "", text)
    text = re.sub(r"(?<=[-‐’',])\d+|\d+(?=[-‐’',][A-Za-z\u4e00-\u9fff(])|\d+(?=H-)", "", text)
    return Counter(float(x) for x in NUM.findall(text))


def snippet(text, pos=None, size=40):
    """A verbatim piece of `text` without tabs, line breaks or non-breaking spaces, near `pos`."""
    pieces = [(m.start(), m.group()) for m in re.finditer(r"[^\t\n\xa0]+", text)]
    pieces = [(s, p) for s, p in pieces if p.strip()]
    if not pieces:
        return ""
    chosen = max(pieces, key=lambda sp: len(sp[1]))
    if pos is not None:
        for s, p in pieces:
            if s <= pos < s + len(p):
                chosen = (s, p)
                break
    start, piece = chosen
    if len(piece) <= size:
        return piece.strip()
    centre = (pos - start) if pos is not None and start <= pos < start + len(piece) else 0
    lo = max(0, min(centre - size // 2, len(piece) - size))
    return piece[lo:lo + size].strip()


class Walker:
    """Tracks section (3.28) and step (3.28.5) while walking body lines."""

    def __init__(self, body):
        self.body = body

    def __iter__(self):
        section = step = None
        for ln in self.body:
            t = ln.text
            m = re.match(r"^\s*(\d+\.\d+)\.\s*\t?[（(]", t) or re.match(r"^\s*(\d+\.\d+)\.?\s*\t?[A-Za-z]", t)
            if m and not re.match(r"^\s*\d+\.\d+\.\d+", t):
                section = m.group(1)
            m = re.match(r"^\s*(\d+\.\d+)\.(\d+)\.?\s+\S", t)
            if m:
                if m.group(2) == "1":
                    section = m.group(1)
                step = f"{m.group(1)}.{m.group(2)}"
            yield section, step, ln


def where(section, step, ln):
    loc = step or section or ""
    return f"{loc}（段落 {ln.index}）" if loc else f"段落 {ln.index}"


def bilingual_pairs(walk):
    """Yield (position, (english, chinese)) for 'English/中文' lines and for an English paragraph
    immediately followed by its Chinese paragraph."""
    for i, item in enumerate(walk):
        ln = item[2]
        pair = split_en_cn(ln.text)
        if pair:
            yield item, pair
            continue
        if CJK.search(ln.text) or len(re.findall(r"[A-Za-z]{3,}", ln.text)) < 3 or i + 1 >= len(walk):
            continue
        nxt = walk[i + 1][2]
        if nxt.index == ln.index + 1 and CJK.search(nxt.text) and not split_en_cn(nxt.text):
            if len(CJK.findall(nxt.text)) >= 4:
                yield item, (ln.text, nxt.text)


def check_encn(walk):
    """Identical differences in repeated cycles are merged into one candidate listing every step."""
    groups = {}
    fmt = lambda xs: ", ".join(f"{x:g}" for x in xs) or "无"
    for (section, step, ln), pair in bilingual_pairs(walk):
        en, cn = numbers(pair[0]), numbers(pair[1])
        if not (en and cn) or en == cn:
            continue
        only_en = tuple(sorted((en - cn).elements()))
        only_cn = tuple(sorted((cn - en).elements()))
        key = (only_en, only_cn, re.sub(r"[\d\s]+", "", pair[1])[:20])
        if key not in groups:
            pos = len(pair[0]) + 1
            groups[key] = [snippet(ln.text, pos), []]
        groups[key][1].append(step or section or f"段落 {ln.index}")
    out = []
    for (only_en, only_cn, _), (q, locs) in groups.items():
        where_ = "、".join(dict.fromkeys(locs))
        note = f"英文独有 {fmt(only_en)}；中文独有 {fmt(only_cn)}。"
        note += "英文漏了数值。" if not only_en else "MOR 以中文为准，核对哪边对。"
        out.append(("encn", f"{where_}（{len(locs)} 处）" if len(locs) > 1 else where_, q, note))
    return out


def check_refs(walk):
    out = []
    for section, step, ln in walk:
        en = re.findall(r"Ins\.?\s*(\d+(?:\.\d+)+)", ln.text)
        cn = re.findall(r"(?:记录于|操作|见)\s*(?:操作)?\s*(\d+(?:\.\d+)+)", ln.text)
        if en and cn and en != cn:
            out.append(("ref", where(section, step, ln), snippet(ln.text, ln.text.find("Ins")),
                        f"英文引用 {', '.join(en)}，中文引用 {', '.join(cn)}。"))
    return out


def check_placeholders(walk):
    out = []
    pat = re.compile(r"(?<![A-Za-z])(X{1,3}|x{1,3})\s*(?:℃|oC|°C|kg|min|h\b|%|V\b)|\bTBD\b|待定|\?\?")
    for section, step, ln in walk:
        m = pat.search(ln.text)
        if m:
            out.append(("placeholder", where(section, step, ln), snippet(ln.text, m.start()), f"占位符 “{m.group()}” 未填。"))
    return out


LIMIT = re.compile(
    r"(?:Lower|Low)\s+limit\s*/\s*低限\s*[:：]\s*([\dXx.]+)" + r"\s*(?:℃|oC|°C)?\s*[，,]\s*High\s+limit\s*/\s*高限\s*[:：]\s*([\dXx.]+)"
)
SETPOINT = re.compile(r"(\d+(?:\.\d+)?)\s*±\s*(\d+(?:\.\d+)?)" + TEMP_UNIT + r"|(\d+(?:\.\d+)?)\s*[~～–-]\s*(\d+(?:\.\d+)?)" + TEMP_UNIT)


def check_limits(walk):
    by_section = defaultdict(lambda: {"set": [], "star": [], "warn": []})
    prev = ""
    for section, step, ln in walk:
        t = ln.text.replace("\xa0", " ")
        for m in SETPOINT.finditer(t):
            if m.group(1):
                c, d = float(m.group(1)), float(m.group(2))
                by_section[section]["set"].append((c - d, c + d))
            else:
                by_section[section]["set"].append((float(m.group(3)), float(m.group(4))))
        m = LIMIT.search(t)
        if m:
            kind = "warn" if "warning" in prev.lower() or "预警" in prev else "star"
            by_section[section][kind].append((m.group(1), m.group(2), step, ln))
        prev = t
    out = []
    for section, d in by_section.items():
        if not d["set"]:
            continue
        lo, hi = min(a for a, _ in d["set"]), max(b for _, b in d["set"])
        stars = []
        for a, b, step, ln in d["star"]:
            try:
                a, b = float(a), float(b)
            except ValueError:
                continue
            stars.append((a, b))
            if a < lo or b > hi:
                out.append(("limits", where(section, step, ln), snippet(ln.text, ln.text.find("★")),
                            f"★ 限 {a:g}–{b:g} ℃ 比本节设定 {lo:g}–{hi:g} ℃ 宽；是否抄了别的步骤的 PAR？"))
        s_lo = max([a for a, _ in stars] + [lo]) if stars else lo
        s_hi = min([b for _, b in stars] + [hi]) if stars else hi
        for a, b, step, ln in d["warn"]:
            try:
                a, b = float(a), float(b)
            except ValueError:
                continue
            if a < s_lo or b > s_hi:
                out.append(("limits", where(section, step, ln), snippet(ln.text, ln.text.find("limit")),
                            f"超温预警 {a:g}–{b:g} ℃ 不在 NOR {s_lo:g}–{s_hi:g} ℃ 内，预警会晚于越界。"))
    return out


def check_time(walk):
    nmt = defaultdict(list)
    windows = defaultdict(list)
    all_nmt = Counter()
    for section, step, ln in walk:
        t = ln.text
        for m in re.finditer(r"NMT\)?\s*(\d+(?:\.\d+)?)\s*(min|h)\b", t):
            minutes = float(m.group(1)) * (60 if m.group(2) == "h" else 1)
            nmt[section].append((minutes, step, ln))
            all_nmt[f"{m.group(1)} {m.group(2)}"] += 1
        for m in re.finditer(r"(\d+(?:\.\d+)?)\s*h\s*≤\s*[A-Za-z ]+≤\s*(\d+(?:\.\d+)?)\s*h|(\d+(?:\.\d+)?)\s*[~～–-]\s*(\d+(?:\.\d+)?)\s*h\b", t):
            upper = float(m.group(2) or m.group(4)) * 60
            if upper >= 120:
                windows[section].append((upper, m, step, ln))
    out = []
    for section, ws in windows.items():
        limits = [x for x, _, _ in nmt.get(section, [])]
        if not limits:
            continue
        cap = min(limits)
        for upper, m, step, ln in ws:
            if upper > cap:
                out.append(("time", where(section, step, ln), snippet(ln.text, m.start()),
                            f"时间窗口上限 {upper / 60:g} h（{upper:g} min）超过本节 NMT {cap:g} min。"))
    if len(all_nmt) > 1:
        out.append(("time", "全文", "NMT", "全文 NMT 取值不止一个：" + "，".join(f"{k}×{v}" for k, v in all_nmt.most_common())
                    + "。逐个确认是不同步骤的不同限值还是笔误。"))
    return out


def check_steps(body):
    out = []
    section = None
    for ln in body:
        t = ln.text
        m = re.match(r"^\s*(\d+\.\d+)\.\s*\t?[（(]", t)
        if m:
            section = m.group(1)
            continue
        m = re.match(r"^\s*(\d+\.\d+)\.(\d+)\.?\s+\S", t)
        if not m:
            continue
        if m.group(2) == "1":
            section = m.group(1)
        elif section and m.group(1) != section:
            out.append(("step", f"{section} 节（段落 {ln.index}）", snippet(t, 0),
                        f"步骤号 {m.group(1)}.{m.group(2)} 出现在 {section} 节，应为 {section}.{m.group(2)}。"))
    return out


def check_docno(all_lines):
    found = defaultdict(list)
    for ln in all_lines:
        for m in re.finditer(r"MOR-[A-Z]{2}[\dA-Z]+(?:\.\d+)?", ln.text):
            found[m.group()].append(ln)
    if len(found) <= 1:
        return []
    main = max(found, key=lambda k: len(found[k]))
    out = []
    for no, lns in found.items():
        if no == main:
            continue
        parts = Counter(ln.part for ln in lns)
        out.append(("docno", "、".join(f"{p}×{c}" for p, c in parts.items()), no,
                    f"MOR 号 {no} 与主号 {main} 不同（{len(lns)} 处）。"))
    return out


def _f(s):
    try:
        return float(str(s).replace(",", ""))
    except ValueError:
        return None


def tolerance_ok(target, lo, hi):
    """Symmetric about the target and no wider than ±3 % (or ±2 kg for small 1 V rinses)."""
    symmetric = abs((lo + hi) / 2 - target) <= 0.005 * target + 0.02
    return symmetric and (hi - lo) / 2 <= max(0.033 * target, 2.02)


def check_weighing(walk, resin):
    out = []
    pat = re.compile(
        r"(\d+(?:\.\d+)?)\s*kg(?!\s*/\s*kg)(?:(?!\d+(?:\.\d+)?\s*kg).){0,80}?(?:weighing range|称量范围)[^0-9]{0,15}"
        r"(\d+(?:\.\d+)?)\s*(?:kg)?\s*[~～–-]\s*(\d+(?:\.\d+)?)\s*kg"
    )
    for section, step, ln in walk:
        t = ln.text.replace("\xa0", " ")
        for m in pat.finditer(t):
            x, lo, hi = float(m.group(1)), float(m.group(2)), float(m.group(3))
            problems = []
            if not lo <= x <= hi:
                problems.append(f"标准量 {x:g} kg 不在称量范围 {lo:g}–{hi:g} kg 内")
            elif not tolerance_ok(x, lo, hi):
                problems.append(f"称量范围 {lo:g}–{hi:g} kg 不对称或宽于 {x:g} kg 的 ±3% / ±2 kg")
            if problems:
                out.append(("weigh", where(section, step, ln), snippet(ln.text, m.start()), "；".join(problems) + "。"))
        if resin:
            for m in re.finditer(r"(\d+(?:\.\d+)?)\s*kg\b[^/]{0,60}?\(\s*(\d+(?:\.\d+)?)\s*kg/kg", t):
                x, r = float(m.group(1)), float(m.group(2))
                if abs(x - r * resin) > max(0.01 * x, 0.05):
                    out.append(("weigh", where(section, step, ln), snippet(ln.text, m.start()),
                                f"{r:g} kg/kg × {resin:g} kg = {r * resin:.2f} kg，正文写 {x:g} kg。"))
    return out


def check_material_tables(path, resin):
    import docx

    out = []
    for t_i, table in enumerate(docx.Document(path).tables):
        if not table.rows:
            continue
        header = [c.text for c in table.rows[0].cells]
        col = {}
        for i, h in enumerate(header):
            if "标准投料量" in h:
                col["std"] = i
            elif "可接受称量范围" in h:
                col["range"] = i
            elif "Density" in h or "密度" in h:
                col["density"] = i
            elif "Volume Ratio" in h:
                col["v"] = i
            elif "Qty Ratio" in h:
                col["q"] = i
        if not {"std", "range"} <= set(col):
            continue
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            op, name = cells[0], cells[1].replace("\n", " ")[:30]
            std = _f(re.sub(r"\s*\(.*", "", cells[col["std"]]))
            rng = re.match(r"([\d.]+)\s*[-~～–]\s*([\d.]+)", cells[col["range"]])
            loc = f"原料信息表（表 {t_i}）操作 {op} {name}"
            if std and rng:
                lo, hi = float(rng.group(1)), float(rng.group(2))
                if not lo <= std <= hi:
                    out.append(("weigh", loc, cells[col["range"]], f"标准投料量 {std:g} kg 不在称量范围 {lo:g}–{hi:g} kg 内。"))
                elif not tolerance_ok(std, lo, hi):
                    out.append(("weigh", loc, cells[col["range"]], f"称量范围 {lo:g}–{hi:g} kg 不对称或宽于 {std:g} kg 的 ±3% / ±2 kg。"))
            v, d, q = (_f(cells[col[k]]) if k in col else None for k in ("v", "density", "q"))
            if v and d and q and abs(v * d - q) > 0.01:
                out.append(("weigh", loc, cells[col["q"]], f"{v:g} V × 密度 {d:g} = {v * d:.3f} kg/kg，表中写 {q:g}。"))
            if resin and q and std and abs(q * resin - std) > max(0.005 * std, 0.05):
                out.append(("weigh", loc, cells[col["std"]], f"{q:g} kg/kg × {resin:g} kg = {q * resin:.2f} kg，表中写 {std:g} kg。"))
    return out


def check_confirm(walk):
    hits = []
    for item, (en, cn) in bilingual_pairs(walk):
        if re.search(r"\bconfirm", en, re.I) and "执行" in cn and "确认" not in cn:
            hits.append(item)
    seen = {id(item[2]) for item in hits}
    hits += [item for item in walk if "工艺性能执行" in item[2].text and id(item[2]) not in seen]
    if not hits:
        return []
    section, step, ln = hits[0]
    pos = ln.text.find("执行")
    return [("confirm", f"共 {len(hits)} 处，首处 {where(section, step, ln)}", snippet(ln.text, pos),
             "英文 confirm 译成“执行”，应为“确认”；“工艺性能执行”应为“工艺性能确认”。建议全文查找替换后人工过一遍。")]


def check_spelling(body):
    try:
        from spellchecker import SpellChecker
    except ImportError:
        return []
    sc = SpellChecker()
    words = Counter()
    first = {}
    for ln in body:
        pair = split_en_cn(ln.text)
        en = pair[0] if pair else ("" if CJK.search(ln.text) else ln.text)
        for w in re.findall(r"\b[a-z]{4,}\b", en):
            if w not in ALLOWED_WORDS:
                words[w] += 1
                first.setdefault(w, ln)
    unknown = sc.unknown(words)
    out = []
    for w in sorted(unknown, key=lambda w: -words[w]):
        ln = first[w]
        out.append(("spelling", f"段落 {ln.index}", w, f"疑似拼写错误 “{w}”（{words[w]} 处），建议 “{sc.correction(w)}”。专业词请忽略。"))
    return out


def run(path, resin=None):
    all_lines = docx_lines(path)
    body = [ln for ln in all_lines if ln.part == "document"]
    walk = list(Walker(body))
    results = []
    results += check_encn(walk)
    results += check_refs(walk)
    results += check_placeholders(walk)
    results += check_limits(walk)
    results += check_time(walk)
    results += check_steps(body)
    results += check_docno(all_lines)
    results += check_weighing(walk, resin)
    results += check_material_tables(path, resin)
    results += check_confirm(walk)
    results += check_spelling(body)
    return results


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("mor")
    ap.add_argument("--resin", type=float, help="limiting-material charge in kg, for kg/kg × resin checks (L2 Step 1: 48)")
    ap.add_argument("--xlsx", help="also write the candidates to this workbook")
    ap.add_argument("--only", help="comma-separated check names to keep, e.g. encn,ref,step")
    args = ap.parse_args()
    results = run(args.mor, args.resin)
    if args.only:
        keep = set(args.only.split(","))
        results = [r for r in results if r[0] in keep]
    counts = Counter(r[0] for r in results)
    print("candidates: " + ", ".join(f"{k} {v}" for k, v in counts.items()))
    for kind, loc, q, note in results:
        print(f"[{kind}] {loc} | {q} | {note}")
    if args.xlsx:
        import openpyxl
        from write_review_xlsx import style_sheet

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "自动检查候选（待人工筛）"
        style_sheet(ws, ["序号", "检查项", "位置", "内容（Word 中搜索）", "说明"], [6, 12, 34, 40, 90],
                    [(i, *r) for i, r in enumerate(results, 1)])
        wb.save(args.xlsx)
        print(f"wrote {args.xlsx}")


if __name__ == "__main__":
    main()
